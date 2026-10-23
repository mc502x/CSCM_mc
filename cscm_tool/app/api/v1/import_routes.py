"""FR-048a-d, EPIC-14, I-05, I-32. See docs/artifacts/openapi.yaml /imports*."""

import csv
import io

from flask import Response, jsonify, request, send_file

from app.api.v1 import api_v1_bp
from app.api.v1._errors import error_response
from app.api.v1.serializers import import_job_payload
from app.auth.decorators import current_user, require_role
from app.domain.errors import DomainError
from app.services.import_service import ImportService

_IMPORT_ROLES = ("ADMINISTRATOR", "ENGINEER")
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _flag(name: str) -> bool:
    return (request.form.get(name) or request.args.get(name) or "").lower() in (
        "1",
        "true",
        "yes",
        "on",
    )


@api_v1_bp.get("/imports")
@require_role(*_IMPORT_ROLES)
def list_imports():
    jobs = ImportService.list_jobs()
    return jsonify([import_job_payload(j) for j in jobs])


@api_v1_bp.post("/imports")
@require_role(*_IMPORT_ROLES)
def create_import():
    """Upload a library file. Form fields: `file`; `mode` CREATE (default) or UPDATE (match on
    legacy_reference_number); `dry_run=true` validates without creating anything (200 with the
    per-row report); `async=true` forces background processing (large files get it automatically).
    """
    upload = request.files.get("file")
    if upload is None or not upload.filename:
        return jsonify(error="validation_error", message="file is required"), 422
    mode = request.form.get("mode") or request.args.get("mode") or "CREATE"

    try:
        if _flag("dry_run"):
            return jsonify(
                ImportService.dry_run(upload.filename, upload.read(), current_user, mode)
            )
        job = ImportService.process_upload(
            upload.filename,
            upload.read(),
            current_user,
            mode=mode,
            asynchronous=True if _flag("async") else None,
        )
    except (DomainError, ValueError) as exc:
        return error_response(exc)

    return jsonify(import_job_payload(job)), 202


@api_v1_bp.get("/imports/template")
@require_role(*_IMPORT_ROLES)
def import_template():
    """I-05b: the exact headers (CSV) or headers + instructions + lookup codes (XLSX)."""
    fmt = (request.args.get("format") or "xlsx").lower()
    if fmt == "csv":
        return Response(
            ImportService.template_csv(),
            mimetype="text/csv",
            headers={"Content-Disposition": "attachment; filename=cscm_import_template.csv"},
        )
    if fmt == "xlsx":
        return Response(
            ImportService.template_xlsx(),
            mimetype=_XLSX,
            headers={"Content-Disposition": "attachment; filename=cscm_import_template.xlsx"},
        )
    return jsonify(error="validation_error", message="format must be csv or xlsx"), 422


def _legacy_numbers_from_upload(upload) -> list[str]:
    name = (upload.filename or "").lower()
    content = upload.read()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        header = [str(c).strip() if c is not None else "" for c in (rows[0] if rows else [])]
        index = (
            header.index("legacy_reference_number") if "legacy_reference_number" in header else 0
        )
        return [str(r[index]).strip() for r in rows[1:] if r and r[index] is not None]
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    if not rows:
        return []
    header = [c.strip() for c in rows[0]]
    index = header.index("legacy_reference_number") if "legacy_reference_number" in header else 0
    return [r[index].strip() for r in rows[1:] if len(r) > index and r[index].strip()]


@api_v1_bp.get("/imports/reconciliation")
@require_role(*_IMPORT_ROLES)
def import_reconciliation():
    """I-05d: imported / in review / released / retired, vs. an expected total."""
    expected_total = request.args.get("expected_total", type=int)
    return jsonify(ImportService.reconciliation(expected_total=expected_total))


@api_v1_bp.post("/imports/reconciliation")
@require_role(*_IMPORT_ROLES)
def import_reconciliation_with_master_list():
    """Same report, with the legacy master list uploaded (`file`: CSV/XLSX with a
    legacy_reference_number column) so the missing numbers are named."""
    upload = request.files.get("file")
    if upload is None:
        return jsonify(error="validation_error", message="file is required"), 422
    try:
        numbers = _legacy_numbers_from_upload(upload)
    except (ValueError, UnicodeDecodeError, OSError) as exc:
        return (
            jsonify(error="validation_error", message=f"The master list could not be read: {exc}"),
            422,
        )
    return jsonify(ImportService.reconciliation(expected_legacy_references=numbers))


@api_v1_bp.get("/imports/<int:import_job_id>")
@require_role(*_IMPORT_ROLES)
def get_import(import_job_id: int):
    job = ImportService.get(import_job_id)
    if job is None:
        return jsonify(error="not_found", message="Import job not found"), 404
    return jsonify(import_job_payload(job))


@api_v1_bp.get("/imports/<int:import_job_id>/errors")
@require_role(*_IMPORT_ROLES)
def get_import_errors(import_job_id: int):
    job = ImportService.get(import_job_id)
    if job is None:
        return jsonify(error="not_found", message="Import job not found"), 404
    if not job.error_report_reference:
        return jsonify(error="not_found", message="No error report for this import job"), 404
    if (request.args.get("format") or "csv").lower() == "xlsx":
        if not job.error_report_xlsx_reference:
            return jsonify(error="not_found", message="No XLSX error report for this job"), 404
        return send_file(
            job.error_report_xlsx_reference,
            mimetype=_XLSX,
            as_attachment=True,
            download_name=f"import_{import_job_id}_errors.xlsx",
        )
    return send_file(
        job.error_report_reference,
        mimetype="text/csv",
        as_attachment=True,
        download_name=f"import_{import_job_id}_errors.csv",
    )
