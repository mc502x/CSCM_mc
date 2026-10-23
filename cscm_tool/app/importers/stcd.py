"""Reader for the generated status code programs (`PROGRAM StCd_<n>`, IEC 61131-3 Structured Text).

Each file written by StatuscodeGeneratorEuros holds one or more programs. Per program the reader
extracts, **without executing or interpreting anything** (it is a text scanner, never a compiler):

  * the header comment: legacy number and name (`StCd: 3000, 'Umrichter nicht bereit'`), the
    source database and the generator,
  * the attribute block between `(* Attribute aus DB *)` and `FirstRun := FALSE;`: the number
    (`StCdArUi[Idx]`), the repeat code and limits, the operating-mode mask, the alarm and
    availability settings, the per-subsystem `*PrgArUs` programs and the read/write level of every
    user group,
  * the set condition (`StatusSetArGBo[Idx] := ...;`), the set and reset delays, the
    `StCd<n>DisableGpBo` parameter and the other StCd programs it references,
  * the block identification `PrgIdx_GUi := <n>` (docs/06-data-dictionary.md section 1).

Main Controller (MC) numbers in the files are the current ones and the old number is only in the
header (`legacy_number`). The numbers in Hub Controller (HC) files are the *legacy* HC numbers;
`apply_mapping` replaces them with the new number from the numbering workbook
(docs/artifacts/hc-numbering-proposal.xlsx, sheet "Mapping"). Both give a **legacy -> new** pair per
program; the new set holds more codes than the legacy one, so a program without a legacy number
simply has none (the mapping is not meant to be inverted, 06-data-dictionary.md section 5a).

scripts/read_stcd.py is the command-line front end; the application
(app/services/st_import_service.py) uses the same functions. Attributes the reader does not know
by name are kept under `other`."""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

SUFFIXES = {".st", ".scl", ".exp", ".txt"}
HC_FROM = 10000
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_PROGRAMS_PER_FILE = 5000

_PROGRAM_RE = re.compile(r"\bPROGRAM\s+StCd_(\d+)\b", re.IGNORECASE)
_END_PROGRAM_RE = re.compile(r"\bEND_PROGRAM\b\s*;?", re.IGNORECASE)
_HEADER_NAME_RE = re.compile(r"StCd:\s*(\d+)\s*,\s*'(.*)'[ \t\r]*$", re.MULTILINE)
_HEADER_DB_RE = re.compile(r"Created by database \(([^)]*)\)")
_HEADER_GEN_RE = re.compile(r"StCd generator:\s*(\S+)")
_INIT_RE = re.compile(r"\(\*\s*Attribute aus DB\s*\*\)(.*?)\bFirstRun\s*:=\s*FALSE\s*;", re.DOTALL)
_ATTR_RE = re.compile(
    r"(?P<name>\w+)\[Idx(?:\s*,\s*(?P<sub>\d+))?\]\s*:=\s*(?P<value>[^;]+);"
    r"[ \t]*(?:\(\*\s*(?P<note>.*?)\s*\*\))?"
)
_SNAP_RE = re.compile(r"\bStCdWithSnap_Bo\s*:=\s*(TRUE|FALSE)\s*;", re.IGNORECASE)
_PRGIDX_RE = re.compile(r"^\s*PrgIdx_GUi\s*:=\s*(\d+)\s*;", re.MULTILINE)
_SET_RE = re.compile(r"\bStatusSetArGBo\[Idx\]\s*:=\s*(.*?);", re.DOTALL)
_SET_TIMER_RE = re.compile(r"\bSetTimer\s*\(\s*PT\s*:=\s*([^,)]+)")
_RESET_TIMER_RE = re.compile(r"\bResetTimer\s*\(\s*PT\s*:=\s*([^,)]+)")
_DISABLE_RE = re.compile(r"\bStCd\d+DisableGpBo\b")
_OTHER_PROGRAM_RE = re.compile(r"\bStCd_(\d+)\b")

# generator attribute -> field
_SCALARS = {
    "RepeatStCdArUi": "repeat_code",
    "MaxRepeatNbArBy": "max_repeats",
    "RepeatTimeArDw": "repeat_time_s",
    "TaskCallFrequencyArBy": "task_call_frequency",
    "ZeroLoadOperation_ArBo": "zero_load_operation",
    "AlarmArBo": "alarm",
    "AlarmTypArBy": "alarm_type",
    "AvailabilityGroupArUs": "availability_group",
    "AvailabilityCalcArBo": "availability_calc",
}
_TIME_UNITS = {"d": 86_400_000, "h": 3_600_000, "m": 60_000, "s": 1_000, "ms": 1}


class StParseError(ValueError):
    """A file cannot be read as generated status code programs."""


@dataclass
class StatusCode:
    """One program. `identifier` is already in the tool's format (`StCd-01000`)."""

    identifier: str
    number: int
    controller: str  # "MC" | "HC"
    name: str = ""
    legacy_number: int | None = None
    source_database: str = ""
    generator: str = ""
    file: str = ""
    line: int = 0
    prg_idx: int | None = None
    proposal_1: int | None = None
    proposal_2: int | None = None
    logical_node: str = ""
    hc_subgroup: str = ""
    mapping_description: str = ""
    mapping_note: str = ""
    repeat_code: int | None = None
    repeat_code_legacy: int | None = None
    max_repeats: int | None = None
    repeat_time_s: int | None = None
    task_call_frequency: int | None = None
    op_mode_bits: str = ""
    zero_load_operation: bool | None = None
    snapshot: bool | None = None
    alarm: bool | None = None
    alarm_type: int | None = None
    availability_group: int | None = None
    availability_calc: bool | None = None
    set_condition: str = ""
    set_delay: str = ""
    set_delay_ms: int | None = None
    reset_delay: str = ""
    reset_delay_ms: int | None = None
    disable_variable: str = ""
    depends_on: list[int] = field(default_factory=list)
    programs: dict[str, int | str | bool] = field(default_factory=dict)
    user_groups: dict[int, dict[str, int]] = field(default_factory=dict)
    other: dict[str, int | str | bool] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return asdict(self)


def identifier_for(number: int) -> str:
    """The tool's identifier format: five digits, zero padded (`StCd-01000`, `StCd-14000`)."""
    return f"StCd-{number:05d}"


def _strip_comments(text: str) -> str:
    """Returns the text without (* *), /* */ and // comments, keeping line breaks."""
    out: list[str] = []
    closer: str | None = None
    for line in text.splitlines():
        i, n = 0, len(line)
        while i < n:
            if closer:
                j = line.find(closer, i)
                if j < 0:
                    i = n
                else:
                    i, closer = j + len(closer), None
                continue
            two = line[i : i + 2]
            if two == "//":
                break
            if two in ("(*", "/*"):
                closer = "*)" if two == "(*" else "*/"
                i += 2
                continue
            if line[i] == "'":
                j = i + 1
                while j < n and line[j] != "'":
                    j += 2 if line[j] == "$" else 1
                out.append(line[i : j + 1])
                i = j + 1
                continue
            out.append(line[i])
            i += 1
        out.append("\n")
    return "".join(out)


def _value(text: str) -> int | str | bool:
    text = text.strip()
    if text.upper() in ("TRUE", "FALSE"):
        return text.upper() == "TRUE"
    if re.fullmatch(r"\d[\d_]*", text):
        return int(text.replace("_", ""))
    return text


def _time_ms(text: str) -> int | None:
    """Milliseconds of a time literal such as T#10ms or T#1m30s; None for a variable."""
    m = re.fullmatch(r"(?i)t(?:ime)?#((?:\d+(?:\.\d+)?(?:ms|d|h|m|s))+)", text.strip())
    if not m:
        return None
    parts = re.findall(r"(\d+(?:\.\d+)?)(ms|d|h|m|s)", m.group(1).lower())
    return round(sum(float(n) * _TIME_UNITS[u] for n, u in parts))


def _parse_program(
    header: str, body: str, program: int, line: int, source: str, hc: bool
) -> StatusCode:
    attrs: dict[str, tuple[int | str | bool, str | None]] = {}
    groups: dict[int, dict[str, int]] = {}
    programs: dict[str, int | str | bool] = {}
    other: dict[str, int | str | bool] = {}
    warnings: list[str] = []
    number: int | None = None
    repeat_note: int | None = None
    op_mode = ""

    init = _INIT_RE.search(body)
    for attr in _ATTR_RE.finditer(init.group(1) if init else ""):
        name, sub, value, note = attr.group("name", "sub", "value", "note")
        value = _value(value)
        if name == "StCdArUi":
            number = int(value) if isinstance(value, int) else number
        elif name == "OpModeArDw":
            op_mode = str(value).split("#", 1)[-1]
        elif name in ("UserGroupReadArBy", "UserGroupWriteArBy") and sub is not None:
            groups.setdefault(int(sub), {})["read" if name == "UserGroupReadArBy" else "write"] = (
                int(value)
            )
        elif name in _SCALARS:
            attrs[_SCALARS[name]] = (value, note)
        elif name.endswith("PrgArUs"):
            programs[name.removesuffix("PrgArUs")] = value
        else:
            other[name] = value
        if name == "RepeatStCdArUi" and note and note.isdigit():
            repeat_note = int(note)

    if number is None:
        warnings.append("no StCdArUi[Idx] assignment; the number of the program name is used")
        number = program
    elif number != program:
        warnings.append(f"program StCd_{program} sets StCdArUi to {number}")

    code = StatusCode(identifier_for(number), number, "HC" if hc else "MC", file=source, line=line)
    for key, (value, _) in attrs.items():
        setattr(code, key, value)
    code.repeat_code_legacy = repeat_note
    code.op_mode_bits = op_mode
    code.programs, code.other = programs, other
    code.user_groups = dict(sorted(groups.items()))
    snap = _SNAP_RE.search(body)
    code.snapshot = snap.group(1).upper() == "TRUE" if snap else None
    if idx := _PRGIDX_RE.search(body):
        code.prg_idx = int(idx.group(1))

    if m := _HEADER_NAME_RE.search(header):
        code.name = m.group(2)
        legacy = int(m.group(1))
        code.legacy_number = legacy if legacy != number else None
    if m := _HEADER_DB_RE.search(header):
        code.source_database = m.group(1)
    if m := _HEADER_GEN_RE.search(header):
        code.generator = m.group(1)

    logic = _strip_comments(body)
    if m := _SET_RE.search(logic):
        code.set_condition = " ".join(m.group(1).split())
    if m := _SET_TIMER_RE.search(logic):
        code.set_delay = m.group(1).strip()
        code.set_delay_ms = _time_ms(code.set_delay)
    if m := _RESET_TIMER_RE.search(logic):
        code.reset_delay = m.group(1).strip()
        code.reset_delay_ms = _time_ms(code.reset_delay)
    if m := _DISABLE_RE.search(logic):
        code.disable_variable = m.group(0)
    code.depends_on = sorted({int(n) for n in _OTHER_PROGRAM_RE.findall(logic)} - {program})
    code.warnings = warnings
    return code


def parse_text(
    text: str, source: str, hc: bool | None = None, hc_from: int = HC_FROM
) -> list[StatusCode]:
    """Returns the status codes defined by the StCd programs of one source text. `hc` True
    marks them as Hub Controller codes, None decides by `hc_from`. The number of an HC code
    is still the legacy one; see `apply_mapping`."""
    codes: list[StatusCode] = []
    start = 0
    for end in _END_PROGRAM_RE.finditer(text):
        segment = text[start : end.end()]
        m = _PROGRAM_RE.search(segment)
        if m:
            if len(codes) >= MAX_PROGRAMS_PER_FILE:
                raise StParseError(f"{source}: more than {MAX_PROGRAMS_PER_FILE} programs")
            line = text.count("\n", 0, start + m.start()) + 1
            codes.append(
                _parse_program(
                    segment[: m.start()],
                    segment[m.start() :],
                    int(m.group(1)),
                    line,
                    source,
                    bool(hc) if hc is not None else int(m.group(1)) >= hc_from,
                )
            )
        start = end.end()
    return codes


def decode(raw: bytes) -> str:
    """UTF-8 first, then CP1252 (the generator's encoding); never raises."""
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def parse_bytes(
    raw: bytes, source: str, hc: bool | None = None, hc_from: int = HC_FROM
) -> tuple[list[StatusCode], str]:
    """Parse an uploaded file; returns (programs, sha256 of the bytes). Raises StParseError for a
    file that is too large."""
    if len(raw) > MAX_FILE_BYTES:
        raise StParseError(
            f"{source}: {len(raw) // 1024} KB is larger than the {MAX_FILE_BYTES // 1024} KB limit"
        )
    return parse_text(decode(raw), source, hc, hc_from), hashlib.sha256(raw).hexdigest()


# ------------------------------------------------------------------ legacy Hub numbers
def load_mapping(source: Path | bytes) -> dict[int, tuple]:
    """Legacy HC number -> (proposal 1, proposal 2, node, subgroup, description, note), from
    the sheet "Mapping" of the numbering workbook."""
    from openpyxl import load_workbook

    handle = io.BytesIO(source) if isinstance(source, bytes) else source
    wb = load_workbook(handle, read_only=True, data_only=True)
    try:
        rows = wb["Mapping"].iter_rows(min_row=2, max_col=7, values_only=True)
        return {r[0]: r[1:] for r in rows if isinstance(r[0], int)}
    finally:
        wb.close()


BUNDLED_MAPPING = Path(__file__).resolve().parent / "data" / "hc_numbering_mapping.csv"


def load_bundled_mapping() -> dict[int, tuple]:
    """The same table as `load_mapping`, from the copy shipped with the application
    (data/hc_numbering_mapping.csv; a test keeps it equal to the workbook in docs/artifacts)."""
    mapping: dict[int, tuple] = {}
    with BUNDLED_MAPPING.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            mapping[int(row["legacy_number"])] = (
                int(row["proposal_1"]),
                int(row["proposal_2"]),
                row["logical_node"] or None,
                row["hc_subgroup"] or None,
                row["description"] or None,
                row["note"] or None,
            )
    return mapping


def apply_mapping(code: StatusCode, mapping: dict[int, tuple], proposal: int) -> None:
    """Turns the legacy number of an HC code into its new number from the mapping table. A code
    without an entry keeps an empty identifier and a warning."""
    legacy = code.number
    code.legacy_number = legacy
    entry = mapping.get(legacy)
    if entry is None:
        code.warnings.append(f"legacy HC number {legacy} is not in the mapping table")
        code.identifier = ""
        return
    (
        code.proposal_1,
        code.proposal_2,
        code.logical_node,
        code.hc_subgroup,
        code.mapping_description,
        code.mapping_note,
    ) = (entry[0], entry[1], entry[2] or "", entry[3] or "", entry[4] or "", entry[5] or "")
    new_number = code.proposal_1 if proposal == 1 else code.proposal_2
    if not isinstance(new_number, int):
        code.warnings.append(f"the mapping has no new number for proposal {proposal}")
        code.identifier = ""
        return
    code.number = new_number
    code.identifier = identifier_for(new_number)
