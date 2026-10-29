# ruff: noqa: E501  (docstrings and checklist text on long lines)
"""The interface translations as data: every string, the screens that show it, each language's
catalogue, the fingerprint of a catalogue and the review sheet of a language (docs/26-translations.md).

This is the part shared by the generator (scripts/gen_translations_doc.py), which writes the
catalogue document, and the application's translation review section (app/services/
translation_review_service.py), where native speakers check the translations.

`catalogue()` is what the code ships. What the application shows is the catalogue plus the
**overrides** a reviewer's accepted proposals have created (table `translation_override`): see
`effective()`. A string is identified by `string_id`, a short hash of its English text, which
survives re-ordering and is the same in the document, the sheets and the database."""

from __future__ import annotations

import hashlib
import io
import time
from collections import defaultdict
from pathlib import Path

from app import status_code_fields
from app.auth.roles import ROLES
from app.i18n import LANGUAGE_META, RTL_LANGUAGES, TEMPLATE_STRING_RE
from app.i18n_de import DE
from app.i18n_extra import CATALOGUES
from app.i18n_messages import service_messages

TEMPLATES = Path(__file__).resolve().parent / "templates"

ENGLISH_NAMES = {
    "en": "English",
    "de": "German",
    "fr": "French",
    "hi": "Hindi",
    "ka": "Georgian",
    "ar": "Arabic",
    "bn": "Bengali",
    "mr": "Marathi",
    "te": "Telugu",
    "ta": "Tamil",
    "gu": "Gujarati",
    "ur": "Urdu",
    "kn": "Kannada",
    "or": "Odia",
    "ml": "Malayalam",
    "pa": "Punjabi",
    "es": "Spanish",
    "it": "Italian",
    "da": "Danish",
}
VERDICTS = ("OK", "CHANGE", "UNCLEAR")
CHECKLIST = (
    ("Meaning", "Does the translation say what the English says - no more, no less?"),
    (
        "Terminology",
        "Are the fixed terms the same everywhere (Status Code, Draft, Release, Reviewer)? Terms that "
        "stay in English are listed in the concept section.",
    ),
    ("Fit", "Is it short enough for a button, a column heading or a menu entry?"),
    ("Register", "Is the form of address right for a technical tool (formal, neutral, no slang)?"),
    ("Placeholders", "Are markers such as %(n)s kept exactly and in a sensible place?"),
    (
        "Script and direction",
        "Is the script right, are digits and identifiers left in Latin, does it read correctly "
        "right to left (Arabic, Urdu)?",
    ),
    (
        "Spelling and punctuation",
        "Spelling, accents, quotation marks, the full stop at the end of a sentence.",
    ),
)


def string_id(english: str) -> str:
    return "S" + hashlib.sha1(english.encode("utf-8")).hexdigest()[:6]  # nosec B324 - not a secret


_TEMPLATE_STRINGS: tuple[tuple, float, dict[str, frozenset[str]]] | None = None
_TEMPLATE_RECHECK_SECONDS = (
    5.0  # a file system call is slow on a network drive: look at most this often
)


def _template_strings() -> dict[str, frozenset[str]]:
    """English string -> the templates that wrap it in `_()`. Reading the hundred templates costs
    seconds and the Translations page asked for it 37 times per request; the result is kept until a
    template changes (its modification time is part of the key, so a development server still
    follows an edit within a few seconds)."""
    global _TEMPLATE_STRINGS
    now = time.monotonic()
    if _TEMPLATE_STRINGS is not None and now - _TEMPLATE_STRINGS[1] < _TEMPLATE_RECHECK_SECONDS:
        return _TEMPLATE_STRINGS[2]
    paths = sorted(TEMPLATES.rglob("*.html"))
    stamp = tuple((str(p), p.stat().st_mtime_ns) for p in paths)
    if _TEMPLATE_STRINGS is not None and _TEMPLATE_STRINGS[0] == stamp:
        _TEMPLATE_STRINGS = (stamp, now, _TEMPLATE_STRINGS[2])
        return _TEMPLATE_STRINGS[2]
    found: dict[str, set[str]] = defaultdict(set)
    for path in paths:
        text = path.read_text(encoding="utf-8")
        for match in TEMPLATE_STRING_RE.finditer(text):
            english = (match.group(1) or match.group(2)).replace("\\'", "'").replace('\\"', '"')
            found[english].add(path.stem)
    frozen = {k: frozenset(v) for k, v in found.items()}
    _TEMPLATE_STRINGS = (stamp, now, frozen)
    return frozen


def screens() -> dict[str, list[str]]:
    """English string -> the templates that show it (role names and help texts have their own)."""
    found: dict[str, set[str]] = defaultdict(set)
    for english, files in _template_strings().items():
        found[english].update(files)
    for role in ROLES.values():
        found[role.label].add("role badge")
    for english in status_code_fields.LABELS.values():  # labels the templates take from the module
        found[english].add("status code fields")
    from app.export.pdf_exporter import TRANSLATED_TEXT
    from app.services.export_service import TITLE_PAGE_TEXT
    from app.ui import nav  # here: the UI package imports this module

    for english in (*TRANSLATED_TEXT, *TITLE_PAGE_TEXT):  # the words around an export
        found[english].add("export")

    for english in nav.LABELS:  # the menu is data
        found[english].add("menu")
    for english, modules in service_messages().items():
        found[english].update(f"message: {f}" for f in modules[:2])
    for english in DE:  # help text and other strings that are not template literals
        found.setdefault(english, {"help and messages"})
    return {k: sorted(v) for k, v in found.items()}


def catalogue() -> dict[str, dict[str, str]]:
    """language code -> {English: translation} as shipped in the code, for every language but English."""
    out = {"de": dict(DE)}
    out.update({code: dict(texts) for code, texts in CATALOGUES.items()})
    return out


def languages() -> list[str]:
    return list(LANGUAGE_META)


def reviewable_languages() -> list[str]:
    return [c for c in languages() if c != "en"]


def overrides(language: str) -> dict[str, str]:
    """English -> text of the accepted proposals of a language (needs an application context)."""
    from sqlalchemy import text as sql
    from sqlalchemy.exc import OperationalError

    from app.extensions import db

    try:
        rows = db.session.execute(
            sql("SELECT english, text FROM translation_override WHERE language_code = :l"),
            {"l": language},
        ).all()
    except OperationalError:  # a database that has not been upgraded yet
        db.session.rollback()
        return {}
    return {english: translated for english, translated in rows}


def effective(language: str) -> dict[str, str]:
    """What the application shows: the shipped catalogue with the accepted proposals on top."""
    merged = dict(catalogue().get(language, {}))
    merged.update(overrides(language))
    return merged


def fingerprint(translations: dict[str, str]) -> str:
    lines = sorted(f"{string_id(k)}\t{v}" for k, v in translations.items())
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()[:12]


def label(code: str) -> str:
    cc, name, flag = LANGUAGE_META[code]
    return f"{cc} {name} {flag}"


def is_rtl(code: str) -> bool:
    return code in RTL_LANGUAGES


# ------------------------------------------------------------------ the offline sheet
def build_sheet(code: str, translations: dict[str, str] | None = None) -> bytes:
    """The review sheet of one language as XLSX: tabs Review, Missing and Checklist."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    known = screens()
    t = translations if translations is not None else catalogue()[code]
    rtl = is_rtl(code)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Review"
    sheet.append(
        [
            "Id",
            "Screens",
            "English",
            f"{ENGLISH_NAMES[code]} (current)",
            "Verdict",
            "Proposal",
            "Comment",
        ]
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    ordered = sorted(t, key=lambda e: (known.get(e, ["~"])[0], e.lower()))
    for english in ordered:
        sheet.append(
            [
                string_id(english),
                ", ".join(known.get(english, [])[:3]),
                english,
                t[english],
                None,
                None,
                None,
            ]
        )
    for index, width in enumerate((9, 26, 46, 46, 11, 46, 40)):
        sheet.column_dimensions["ABCDEFG"[index]].width = width
    wrap = Alignment(wrap_text=True, vertical="top")
    side = Alignment(
        wrap_text=True,
        vertical="top",
        horizontal="right" if rtl else "left",
        readingOrder=2 if rtl else 1,
    )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = side if cell.column_letter in "DF" else wrap
    sheet.freeze_panes = "D2"
    validation = DataValidation(
        type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True
    )
    sheet.add_data_validation(validation)
    validation.add(f"E2:E{len(ordered) + 1}")
    sheet.auto_filter.ref = f"A1:G{len(ordered) + 1}"

    missing = [e for e in sorted(known, key=str.lower) if e not in t]
    gap = workbook.create_sheet("Missing")
    gap.append(["Id", "Screens", "English", "Your translation"])
    for cell in gap[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7F6000")
    for english in missing:
        gap.append([string_id(english), ", ".join(known[english][:3]), english, None])
    for index, width in enumerate((9, 26, 60, 60)):
        gap.column_dimensions["ABCD"[index]].width = width
    gap.freeze_panes = "A2"

    guide = workbook.create_sheet("Checklist")
    guide.append([f"{label(code)} - {ENGLISH_NAMES[code]}: what to check"])
    guide["A1"].font = Font(bold=True, size=13)
    for name, question in CHECKLIST:
        guide.append([name, question])
    guide.append([])
    guide.append(
        [
            "Verdict",
            "OK = fine as it is; CHANGE = write your version in Proposal; UNCLEAR = the English or the "
            "context is not clear, say why in Comment.",
        ]
    )
    guide.append(["Fingerprint", f"{fingerprint(t)} - it identifies the version you reviewed."])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 110
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
