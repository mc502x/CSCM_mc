# ruff: noqa: E501  (messages)
"""The legacy -> new status code mapping, kept in the database and changed only with three approvals
(FR-191..FR-198, docs/06-data-dictionary.md section 5b).

The mapping holds only what the generated .st programs do **not** carry: the legacy -> new pair, the
English title, the type, the up/down counter, the eight named access levels, the cause and the
directive (and the value a parameter delay is known to have). Everything a program carries - the
trigger snapshot, the alarm, the repeats, the programs, the user-group numbers - stays in the program
and is never duplicated here.

Its initial content is the workbook `seed_data/legacy_mapping.xlsx` (`flask mapping-seed`). From then on
a row changes only through a **mapping change**: one request with one or several rows, decided as a whole
by the StCd Owner, the Chief Engineer and one Administrator - each a different person, none of them the
requester. One rejection closes the request; the third approval applies it, in one transaction.
Delegation does not apply: the three approvals are given by the persons holding the three roles.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, text
from sqlalchemy.exc import IntegrityError

from app import status_code_fields
from app.domain.errors import (
    DomainError,
    ForbiddenError,
    InvalidTransitionError,
    NotFoundError,
    SegregationOfDutiesViolation,
)
from app.extensions import db
from app.models.mapping import (
    ACCESS_COLUMNS,
    MappingApproval,
    MappingChange,
    MappingChangeItem,
    StatusCodeMapping,
)
from app.models.user import User
from app.services.audit_service import AuditService
from app.services.numbering_service import NumberingService
from app.services.validation import is_delay
from app.utils import utcnow_iso

SEED_WORKBOOK = Path(__file__).resolve().parents[2] / "seed_data" / "legacy_mapping.xlsx"

FAMILIES = ("MC", "HC")
STATUSES = ("PROPOSED", "FROM_PROGRAM", "APPROVED", "ISSUED", "REJECTED")
TYPES = ("ERROR", "WARNING", "INFO")
ACTIONS = ("ADD", "CHANGE", "CONFIRM", "REJECT")
APPROVER_ROLES = ("STCD_OWNER", "CHIEF_ENGINEER", "ADMINISTRATOR")
VIEW_ROLES = ("ADMINISTRATOR", "ENGINEER", "REVIEWER", "CHIEF_ENGINEER", "STCD_OWNER")
REQUEST_ROLES = VIEW_ROLES
# a status that lets the import use the number without a warning
SETTLED = ("APPROVED", "ISSUED")

MAX_ITEMS = 3000
MAX_WORKBOOK_BYTES = 5 * 1024 * 1024
MAX_WORKBOOK_UNPACKED = 80 * 1024 * 1024
MAX_SHEET_ROWS = 20_000
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

# the columns of the workbook and of a row, in this order
COLUMNS: tuple[str, ...] = (
    "legacy_number",
    "new_number",
    "description",
    "functional_group",
    "logical_node",
    "type",
    "up_down_counter",
    "parameter_set_delay",
    "parameter_reset_delay",
    "development",
    "sales",
    "tcc",
    "service",
    "turbine_operator_package",
    "grid_operator",
    "service_partner",
    "customer",
    "cause",
    "directive",
    "note",
    "status",
    "source",
)
# workbook column -> database column
_COLUMN_FIELD = {"type": "status_type"}
_COLUMN_FIELD.update({c.removesuffix("_access"): c for c in ACCESS_COLUMNS})
# the fields a mapping row has (the columns of the workbook without family; the family is the sheet)
FIELDS: tuple[str, ...] = ("family",) + tuple(_COLUMN_FIELD.get(c, c) for c in COLUMNS)
# what a person edits; the status is the result of the approvals and the source is a record
EDITABLE: tuple[str, ...] = tuple(f for f in FIELDS if f not in ("status", "source"))
_TEXT_LIMITS = {
    "description": 120,
    "functional_group": 80,
    "logical_node": 12,
    "parameter_set_delay": 40,
    "parameter_reset_delay": 40,
    "cause": 4000,
    "directive": 4000,
    "note": 2000,
    "source": 200,
}
_FIELD_LABELS = {
    "legacy_number": "legacy number",
    "new_number": "new number",
    "description": "description",
    "functional_group": "functional group",
    "logical_node": "logical node",
    "status_type": "type",
    "up_down_counter": "up/down counter",
    "parameter_set_delay": "set delay",
    "parameter_reset_delay": "reset delay",
    "cause": "cause",
    "directive": "directive",
    "note": "note",
}
_FIELD_LABELS.update(
    {c: c.removesuffix("_access").replace("_", " ") + " access" for c in ACCESS_COLUMNS}
)


class MappingError(DomainError):
    """A mapping change that cannot be made: stale rows, an invalid row, a rule of the approvals."""


# ------------------------------------------------------------------------------ the rows
@dataclass(frozen=True)
class MappingRow:
    """A row as the import reads it."""

    id: int
    family: str
    legacy_number: int | None
    new_number: int
    description: str
    functional_group: str | None
    logical_node: str | None
    status_type: str | None
    up_down_counter: bool | None
    parameter_set_delay: str | None
    parameter_reset_delay: str | None
    access: dict[str, str | None]
    cause: str | None
    directive: str | None
    note: str | None
    status: str

    @property
    def identifier(self) -> str:
        return f"StCd-{self.new_number:05d}"


def row_dict(model: StatusCodeMapping) -> dict[str, Any]:
    """The editable and recorded fields of a stored row, as the plain values a change carries."""
    return {name: getattr(model, name) for name in FIELDS}


def row_view(model: StatusCodeMapping) -> MappingRow:
    return MappingRow(
        id=model.id,
        family=model.family,
        legacy_number=model.legacy_number,
        new_number=model.new_number,
        description=model.description,
        functional_group=model.functional_group,
        logical_node=model.logical_node,
        status_type=model.status_type,
        up_down_counter=None if model.up_down_counter is None else bool(model.up_down_counter),
        parameter_set_delay=model.parameter_set_delay,
        parameter_reset_delay=model.parameter_reset_delay,
        access={c: getattr(model, c) for c in ACCESS_COLUMNS},
        cause=model.cause,
        directive=model.directive,
        note=model.note,
        status=model.status,
    )


class MappingTable:
    """The mapping as the import needs it: by new number (Main Controller programs carry the new
    number) and by legacy number (Hub Controller programs carry the legacy number)."""

    def __init__(self, rows: list[MappingRow]):
        self.rows = rows
        self.by_new: dict[int, MappingRow] = {r.new_number: r for r in rows}
        self.by_legacy: dict[tuple[str, int], MappingRow] = {
            (r.family, r.legacy_number): r for r in rows if r.legacy_number is not None
        }

    def __bool__(self) -> bool:
        return bool(self.rows)

    def for_main(self, number: int) -> MappingRow | None:
        row = self.by_new.get(number)
        return row if row is not None and row.family == "MC" else None

    def for_hub(self, legacy: int) -> MappingRow | None:
        return self.by_legacy.get(("HC", legacy))


# ------------------------------------------------------------------------------ cleaning
def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str | None:
    if _blank(value):
        return None
    if isinstance(value, float) and value == int(value):
        value = int(value)
    return str(value).strip()


def _int(value: Any, label: str, errors: list[str]) -> int | None:
    if _blank(value):
        return None
    try:
        if isinstance(value, bool):
            raise ValueError
        number = float(str(value).strip())
        if number != int(number):
            raise ValueError
        return int(number)
    except ValueError:
        errors.append(f"The {label} must be a whole number, not {value!r}.")
        return None


def audience_labels() -> dict[str, dict[str, str]]:
    """Per access column: the active named levels, keyed by the lower-case label."""
    out: dict[str, dict[str, str]] = {}
    for column, (_slot, table) in status_code_fields.AUDIENCE_SLOTS.items():
        # the table names are the constants of AUDIENCE_SLOTS, never input
        statement = "SELECT label FROM " + table + " WHERE is_active = 1"  # nosec B608
        out[column] = {r[0].lower(): r[0] for r in db.session.execute(text(statement))}
    return out


def clean_row(
    raw: dict[str, Any], labels: dict[str, dict[str, str]] | None = None
) -> tuple[dict[str, Any], list[str], list[str]]:
    """Checks and normalises one row. `raw` may use the names of the workbook ("type", "development")
    or of the fields. Returns (row, errors, warnings). An error stops the row; a warning is shown to the
    people who decide."""
    labels = labels if labels is not None else audience_labels()
    errors: list[str] = []
    warnings: list[str] = []
    values: dict[str, Any] = {}
    for key, value in raw.items():
        values[_COLUMN_FIELD.get(key, key)] = value
    row: dict[str, Any] = {}

    family = str(values.get("family") or "").strip().upper()
    if family not in FAMILIES:
        errors.append("The family must be MC or HC.")
    row["family"] = family
    row["legacy_number"] = _int(values.get("legacy_number"), "legacy number", errors)
    if row["legacy_number"] is not None and row["legacy_number"] < 0:
        errors.append("The legacy number must not be negative.")
    row["new_number"] = _int(values.get("new_number"), "new number", errors)
    if row["new_number"] is None:
        if not any("new number" in e for e in errors):
            errors.append("The new number is missing.")
    elif not 1 <= row["new_number"] <= 99999:
        errors.append("The new number must lie between 1 and 99999.")

    for name, limit in _TEXT_LIMITS.items():
        if name == "source":
            continue
        value = _text(values.get(name))
        if value is not None and _CONTROL.search(value):
            errors.append(f"The {_FIELD_LABELS.get(name, name)} contains control characters.")
            value = None
        if value is not None and len(value) > limit:
            errors.append(f"The {_FIELD_LABELS.get(name, name)} is longer than {limit} characters.")
        row[name] = value
    if row["description"] is None:
        errors.append("The description (the English title) is missing.")

    kind = _text(values.get("status_type"))
    if kind is None:
        row["status_type"] = None
    elif kind.upper() in TYPES:
        row["status_type"] = kind.upper()
    else:
        row["status_type"] = None
        errors.append(f"The type must be Error, Warning or Info, not {kind!r}.")

    counter = values.get("up_down_counter")
    if _blank(counter):
        row["up_down_counter"] = None
    else:
        word = str(counter).strip().lower()
        if word in ("yes", "y", "1", "true"):
            row["up_down_counter"] = 1
        elif word in ("no", "n", "0", "false"):
            row["up_down_counter"] = 0
        else:
            row["up_down_counter"] = None
            errors.append(f"The up/down counter must be Yes or No, not {counter!r}.")

    for name in ("parameter_set_delay", "parameter_reset_delay"):
        if row[name] is not None and not is_delay(row[name]):
            errors.append(
                f"The {_FIELD_LABELS[name]} {row[name]!r} is neither a time (10ms, *5s) nor the name of a parameter."
            )

    for column in ACCESS_COLUMNS:
        value = _text(values.get(column))
        if value is None:
            row[column] = None
        elif value.lower() in labels.get(column, {}):
            row[column] = labels[column][value.lower()]
        else:
            row[column] = None
            errors.append(
                f"{_FIELD_LABELS[column].capitalize()}: {value!r} is not a named level of this audience."
            )

    status = _text(values.get("status")) or "PROPOSED"
    row["status"] = status.upper()
    if row["status"] not in STATUSES:
        errors.append("The status must be one of " + ", ".join(STATUSES) + ".")
        row["status"] = "PROPOSED"
    row["source"] = (_text(values.get("source")) or "")[: _TEXT_LIMITS["source"]] or None

    # the number against the numbering bands of the library
    if row["new_number"] is not None and family in FAMILIES and 1 <= row["new_number"] <= 99999:
        band = NumberingService.band_of(row["new_number"])
        if band is None:
            warnings.append(
                f"No numbering subgroup covers {row['new_number']}; Controls Engineering must open one before a program can be imported."
            )
        else:
            from app.models.lookup import LookupFunctionalSystemGroup

            group = db.session.get(LookupFunctionalSystemGroup, band[0])
            expected = "HC" if group is not None and group.controller_type == "HUB" else "MC"
            if expected != family:
                errors.append(
                    f"{row['new_number']} lies in the {group.code if group else '?'} range, which belongs to the {expected} numbering, not to {family}."
                )
            elif group is not None and row["logical_node"] and row["logical_node"] != group.code:
                warnings.append(
                    f"The logical node {row['logical_node']} differs from {group.code}, the group of the number range of {row['new_number']}."
                )
    return row, errors, warnings


def completeness(row: dict[str, Any]) -> list[str]:
    """What a row still lacks before it can be approved: the import needs the type and the counter."""
    missing = []
    if not row.get("status_type"):
        missing.append("type")
    if row.get("up_down_counter") is None:
        missing.append("up/down counter")
    if not row.get("functional_group"):
        missing.append("functional group")
    if not row.get("logical_node"):
        missing.append("logical node")
    return missing


def differences(before: dict[str, Any] | None, after: dict[str, Any]) -> list[str]:
    """The editable fields in which two rows differ."""
    if before is None:
        return [f for f in EDITABLE if after.get(f) is not None]
    return [f for f in EDITABLE if before.get(f) != after.get(f)]


def field_label(field: str) -> str:
    return _FIELD_LABELS.get(field, field)


# ------------------------------------------------------------------------------ the service
class MappingService:
    # ---------------------------------------------------------------- reading
    @staticmethod
    def table() -> MappingTable:
        """Every row except the rejected ones, for the import."""
        models = (
            db.session.query(StatusCodeMapping).filter(StatusCodeMapping.status != "REJECTED").all()
        )
        return MappingTable([row_view(m) for m in models])

    @staticmethod
    def count() -> int:
        return int(db.session.query(func.count(StatusCodeMapping.id)).scalar() or 0)

    @staticmethod
    def get(mapping_id: int) -> StatusCodeMapping:
        model = db.session.get(StatusCodeMapping, mapping_id)
        if model is None:
            raise NotFoundError("Mapping row not found")
        return model

    @staticmethod
    def search(
        family: str | None = None,
        status: str | None = None,
        query: str | None = None,
        page: int = 1,
        page_size: int = 100,
    ) -> tuple[list[StatusCodeMapping], int]:
        q = db.session.query(StatusCodeMapping)
        if family in FAMILIES:
            q = q.filter(StatusCodeMapping.family == family)
        if status in STATUSES:
            q = q.filter(StatusCodeMapping.status == status)
        needle = (query or "").strip()
        if needle:
            like = "%" + needle.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
            clauses = [
                StatusCodeMapping.description.ilike(like, escape="\\"),
                StatusCodeMapping.functional_group.ilike(like, escape="\\"),
                StatusCodeMapping.cause.ilike(like, escape="\\"),
            ]
            if needle.isdigit():
                clauses += [
                    StatusCodeMapping.legacy_number == int(needle),
                    StatusCodeMapping.new_number == int(needle),
                ]
            q = q.filter(db.or_(*clauses))
        total = q.count()
        page = max(1, page)
        rows = (
            q.order_by(StatusCodeMapping.family, StatusCodeMapping.new_number)
            .offset((page - 1) * page_size)
            .limit(page_size)
            .all()
        )
        return rows, total

    @staticmethod
    def summary() -> dict[str, Any]:
        by_status: dict[str, int] = {s: 0 for s in STATUSES}
        by_family = {f: 0 for f in FAMILIES}
        for family, status, count in db.session.query(
            StatusCodeMapping.family, StatusCodeMapping.status, func.count(StatusCodeMapping.id)
        ).group_by(StatusCodeMapping.family, StatusCodeMapping.status):
            by_status[status] += count
            by_family[family] += count
        return {
            "total": sum(by_family.values()),
            "by_family": by_family,
            "by_status": by_status,
            "open_changes": int(
                db.session.query(func.count(MappingChange.id)).filter_by(state="OPEN").scalar() or 0
            ),
        }

    @staticmethod
    def changes(state: str | None = None) -> list[MappingChange]:
        q = db.session.query(MappingChange)
        if state:
            q = q.filter(MappingChange.state == state)
        return q.order_by(MappingChange.id.desc()).all()

    @staticmethod
    def get_change(change_id: int) -> MappingChange:
        change = db.session.get(MappingChange, change_id)
        if change is None:
            raise NotFoundError("Mapping change not found")
        return change

    @staticmethod
    def awaiting(user: Any) -> int:
        """Open requests the user can still decide (nav counter)."""
        role = getattr(user, "role_code", None)
        if role not in APPROVER_ROLES:
            return 0
        decided = {a.change_id for a in db.session.query(MappingApproval).filter_by(role_code=role)}
        return int(
            sum(
                1
                for c in db.session.query(MappingChange).filter_by(state="OPEN")
                if c.id not in decided and c.requested_by != user.id
            )
        )

    @staticmethod
    def who_decides(change: MappingChange) -> dict[str, MappingApproval | None]:
        got = {a.role_code: a for a in change.approvals}
        return {role: got.get(role) for role in APPROVER_ROLES}

    # ---------------------------------------------------------------- proposing a change
    @staticmethod
    def build_item(
        action: str, mapping_id: int | None, edits: dict[str, Any] | None = None
    ) -> dict[str, Any]:
        """An item for `propose`: the action, the stored row (None for ADD) and the edits applied to it."""
        if action not in ACTIONS:
            raise MappingError("Unknown action " + action)
        before = None
        if mapping_id is not None:
            before = row_dict(MappingService.get(mapping_id))
        after = dict(before or {})
        after.update(edits or {})
        return {"action": action, "mapping_id": mapping_id, "before": before, "after": after}

    @staticmethod
    def propose(
        actor: Any,
        title: str,
        reason: str,
        items: list[dict[str, Any]],
        origin: str = "FORM",
    ) -> MappingChange:
        """Opens a change request. Every item is checked against the table as it would be after the
        whole request; any error stops the request with every message."""
        if getattr(actor, "role_code", None) not in REQUEST_ROLES:
            raise ForbiddenError("Your role may not propose a change of the mapping.")
        title = (title or "").strip()
        reason = (reason or "").strip()
        if not title or len(title) > 120:
            raise MappingError("Give the request a title of at most 120 characters.")
        if not reason:
            raise MappingError("Give the reason for the change; the three approvers read it first.")
        if len(reason) > 2000:
            raise MappingError("The reason is longer than 2000 characters.")
        if not items:
            raise MappingError("The request contains no row.")
        if len(items) > MAX_ITEMS:
            raise MappingError(f"At most {MAX_ITEMS} rows in one request.")
        cleaned, problems = MappingService._check_items(items)
        if problems:
            raise MappingError(
                "\n".join(problems[:30])
                + (f"\n... and {len(problems) - 30} more" if len(problems) > 30 else "")
            )
        change = MappingChange(
            title=title,
            reason=reason,
            origin=origin if origin in ("FORM", "WORKBOOK", "SELECTION") else "FORM",
            requested_by=actor.id,
            requested_at=utcnow_iso(),
            state="OPEN",
        )
        db.session.add(change)
        db.session.flush()
        for item in cleaned:
            db.session.add(
                MappingChangeItem(
                    change_id=change.id,
                    mapping_id=item["mapping_id"],
                    action=item["action"],
                    base_version=item["base_version"],
                    before_json=(
                        json.dumps(item["before"], ensure_ascii=False) if item["before"] else None
                    ),
                    after_json=json.dumps(item["after"], ensure_ascii=False),
                )
            )
        AuditService.log(
            "MappingChange",
            change.id,
            "MAPPING_CHANGE_REQUESTED",
            actor=actor,
            after={"title": title, "rows": len(cleaned), "origin": change.origin},
        )
        db.session.commit()
        return change

    @staticmethod
    def _check_items(items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
        """Cleans the items and checks the table that would result: unique new numbers and unique
        legacy numbers, no change of a number that a code of the catalogue already carries, nothing
        approved that lacks a type."""
        labels = audience_labels()
        problems: list[str] = []
        cleaned: list[dict[str, Any]] = []
        stored = {m.id: m for m in db.session.query(StatusCodeMapping)}
        held = numbers_in_catalogue()
        result: dict[int, dict[str, Any]] = {mid: row_dict(m) for mid, m in stored.items()}
        added: list[dict[str, Any]] = []
        touched: set[int] = set()
        for index, item in enumerate(items, 1):
            action = item.get("action")
            mapping_id = item.get("mapping_id")
            where = f"Row {index}"
            if action not in ACTIONS:
                problems.append(f"{where}: unknown action {action!r}.")
                continue
            if action == "ADD":
                if mapping_id is not None:
                    problems.append(f"{where}: a new row has no stored row.")
                    continue
                before = None
            else:
                model = stored.get(mapping_id) if mapping_id is not None else None
                if model is None:
                    problems.append(f"{where}: the row to change does not exist.")
                    continue
                if mapping_id in touched:
                    problems.append(f"{where}: the row {model.new_number} is in the request twice.")
                    continue
                touched.add(model.id)
                before = row_dict(model)
            raw = dict(item.get("after") or {})
            row, errors, _warnings = clean_row(raw, labels)
            name = f"Row {row.get('new_number') or index}"
            problems.extend(f"{name}: {e}" for e in errors)
            if action == "REJECT":
                row = dict(before or {})
                row["status"] = "REJECTED"
            elif action in ("ADD", "CHANGE", "CONFIRM"):
                missing = completeness(row)
                if missing:
                    problems.append(
                        f"{name}: cannot be approved without " + ", ".join(missing) + "."
                    )
            if before is not None:
                current = stored[mapping_id]  # type: ignore[index]
                if current.new_number in held:
                    if action == "REJECT":
                        problems.append(
                            f"{name}: a code of the catalogue carries this number; it cannot be rejected."
                        )
                    if (
                        row["new_number"] != current.new_number
                        or row["legacy_number"] != current.legacy_number
                        or row["family"] != current.family
                    ):
                        problems.append(
                            f"{name}: a code of the catalogue carries this number; the numbers cannot change."
                        )
                if action == "CHANGE" and not differences(before, row):
                    problems.append(f"{name}: nothing is changed.")
                if action == "CONFIRM":
                    if differences(before, row):
                        problems.append(
                            f"{name}: confirming does not change content; propose a change instead."
                        )
                    if current.status in SETTLED:
                        problems.append(f"{name}: already {current.status.lower()}.")
                result[mapping_id] = row  # type: ignore[index]
            else:
                added.append(row)
            cleaned.append(
                {
                    "action": action,
                    "mapping_id": mapping_id,
                    "base_version": (
                        stored[mapping_id].version
                        if mapping_id is not None and mapping_id in stored
                        else None
                    ),
                    "before": before,
                    "after": row,
                }
            )
        # the table after the request: unique new numbers and unique (family, legacy)
        seen_new: dict[int, str] = {}
        seen_legacy: dict[tuple[str, int], str] = {}
        # a rejected row keeps its numbers (the database holds them unique), so it counts as well
        for r in [*result.values(), *added]:
            if r.get("new_number") is None:
                continue
            label = f"{r['family']} {r['new_number']}"
            if r["new_number"] in seen_new:
                problems.append(
                    f"The new number {r['new_number']} would belong to two rows ({seen_new[r['new_number']]} and {label})."
                )
            seen_new[r["new_number"]] = label
            if r.get("legacy_number") is not None:
                key = (r["family"], r["legacy_number"])
                if key in seen_legacy:
                    problems.append(
                        f"The legacy {r['family']} number {r['legacy_number']} would belong to two rows ({seen_legacy[key]} and {label})."
                    )
                seen_legacy[key] = label
        return cleaned, problems

    # ---------------------------------------------------------------- deciding
    @staticmethod
    def decide(
        change_id: int,
        actor: Any,
        decision: str,
        comment: str | None = None,
        reauth_password: str | None = None,
    ) -> MappingChange:
        """The StCd Owner, the Chief Engineer and one Administrator each decide once. A rejection
        closes the request; the third approval applies it. The decision asks for the password again
        (I-17), like every sensitive decision."""
        change = MappingService.get_change(change_id)
        role = getattr(actor, "role_code", None)
        if role not in APPROVER_ROLES:
            raise ForbiddenError(
                "Only the StCd Owner, the Chief Engineer and an Administrator decide on the mapping."
            )
        if decision not in ("APPROVE", "REJECT"):
            raise MappingError("The decision must be APPROVE or REJECT.")
        if change.state != "OPEN":
            raise InvalidTransitionError("The request is already closed.")
        if change.requested_by == actor.id:
            raise SegregationOfDutiesViolation(
                "You requested this change; another person must decide."
            )
        comment = (comment or "").strip() or None
        if comment and len(comment) > 1000:
            raise MappingError("The comment is longer than 1000 characters.")
        if decision == "REJECT" and not comment:
            raise MappingError("Say why you reject the request.")
        if any(a.role_code == role for a in change.approvals):
            raise InvalidTransitionError(
                f"The {role.replace('_', ' ').title()} has decided already."
            )
        from app.services.auth_service import AuthService

        AuthService.step_up(actor, reauth_password, "mapping_decision")
        db.session.add(
            MappingApproval(
                change_id=change.id,
                role_code=role,
                user_id=actor.id,
                decision=decision,
                comment=comment,
                decided_at=utcnow_iso(),
            )
        )
        try:
            db.session.flush()
        except IntegrityError as exc:  # the same role decided in the very same moment
            db.session.rollback()
            raise InvalidTransitionError("The role has decided already.") from exc
        db.session.refresh(change)
        AuditService.log(
            "MappingChange",
            change.id,
            "MAPPING_APPROVAL" if decision == "APPROVE" else "MAPPING_REJECTION",
            actor=actor,
            after={"role": role, "decision": decision, "comment": comment},
        )
        if decision == "REJECT":
            change.state, change.closed_at = "REJECTED", utcnow_iso()
            AuditService.log(
                "MappingChange",
                change.id,
                "MAPPING_CHANGE_REJECTED",
                actor=actor,
                after={"by": role},
            )
        elif all(a.decision == "APPROVE" for a in change.approvals) and {
            a.role_code for a in change.approvals
        } == set(APPROVER_ROLES):
            try:
                MappingService._apply(change, actor)
            except MappingError:
                db.session.rollback()
                raise
        db.session.commit()
        return change

    @staticmethod
    def withdraw(change_id: int, actor: Any) -> MappingChange:
        change = MappingService.get_change(change_id)
        if change.requested_by != actor.id:
            raise ForbiddenError("Only the person who made the request can withdraw it.")
        if change.state != "OPEN":
            raise InvalidTransitionError("The request is already closed.")
        change.state, change.closed_at = "WITHDRAWN", utcnow_iso()
        AuditService.log("MappingChange", change.id, "MAPPING_CHANGE_WITHDRAWN", actor=actor)
        db.session.commit()
        return change

    @staticmethod
    def stale_items(change: MappingChange) -> list[str]:
        """The rows that changed after the request was made."""
        versions = dict(db.session.query(StatusCodeMapping.id, StatusCodeMapping.version))
        out = []
        for item in change.items:
            if item.mapping_id is None:
                continue
            if versions.get(item.mapping_id) != item.base_version:
                out.append(str(json.loads(item.after_json).get("new_number")))
        return out

    @staticmethod
    def _apply(change: MappingChange, actor: Any) -> None:
        """Writes the rows. The table is checked once more as it is now: another request may have been
        applied since this one was made."""
        stale = MappingService.stale_items(change)
        if stale:
            raise MappingError(
                "These rows changed after the request was made: "
                + ", ".join(stale[:10])
                + ". Withdraw the request and make it again."
            )
        items = [
            {
                "action": i.action,
                "mapping_id": i.mapping_id,
                "after": json.loads(i.after_json),
            }
            for i in change.items
        ]
        _cleaned, problems = MappingService._check_items(items)
        if problems:
            raise MappingError(
                "The request no longer fits the mapping:\n" + "\n".join(problems[:10])
            )
        now = utcnow_iso()
        held = numbers_in_catalogue()
        for item in change.items:
            row = json.loads(item.after_json)
            if item.action == "ADD":
                model = StatusCodeMapping(created_at=now)
                db.session.add(model)
                status = "APPROVED"
            else:
                model = MappingService.get(item.mapping_id)  # type: ignore[arg-type]
                status = (
                    "REJECTED"
                    if item.action == "REJECT"
                    else ("ISSUED" if model.new_number in held else "APPROVED")
                )
            for name in FIELDS:
                if name in ("status", "source"):
                    continue
                setattr(model, name, row.get(name))
            model.status = status
            model.source = f"mapping change {change.id}"
            model.version = 1 if item.action == "ADD" else model.version + 1
            model.updated_at = now
        change.state, change.closed_at = "APPLIED", now
        AuditService.log(
            "MappingChange",
            change.id,
            "MAPPING_CHANGE_APPLIED",
            actor=actor,
            after={"rows": len(change.items)},
        )

    @staticmethod
    def mark_issued(new_number: int, actor: Any = None) -> bool:
        """The catalogue now holds a code with this number: an approved row becomes ISSUED, after which
        its numbers cannot change. A row that is not approved stays as it is."""
        model = db.session.query(StatusCodeMapping).filter_by(new_number=new_number).first()
        if model is None or model.status != "APPROVED":
            return False  # an unapproved row stays what it is: the import only warned
        model.status = "ISSUED"
        model.version += 1
        model.updated_at = utcnow_iso()
        AuditService.log(
            "StatusCodeMapping",
            model.id,
            "MAPPING_ROW_ISSUED",
            actor=actor,
            after={"new_number": new_number},
        )
        return True

    # ---------------------------------------------------------------- workbook
    @staticmethod
    def read_workbook(source: Path | bytes) -> tuple[list[dict[str, Any]], list[str]]:
        """The rows of the sheets MC and HC. Returns (rows, problems); a problem names the sheet and row."""
        if isinstance(source, bytes) and len(source) > MAX_WORKBOOK_BYTES:
            raise MappingError("The workbook is larger than 5 MB.")
        handle: Any = io.BytesIO(source) if isinstance(source, bytes) else source
        try:  # an .xlsx is a ZIP archive: refuse one that unpacks to something absurd
            with zipfile.ZipFile(handle) as archive:
                if sum(i.file_size for i in archive.infolist()) > MAX_WORKBOOK_UNPACKED:
                    raise MappingError("The workbook unpacks to more than 80 MB.")
        except zipfile.BadZipFile as exc:
            raise MappingError("The file is not a readable .xlsx workbook.") from exc
        if hasattr(handle, "seek"):
            handle.seek(0)
        try:
            workbook = load_workbook(handle, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any unreadable file
            raise MappingError("The file is not a readable .xlsx workbook.") from exc
        rows: list[dict[str, Any]] = []
        problems: list[str] = []
        try:
            for family in FAMILIES:
                if family not in workbook.sheetnames:
                    problems.append(f"The workbook has no sheet {family}.")
                    continue
                sheet = workbook[family]
                header: list[str] = []
                for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    if number == 1:
                        header = [str(v).strip() if v is not None else "" for v in values]
                        missing = [
                            c
                            for c in ("legacy_number", "new_number", "description")
                            if c not in header
                        ]
                        if missing:
                            problems.append(
                                f"Sheet {family}: the columns {', '.join(missing)} are missing in row 1."
                            )
                            break
                        continue
                    if number > MAX_SHEET_ROWS:
                        raise MappingError(f"Sheet {family} has more than {MAX_SHEET_ROWS} rows.")
                    if all(_blank(v) for v in values):
                        continue
                    record = {h: v for h, v in zip(header, values, strict=False) if h in COLUMNS}
                    record["family"] = family
                    record["_where"] = f"{family} row {number}"
                    rows.append(record)
                    if len(rows) > MAX_ITEMS:
                        raise MappingError(f"More than {MAX_ITEMS} rows in the workbook.")
        finally:
            workbook.close()
        return rows, problems

    @staticmethod
    def plan_from_workbook(source: bytes, confirm_unchanged: bool = False) -> dict[str, Any]:
        """Compares an uploaded workbook with the mapping. Returns the items of a change request
        (ADD for a new number, CHANGE where a row differs, CONFIRM for an unchanged proposed row when
        asked), the rows that need no change and the problems of the file. Rows of the mapping that the
        workbook does not mention stay as they are."""
        rows, problems = MappingService.read_workbook(source)
        labels = audience_labels()
        stored = list(db.session.query(StatusCodeMapping))
        by_new = {m.new_number: m for m in stored}
        by_legacy = {(m.family, m.legacy_number): m for m in stored if m.legacy_number is not None}
        items: list[dict[str, Any]] = []
        unchanged = 0
        used: set[int] = set()
        for record in rows:
            where = record.pop("_where")
            row, errors, _warnings = clean_row(record, labels)
            if errors:
                problems.extend(f"{where}: {e}" for e in errors)
                continue
            model = by_new.get(row["new_number"])
            if model is None and row["legacy_number"] is not None:
                model = by_legacy.get((row["family"], row["legacy_number"]))
            if model is not None and model.id in used:
                problems.append(
                    f"{where}: the number {row['new_number']} appears twice in the workbook."
                )
                continue
            if model is None:
                items.append({"action": "ADD", "mapping_id": None, "after": row})
                continue
            used.add(model.id)
            before = row_dict(model)
            row["status"], row["source"] = before["status"], before["source"]
            if differences(before, row):
                items.append({"action": "CHANGE", "mapping_id": model.id, "after": row})
            elif model.status in ("PROPOSED", "FROM_PROGRAM") and confirm_unchanged:
                items.append({"action": "CONFIRM", "mapping_id": model.id, "after": row})
            else:
                unchanged += 1
        counts = {a: sum(1 for i in items if i["action"] == a) for a in ACTIONS}
        return {
            "items": items,
            "problems": problems,
            "unchanged": unchanged,
            "counts": counts,
            "rows": len(rows),
        }

    @staticmethod
    def seed(path: Path | None = None) -> int:
        """Loads the initial mapping from the workbook when the table is empty. Returns the number of
        rows (0 when the table already holds rows). Rows that fail the checks are skipped and reported by
        raising MappingError with all of them, so a bad seed workbook is never half loaded."""
        if MappingService.count():
            return 0
        source = path or SEED_WORKBOOK
        if not source.exists():
            raise MappingError(f"The seed workbook {source.name} is not installed.")
        rows, problems = MappingService.read_workbook(source)
        labels = audience_labels()
        now = utcnow_iso()
        cleaned: list[dict[str, Any]] = []
        for record in rows:
            where = record.pop("_where")
            row, errors, _warnings = clean_row(record, labels)
            if errors:
                problems.extend(f"{where}: {e}" for e in errors)
            else:
                cleaned.append(row)
        seen_new: set[int] = set()
        seen_legacy: set[tuple[str, int]] = set()
        for row in cleaned:
            if row["new_number"] in seen_new:
                problems.append(f"The new number {row['new_number']} appears twice.")
            seen_new.add(row["new_number"])
            if row["legacy_number"] is not None:
                key = (row["family"], row["legacy_number"])
                if key in seen_legacy:
                    problems.append(
                        f"The legacy {row['family']} number {row['legacy_number']} appears twice."
                    )
                seen_legacy.add(key)
        if problems:
            raise MappingError("The seed workbook has problems:\n" + "\n".join(problems[:30]))
        for row in cleaned:
            model = StatusCodeMapping(created_at=now, updated_at=now, version=1)
            for name in FIELDS:
                setattr(model, name, row.get(name))
            model.source = row.get("source") or f"{source.name}"
            db.session.add(model)
        AuditService.log(
            "StatusCodeMapping",
            0,
            "MAPPING_SEEDED",
            actor=None,
            after={
                "rows": len(cleaned),
                "file": source.name,
                "by_status": {s: sum(1 for r in cleaned if r["status"] == s) for s in STATUSES},
            },
        )
        db.session.commit()
        return len(cleaned)

    @staticmethod
    def export_xlsx(family: str | None = None) -> bytes:
        """The mapping as a workbook with the sheets MC and HC, in the columns the upload reads."""
        workbook = Workbook()
        workbook.remove(workbook.active)
        for fam in FAMILIES:
            if family and family != fam:
                continue
            sheet = workbook.create_sheet(fam)
            sheet.append(list(COLUMNS))
            sheet.freeze_panes = "A2"
            q = (
                db.session.query(StatusCodeMapping)
                .filter_by(family=fam)
                .order_by(StatusCodeMapping.new_number)
            )
            for m in q:
                values = []
                for column in COLUMNS:
                    value = getattr(m, _COLUMN_FIELD.get(column, column))
                    if column == "up_down_counter":
                        value = None if value is None else ("Yes" if value else "No")
                    elif column == "type":
                        value = value.capitalize() if value else None
                    values.append(_safe(value))
                sheet.append(values)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()


def _safe(value: Any) -> Any:
    """A cell never starts a formula in a spreadsheet program."""
    return (
        "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
    )


def numbers_in_catalogue() -> set[int]:
    """The new numbers that a code of the catalogue holds now, as an issued identifier or as the
    number an import reserved (a withdrawn Draft that gave its number back is not among them)."""
    held: set[int] = set()
    for issued, reserved in db.session.execute(
        text(
            "SELECT status_code_identifier, reserved_identifier FROM status_code "
            "WHERE status_code_identifier IS NOT NULL OR reserved_identifier IS NOT NULL"
        )
    ):
        for identifier in (issued, reserved):
            if identifier and identifier.startswith("StCd-") and identifier[5:].isdigit():
                held.add(int(identifier[5:]))
    return held


def user_name(user_id: int | None) -> str:
    user = db.session.get(User, user_id) if user_id else None
    return user.full_name if user else "-"
