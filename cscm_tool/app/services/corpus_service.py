# ruff: noqa: E501  (docstring lines)
"""A report on a folder of real programs: what would the import do with each of them? (I-68)

The import was proved on six fixtures; the first run on the real corpus - about 1,200 generated programs -
will show the cases nobody thought of. This reads a folder (and its sub-folders) the way an upload would,
checks every program against the catalogue and the mapping exactly as the preview does, and reports how
many are ready, how many are blocked and *why*, grouped by reason with the numbers taken out, so that
fifty programs with the same cause are one line. Nothing is created: the preview writes only a staged copy,
which is removed again.
"""

from __future__ import annotations

import io
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from app.importers import stcd
from app.services.st_import_service import MAX_UPLOAD_FILES, StImportService
from app.utils import xlsx_text

BATCH = min(500, MAX_UPLOAD_FILES)
_DIGITS = re.compile(r"\d+")
_QUOTED = re.compile(r"'[^']*'")


def normalise(message: str) -> str:
    """The reason of a message without what differs from program to program (numbers, names)."""
    return _DIGITS.sub("#", _QUOTED.sub("'…'", message))[:200]


class CorpusService:
    @staticmethod
    def files(folder: Path) -> list[Path]:
        return sorted(
            p
            for p in folder.rglob("*")
            if p.is_file() and p.suffix.lower() in stcd.SUFFIXES and p.stat().st_size > 0
        )

    @staticmethod
    def run(folder: Path, actor: Any) -> dict[str, Any]:
        """Checks every program under `folder`. Returns the counts, the reasons and one row per program."""
        paths = CorpusService.files(folder)
        rows: list[dict[str, Any]] = []
        file_errors: list[dict[str, str]] = []
        notes: list[str] = []
        for start in range(0, len(paths), BATCH):
            batch = paths[start : start + BATCH]
            uploads = [
                (str(p.relative_to(folder)).replace("\\", "/"), p.read_bytes()) for p in batch
            ]
            try:
                preview = StImportService.preview(uploads, actor, {})
            except ValueError as exc:  # nothing in this batch is a program
                notes.append(f"files {start + 1}-{start + len(batch)}: {exc}")
                continue
            try:
                rows.extend(preview["items"])
                file_errors.extend(preview["file_errors"])
                notes.extend(preview["notes"])
            finally:
                StImportService.discard(preview["token"], actor)
        states = Counter(r["state"] for r in rows)
        reasons: Counter[tuple[str, str, str]] = Counter()
        for row in rows:
            for issue in row["issues"]:
                reasons[("blocks", issue["column"], normalise(issue["message"]))] += 1
            for warning in row["warnings"]:
                reasons[("warns", "", normalise(warning))] += 1
        return {
            "files": len(paths),
            "programs": len(rows),
            "states": dict(states),
            "controllers": dict(Counter(r["controller"] for r in rows)),
            "reasons": [
                {"kind": k, "column": c, "reason": m, "programs": n}
                for (k, c, m), n in reasons.most_common()
            ],
            "file_errors": file_errors,
            "notes": notes,
            "rows": rows,
        }

    @staticmethod
    def workbook(result: dict[str, Any]) -> bytes:
        """The report as a workbook: the summary, the reasons and one row per program."""
        book = Workbook()
        summary = book.active
        summary.title = "Summary"
        summary.append(["files read", result["files"]])
        summary.append(["programs", result["programs"]])
        for state, count in sorted(result["states"].items()):
            summary.append([f"state {state}", count])
        for controller, count in sorted(result["controllers"].items()):
            summary.append([f"controller {controller}", count])
        summary.append(["files that are not programs", len(result["file_errors"])])
        reasons = book.create_sheet("Reasons")
        reasons.append(["kind", "column", "reason (numbers and names taken out)", "programs"])
        for r in result["reasons"]:
            reasons.append([r["kind"], r["column"], xlsx_text(r["reason"]), r["programs"]])
        programs = book.create_sheet("Programs")
        programs.append(
            [
                "file",
                "line",
                "controller",
                "program",
                "identifier",
                "legacy",
                "state",
                "problems",
                "warnings",
            ]
        )
        for row in result["rows"]:
            programs.append(
                [
                    xlsx_text(row["file"]),
                    row["line"],
                    row["controller"],
                    row["program_number"],
                    row["identifier"],
                    row["legacy_number"],
                    row["state"],
                    xlsx_text(" | ".join(i["message"] for i in row["issues"])[:2000]),
                    xlsx_text(" | ".join(row["warnings"])[:2000]),
                ]
            )
        errors = book.create_sheet("Not programs")
        errors.append(["file", "message"])
        for e in result["file_errors"]:
            errors.append([xlsx_text(e["file"]), xlsx_text(e["message"])])
        for sheet in book.worksheets:
            sheet.freeze_panes = "A2"
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()
