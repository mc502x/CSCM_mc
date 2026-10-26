# ruff: noqa: E501  (docstring and helper on long lines)
"""Translation review in the application (docs/26-translations.md, docs/03-srs.md FR-164).

Native speakers check the interface translations of the language they are assigned to, in the
browser: one row per string with the English source and the current translation, a verdict
(OK / CHANGE / UNCLEAR), a proposal and a comment. The reviewer *completes* the review when every
string has an answer. An Administrator or an StCd Owner - never the reviewer of that proposal -
**accepts** or rejects the proposals; an accepted proposal becomes an *override*, shown by the
application at once (`app/translation_catalogue.py: effective()`), without a new release.

Who may do what:

* any signed-in user sees the overview (an empty one when they review nothing);
* **reviewing** a language needs an assignment to it (`language_reviewer`); an assigned account can be
  of any role, an external native speaker gets a Viewer account;
* **assigning** reviewers and **deciding** on proposals is for the Administrator and the StCd Owner;
* nobody decides on their own proposal (the same segregation as BR-018 for status code translations).

The accepted overrides stay in the database. `docs/26-translations.md` describes what the code
ships; moving the overrides into the catalogue files is a maintainer's step."""

from __future__ import annotations

import io
import re
from typing import Any

from sqlalchemy import func

from app import translation_precheck as precheck
from app.domain.errors import DomainError, ForbiddenError, NotFoundError
from app.extensions import db
from app.models.translation import (
    LanguageReviewer,
    TranslationOverride,
    TranslationReview,
    TranslationReviewItem,
)
from app.models.user import User
from app.services.audit_service import AuditService
from app.translation_catalogue import (
    VERDICTS,
    effective,
    fingerprint,
    is_rtl,
    reviewable_languages,
    screens,
    string_id,
)
from app.utils import utcnow_iso

MANAGER_ROLES = ("ADMINISTRATOR", "STCD_OWNER")
MAX_TEXT = 1000
MAX_SHEET_BYTES = 5 * 1024 * 1024
PER_PAGE = 40
_PLACEHOLDER = re.compile(r"%\(\w+\)[sd]")


def _check_language(language: str) -> None:
    if language not in reviewable_languages():
        raise NotFoundError("Unknown language")


def _placeholders(text: str) -> list[str]:
    return sorted(_PLACEHOLDER.findall(text))


class TranslationReviewService:
    # ------------------------------------------------------------------ access
    @staticmethod
    def is_manager(actor) -> bool:
        return getattr(actor, "role_code", None) in MANAGER_ROLES

    @staticmethod
    def approved(language: str) -> dict[str, str]:
        """English -> the translation a native speaker has approved: an OK answer, an accepted
        proposal (its final text), a rejected proposal or a noted question (the text as it was).
        The newest answer of a completed review wins; an unanswered question or an undecided
        proposal approves nothing. A string that changed since is not approved any more."""
        rows = (
            db.session.query(TranslationReviewItem)
            .join(TranslationReview, TranslationReview.id == TranslationReviewItem.review_id)
            .filter(
                TranslationReview.language_code == language, TranslationReview.status == "COMPLETED"
            )
            .order_by(TranslationReview.completed_at, TranslationReviewItem.id)
            .all()
        )
        overrides = {
            o.review_item_id: o.text
            for o in db.session.query(TranslationOverride).filter_by(language_code=language)
        }
        approved: dict[str, str] = {}
        for item in rows:
            if item.verdict == "OK":
                text = item.current_text
            elif item.verdict == "CHANGE" and item.decision == "ACCEPTED":
                text = overrides.get(item.id) or item.proposal
            elif item.decision in ("REJECTED", "NOTED"):
                text = item.current_text
            else:
                text = None
            if text is None:
                approved.pop(item.english, None)
            else:
                approved[item.english] = text
        return approved

    @staticmethod
    def pending(language: str) -> dict[str, str]:
        """English -> text of the strings of a language that no native speaker has approved yet:
        everything while no review is completed, afterwards only what is new or changed (FR-176)."""
        current = effective(language)
        approved = TranslationReviewService.approved(language)
        return {e: t for e, t in current.items() if approved.get(e) != t}

    @staticmethod
    def reviewed_languages() -> frozenset[str]:
        """Languages with at least one completed review: the ones a Viewer may use (FR-172)."""
        rows = db.session.query(TranslationReview.language_code).filter_by(status="COMPLETED")
        return frozenset(r[0] for r in rows.distinct())

    @staticmethod
    def beta_languages() -> frozenset[str]:
        """Languages (besides the source English) with strings that no native speaker has approved:
        all of them while no review is completed, afterwards only the new or changed ones (FR-170,
        FR-176)."""
        from flask import g, has_request_context

        from app.services.settings_service import SettingsService

        if has_request_context() and g.get("beta_languages") is not None:
            return g.beta_languages
        if SettingsService.get("languages.beta_marker") != "1":
            return frozenset()
        reviewed = TranslationReviewService.reviewed_languages()
        beta = set(reviewable_languages()) - reviewed
        for code in reviewed:
            if TranslationReviewService.pending(code):
                beta.add(code)
        result = frozenset(beta)
        if has_request_context():
            g.beta_languages = result
        return result

    @staticmethod
    def unreviewed_languages() -> frozenset[str]:
        """Languages no Viewer or guest may use: no review has ever been completed."""
        from app.services.settings_service import SettingsService

        if SettingsService.get("languages.beta_marker") != "1":
            return frozenset()
        return frozenset(reviewable_languages()) - TranslationReviewService.reviewed_languages()

    @staticmethod
    def withheld_strings(language: str) -> frozenset[str]:
        """The English texts a Viewer or guest reads in English although the language has been
        reviewed: the strings added or changed since the review, until they are approved (FR-176).
        """
        from flask import g, has_request_context

        if has_request_context():
            cache = g.setdefault("withheld_strings", {})
            if language in cache:
                return cache[language]
        result: frozenset[str] = frozenset()
        if language in TranslationReviewService.reviewed_languages():
            result = frozenset(TranslationReviewService.pending(language))
        if has_request_context():
            cache[language] = result
        return result

    @staticmethod
    def assigned_languages(user_id: int) -> list[str]:
        rows = db.session.query(LanguageReviewer.language_code).filter_by(user_id=user_id).all()
        return sorted(r[0] for r in rows)

    @staticmethod
    def has_access(actor) -> bool:
        return TranslationReviewService.is_manager(actor) or bool(
            TranslationReviewService.assigned_languages(actor.id)
        )

    @staticmethod
    def can_review(actor, language: str) -> bool:
        return language in TranslationReviewService.assigned_languages(actor.id)

    @staticmethod
    def _require_manager(actor) -> None:
        if not TranslationReviewService.is_manager(actor):
            raise ForbiddenError("Only an Administrator or an StCd Owner may do this.")

    # ------------------------------------------------------------------ overview
    @staticmethod
    def overview(actor) -> list[dict[str, Any]]:
        total = len(screens())
        manager = TranslationReviewService.is_manager(actor)
        mine = set(TranslationReviewService.assigned_languages(actor.id))
        rows = []
        for code in reviewable_languages():
            if not manager and code not in mine:
                continue
            current = effective(code)
            reviewers = (
                db.session.query(User.full_name, User.id)
                .join(LanguageReviewer, LanguageReviewer.user_id == User.id)
                .filter(LanguageReviewer.language_code == code)
                .order_by(User.full_name)
                .all()
            )
            last = (
                db.session.query(TranslationReview)
                .filter_by(language_code=code, status="COMPLETED")
                .order_by(TranslationReview.completed_at.desc())
                .first()
            )
            open_count = (
                db.session.query(TranslationReview)
                .filter_by(language_code=code, status="OPEN")
                .count()
            )
            pending = (
                db.session.query(func.count(TranslationReviewItem.id))
                .join(TranslationReview, TranslationReview.id == TranslationReviewItem.review_id)
                .filter(
                    TranslationReview.language_code == code,
                    TranslationReviewItem.decision == "PENDING",
                    TranslationReviewItem.verdict.in_(("CHANGE", "UNCLEAR")),
                )
                .scalar()
            )
            awaiting = len(TranslationReviewService.pending(code))
            if last is None:
                state = "in progress" if open_count else "not reviewed"
            elif awaiting:
                state = "changed since the review"
            else:
                state = "reviewed"
            check = precheck.load().get("languages", {}).get(code)
            my_open = (
                db.session.query(TranslationReview)
                .filter_by(language_code=code, reviewer_id=actor.id, status="OPEN")
                .first()
            )
            rows.append(
                {
                    "code": code,
                    "rtl": is_rtl(code),
                    "translated": len(current),
                    "total": total,
                    "share": 100 * len(current) // total if total else 0,
                    "overrides": db.session.query(TranslationOverride)
                    .filter_by(language_code=code)
                    .count(),
                    "reviewers": [{"id": i, "name": n} for n, i in reviewers],
                    "state": state,
                    "last_completed": last,
                    "pending": int(pending or 0),
                    "awaiting": awaiting,
                    "precheck": (
                        {
                            "confidence": check["confidence"],
                            "bucket": precheck.bucket(check["confidence"]),
                            "errors": check["errors"],
                            "awkward": check["awkward"],
                            "glossary": check["mechanical"]["glossary_mismatches"],
                            "haircut": check["haircut"],
                            "sample": check["sample_size"],
                            "checked_at": check["checked_at"],
                            "changed": len(precheck.changes()["strings"].get(code, [])),
                        }
                        if check
                        else None
                    ),
                    "can_review": code in mine,
                    "my_open": my_open,
                }
            )
        return rows

    # ------------------------------------------------------------------ assignments
    @staticmethod
    def assign(language: str, user_id: int, actor) -> LanguageReviewer:
        TranslationReviewService._require_manager(actor)
        _check_language(language)
        user = db.session.get(User, user_id)
        if user is None or not user.is_active:
            raise DomainError("Choose an active user.")
        existing = (
            db.session.query(LanguageReviewer)
            .filter_by(user_id=user_id, language_code=language)
            .first()
        )
        if existing is not None:
            return existing
        row = LanguageReviewer(
            user_id=user_id, language_code=language, assigned_by=actor.id, assigned_at=utcnow_iso()
        )
        db.session.add(row)
        db.session.flush()
        AuditService.log(
            "LanguageReviewer",
            row.id,
            "LANGUAGE_REVIEWER_ASSIGN",
            actor=actor,
            after={"user_id": user_id, "username": user.username, "language": language},
        )
        db.session.commit()
        return row

    @staticmethod
    def unassign(language: str, user_id: int, actor) -> None:
        TranslationReviewService._require_manager(actor)
        _check_language(language)
        row = (
            db.session.query(LanguageReviewer)
            .filter_by(user_id=user_id, language_code=language)
            .first()
        )
        if row is None:
            raise NotFoundError("That user does not review this language.")
        AuditService.log(
            "LanguageReviewer",
            row.id,
            "LANGUAGE_REVIEWER_REMOVE",
            actor=actor,
            before={"user_id": user_id, "language": language},
        )
        db.session.delete(row)
        db.session.commit()

    # ------------------------------------------------------------------ the review
    @staticmethod
    def start(language: str, actor) -> TranslationReview:
        """The open review of the actor for the language; a new one when there is none."""
        _check_language(language)
        if not TranslationReviewService.can_review(actor, language):
            raise ForbiddenError("You are not assigned to review this language.")
        review = (
            db.session.query(TranslationReview)
            .filter_by(language_code=language, reviewer_id=actor.id, status="OPEN")
            .first()
        )
        if review is not None:
            return review
        current = effective(language)
        review = TranslationReview(
            language_code=language,
            reviewer_id=actor.id,
            status="OPEN",
            started_at=utcnow_iso(),
            fingerprint=fingerprint(current),
            string_count=len(current),
        )
        db.session.add(review)
        db.session.commit()
        return review

    @staticmethod
    def get(review_id: int) -> TranslationReview:
        review = db.session.get(TranslationReview, review_id)
        if review is None:
            raise NotFoundError("Review not found")
        return review

    @staticmethod
    def answers(review_id: int) -> dict[str, TranslationReviewItem]:
        items = db.session.query(TranslationReviewItem).filter_by(review_id=review_id).all()
        return {i.string_id: i for i in items}

    @staticmethod
    def page(
        review: TranslationReview,
        *,
        scope: str = "translated",
        q: str = "",
        screen: str = "",
        state: str = "",
        page: int = 1,
    ) -> dict[str, Any]:
        """The rows of one page of the review screen. `scope` is `translated` (check what is there;
        in a follow-up review only the strings that are new or changed since the last completed
        review, FR-176), `all` (every translated string) or `missing` (strings that still have no
        translation)."""
        known = screens()
        current = effective(review.language_code)
        followup = review.language_code in TranslationReviewService.reviewed_languages()
        required = (
            set(TranslationReviewService.pending(review.language_code))
            if followup
            else set(current)
        )
        answers = TranslationReviewService.answers(review.id)
        rows = []
        needle = q.strip().lower()
        for english in sorted(known, key=lambda e: (known[e][0], e.lower())):
            has = english in current
            if (scope != "missing") != has:
                continue
            if scope == "translated" and followup and english not in required:
                continue
            sid = string_id(english)
            item = answers.get(sid)
            if screen and screen not in known[english]:
                continue
            if (
                needle
                and needle not in english.lower()
                and needle not in (current.get(english) or "").lower()
            ):
                continue
            verdict = item.verdict if item else None
            if state == "unanswered" and verdict:
                continue
            if state == "answered" and not verdict:
                continue
            if state in ("OK", "CHANGE", "UNCLEAR") and verdict != state:
                continue
            rows.append(
                {
                    "id": sid,
                    "english": english,
                    "screens": known[english],
                    "current": current.get(english),
                    "item": item,
                }
            )
        pages = max(1, -(-len(rows) // PER_PAGE))
        page = min(max(page, 1), pages)
        answered = sum(1 for e in required if (a := answers.get(string_id(e))) and a.verdict)
        return {
            "rows": rows[(page - 1) * PER_PAGE : page * PER_PAGE],
            "total": len(rows),
            "page": page,
            "pages": pages,
            "answered": answered,
            "to_answer": len(required),
            "followup": followup,
            "all_total": len(current),
            "missing_total": len(known) - len(current),
            "screens": sorted({s for v in known.values() for s in v}),
        }

    @staticmethod
    def save(review_id: int, entries: list[dict], actor) -> dict[str, Any]:
        """Stores the answers of one page. Rows that are wrong are reported and skipped; the others
        are saved."""
        review = TranslationReviewService.get(review_id)
        if review.reviewer_id != actor.id:
            raise ForbiddenError("This is not your review.")
        if review.status != "OPEN":
            raise DomainError("This review is already completed.")
        known = {string_id(e): e for e in screens()}
        current = effective(review.language_code)
        existing = TranslationReviewService.answers(review.id)
        saved = 0
        errors: list[dict[str, str]] = []
        for entry in entries:
            sid = str(entry.get("string_id") or "")
            english = known.get(sid)
            if english is None:
                errors.append({"string_id": sid, "message": "Unknown string."})
                continue
            verdict = str(entry.get("verdict") or "").strip().upper() or None
            proposal = str(entry.get("proposal") or "").strip() or None
            comment = str(entry.get("comment") or "").strip() or None
            if verdict is not None and verdict not in VERDICTS:
                errors.append(
                    {"string_id": sid, "message": "The verdict must be OK, CHANGE or UNCLEAR."}
                )
                continue
            if proposal and verdict is None:
                verdict = "CHANGE"
            if verdict == "CHANGE" and not proposal:
                errors.append({"string_id": sid, "message": "A CHANGE needs a proposal."})
                continue
            if (proposal and len(proposal) > MAX_TEXT) or (comment and len(comment) > MAX_TEXT):
                errors.append({"string_id": sid, "message": f"At most {MAX_TEXT} characters."})
                continue
            if proposal and _placeholders(proposal) != _placeholders(english):
                errors.append(
                    {
                        "string_id": sid,
                        "message": "Keep the placeholders exactly as in the English text.",
                    }
                )
                continue
            item = existing.get(sid)
            if item is not None and item.decision != "PENDING":
                errors.append(
                    {"string_id": sid, "message": "Already decided; it can no longer be changed."}
                )
                continue
            if not (verdict or proposal or comment):
                if item is not None:
                    db.session.delete(item)
                    saved += 1
                continue
            if item is None:
                item = TranslationReviewItem(
                    review_id=review.id,
                    string_id=sid,
                    english=english,
                    current_text=current.get(english),
                    updated_at=utcnow_iso(),
                )
                db.session.add(item)
            item.verdict, item.proposal, item.comment = verdict, proposal, comment
            item.updated_at = utcnow_iso()
            saved += 1
        db.session.commit()
        return {"saved": saved, "errors": errors}

    @staticmethod
    def complete(review_id: int, actor) -> TranslationReview:
        review = TranslationReviewService.get(review_id)
        if review.reviewer_id != actor.id:
            raise ForbiddenError("This is not your review.")
        if review.status != "OPEN":
            raise DomainError("This review is already completed.")
        answers = TranslationReviewService.answers(review.id)
        current = effective(review.language_code)
        required = (
            set(TranslationReviewService.pending(review.language_code))
            if review.language_code in TranslationReviewService.reviewed_languages()
            else set(current)
        )
        unanswered = [
            e for e in required if not (answers.get(string_id(e)) and answers[string_id(e)].verdict)
        ]
        if unanswered:
            raise ValueError(
                f"{len(unanswered)} of {len(required)} strings have no verdict yet. Answer them, or use "
                "'Mark the rest of this page OK' where they are fine."
            )
        counts = {v: sum(1 for a in answers.values() if a.verdict == v) for v in VERDICTS}
        review.status = "COMPLETED"
        review.completed_at = utcnow_iso()
        AuditService.log(
            "TranslationReview",
            review.id,
            "TRANSLATION_REVIEW_COMPLETE",
            actor=actor,
            after={
                "language": review.language_code,
                "fingerprint_started": review.fingerprint,
                "fingerprint_now": fingerprint(current),
                "strings": len(current),
                **{k.lower(): v for k, v in counts.items()},
            },
        )
        db.session.commit()
        return review

    @staticmethod
    def import_sheet(review_id: int, content: bytes, actor) -> dict[str, Any]:
        """The offline way: the XLSX sheet (tabs Review and Missing) filled in by the reviewer."""
        from openpyxl import load_workbook

        if len(content) > MAX_SHEET_BYTES:
            raise ValueError("The sheet is larger than 5 MB.")
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any unreadable file is the same answer
            raise ValueError(f"The sheet could not be read: {exc}") from exc
        entries: list[dict] = []
        if "Review" not in workbook.sheetnames:
            raise ValueError("The sheet has no 'Review' tab; download the sheet from this page.")
        for row in workbook["Review"].iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                entries.append(
                    {
                        "string_id": row[0],
                        "verdict": row[4] if len(row) > 4 else None,
                        "proposal": row[5] if len(row) > 5 else None,
                        "comment": row[6] if len(row) > 6 else None,
                    }
                )
        if "Missing" in workbook.sheetnames:
            for row in workbook["Missing"].iter_rows(min_row=2, values_only=True):
                if row and row[0] and len(row) > 3 and row[3]:
                    entries.append({"string_id": row[0], "proposal": row[3]})
        if len(entries) > 5000:
            raise ValueError("More than 5000 rows.")
        result = TranslationReviewService.save(
            review_id, [e for e in entries if _answered(e)], actor
        )
        result["rows_read"] = len(entries)
        return result

    # ------------------------------------------------------------------ decisions
    @staticmethod
    def decisions(language: str) -> dict[str, Any]:
        _check_language(language)
        pending = (
            db.session.query(TranslationReviewItem, TranslationReview, User)
            .join(TranslationReview, TranslationReview.id == TranslationReviewItem.review_id)
            .join(User, User.id == TranslationReview.reviewer_id)
            .filter(
                TranslationReview.language_code == language,
                TranslationReviewItem.decision == "PENDING",
                TranslationReviewItem.verdict.in_(("CHANGE", "UNCLEAR")),
            )
            .order_by(TranslationReviewItem.updated_at)
            .all()
        )
        decided = (
            db.session.query(TranslationReviewItem, TranslationReview, User)
            .join(TranslationReview, TranslationReview.id == TranslationReviewItem.review_id)
            .join(User, User.id == TranslationReview.reviewer_id)
            .filter(
                TranslationReview.language_code == language,
                TranslationReviewItem.decision != "PENDING",
            )
            .order_by(TranslationReviewItem.decided_at.desc())
            .limit(30)
            .all()
        )
        return {
            "pending": [{"item": i, "review": r, "reviewer": u} for i, r, u in pending],
            "decided": [{"item": i, "review": r, "reviewer": u} for i, r, u in decided],
            "overrides": db.session.query(TranslationOverride)
            .filter_by(language_code=language)
            .order_by(TranslationOverride.created_at.desc())
            .all(),
        }

    @staticmethod
    def decide(item_id: int, action: str, actor, text: str | None = None) -> TranslationReviewItem:
        TranslationReviewService._require_manager(actor)
        item = db.session.get(TranslationReviewItem, item_id)
        if item is None:
            raise NotFoundError("Proposal not found")
        review = TranslationReviewService.get(item.review_id)
        if review.reviewer_id == actor.id:
            raise ForbiddenError(
                "Nobody decides on their own proposal; another Administrator or StCd Owner does."
            )
        if item.decision != "PENDING":
            raise DomainError("This proposal has been decided already.")
        action = (action or "").lower()
        if action == "accept":
            if item.verdict != "CHANGE" or not item.proposal:
                raise DomainError("Only a proposed change can be accepted.")
            final = (text or item.proposal).strip()
            if not final or len(final) > MAX_TEXT:
                raise ValueError(f"The text must be 1-{MAX_TEXT} characters.")
            if _placeholders(final) != _placeholders(item.english):
                raise ValueError("Keep the placeholders exactly as in the English text.")
            override = (
                db.session.query(TranslationOverride)
                .filter_by(language_code=review.language_code, string_id=item.string_id)
                .first()
            )
            if override is None:
                override = TranslationOverride(
                    language_code=review.language_code,
                    string_id=item.string_id,
                    english=item.english,
                    text=final,
                    review_item_id=item.id,
                    created_by=actor.id,
                    created_at=utcnow_iso(),
                )
                db.session.add(override)
            else:
                override.text, override.review_item_id = final, item.id
                override.created_by, override.created_at = actor.id, utcnow_iso()
            item.decision = "ACCEPTED"
            after: dict[str, Any] = {"decision": "ACCEPTED", "text": final}
        elif action == "reject":
            item.decision = "REJECTED"
            after = {"decision": "REJECTED"}
        elif action == "note":
            if item.verdict != "UNCLEAR":
                raise DomainError("Only an unclear string is noted.")
            item.decision = "NOTED"
            after = {"decision": "NOTED"}
        else:
            raise ValueError("The action must be accept, reject or note.")
        item.decided_by, item.decided_at = actor.id, utcnow_iso()
        AuditService.log(
            "TranslationReviewItem",
            item.id,
            "TRANSLATION_PROPOSAL_DECIDE",
            actor=actor,
            before={
                "language": review.language_code,
                "string_id": item.string_id,
                "verdict": item.verdict,
            },
            after={"language": review.language_code, "english": item.english, **after},
        )
        db.session.commit()
        return item

    @staticmethod
    def remove_override(override_id: int, actor) -> None:
        TranslationReviewService._require_manager(actor)
        override = db.session.get(TranslationOverride, override_id)
        if override is None:
            raise NotFoundError("Override not found")
        AuditService.log(
            "TranslationOverride",
            override.id,
            "TRANSLATION_OVERRIDE_REMOVE",
            actor=actor,
            before={
                "language": override.language_code,
                "english": override.english,
                "text": override.text,
            },
        )
        db.session.delete(override)
        db.session.commit()


def _answered(entry: dict) -> bool:
    return bool(
        str(entry.get("verdict") or "").strip()
        or str(entry.get("proposal") or "").strip()
        or str(entry.get("comment") or "").strip()
    )
