"""Fleet feedback loop (I-28): occurrence counts per status code, imported from the
internal SCADA/CMS side, to show which codes never fire (candidates for deprecation) and
which fire most (candidates for better remedy text).

The interface is a deliberately plain file - a CSV/XLSX with `status_code_identifier` and
`occurrences` (anonymised, aggregated per period; no turbine, site or customer data) - so
whatever the monitoring side can export can be loaded. Whether a scheduled feed replaces the
upload is a business decision (docs/21-improvement-ideas.md §7); this is its agreed shape."""

import csv
import io
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import func

from app.domain.authz import require_roles
from app.extensions import db
from app.models.operations import CodeOccurrence, OccurrenceImport
from app.models.status_code import StatusCode, StatusCodeRevision
from app.services.audit_service import AuditService
from app.utils import utcnow_iso


def _read_rows(filename: str, content: bytes) -> list[dict]:
    lowered = filename.lower()
    if lowered.endswith(".xlsx"):
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [
            dict(zip(headers, r, strict=False)) for r in rows[1:] if any(c is not None for c in r)
        ]
    if lowered.endswith(".csv"):
        return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    raise ValueError("Only .csv and .xlsx files are accepted.")


def _day(value: str, field: str) -> str:
    try:
        return date.fromisoformat(value or "").isoformat()
    except ValueError as exc:
        raise ValueError(f"{field} must be a date like 2027-03-31.") from exc


class OccurrenceService:
    @staticmethod
    def import_counts(
        filename: str,
        content: bytes,
        source_label: str,
        period_start: str,
        period_end: str,
        actor,
    ) -> OccurrenceImport:
        require_roles(actor, "ADMINISTRATOR", "CHIEF_ENGINEER")
        label = (source_label or "").strip()
        if not 1 <= len(label) <= 80:
            raise ValueError("source_label must be 1-80 characters.")
        start, end = _day(period_start, "period_start"), _day(period_end, "period_end")
        if end < start:
            raise ValueError("period_end must not be before period_start.")
        rows = _read_rows(filename, content)
        if rows and not {"status_code_identifier", "occurrences"} <= set(rows[0]):
            raise ValueError("The file needs the columns status_code_identifier and occurrences.")

        identifiers = {
            code.status_code_identifier: code.id
            for code in db.session.query(StatusCode).filter(
                StatusCode.status_code_identifier.isnot(None), StatusCode.is_sandbox == 0
            )
        }
        batch = OccurrenceImport(
            source_label=label,
            period_start=start,
            period_end=end,
            imported_by=actor.id,
            imported_at=utcnow_iso(),
        )
        db.session.add(batch)
        db.session.flush()

        totals: dict[int, int] = {}
        unmatched = 0
        for row in rows:
            identifier = str(row.get("status_code_identifier") or "").strip()
            try:
                count = int(float(row.get("occurrences") or 0))
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"'{row.get('occurrences')}' is not a valid count ({identifier})."
                ) from exc
            if count < 0:
                raise ValueError(f"Negative count for {identifier}.")
            code_id = identifiers.get(identifier)
            if code_id is None:
                unmatched += 1
                continue
            totals[code_id] = totals.get(code_id, 0) + count
        for code_id, count in totals.items():
            db.session.add(
                CodeOccurrence(
                    occurrence_import_id=batch.id, status_code_id=code_id, occurrences=count
                )
            )
        batch.row_count = len(totals)
        batch.unmatched_count = unmatched
        AuditService.log(
            "OccurrenceImport",
            batch.id,
            "OCCURRENCE_IMPORT",
            actor=actor,
            after={
                "source": label,
                "period": f"{start}..{end}",
                "matched_codes": len(totals),
                "unmatched_rows": unmatched,
            },
        )
        db.session.commit()
        return batch

    @staticmethod
    def list_imports() -> list[OccurrenceImport]:
        return db.session.query(OccurrenceImport).order_by(OccurrenceImport.id.desc()).all()

    @staticmethod
    def report(top: int = 20, min_periods: int = 1) -> dict:
        """`never_fired`: Released/Approved codes with zero occurrences across at least
        `min_periods` imported periods. `most_frequent`: the busiest codes, flagged when they
        still have no remedy text."""
        periods = db.session.query(func.count(OccurrenceImport.id)).scalar() or 0
        live = (
            db.session.query(StatusCodeRevision, StatusCode)
            .join(StatusCode, StatusCode.id == StatusCodeRevision.status_code_id)
            .filter(
                StatusCode.is_sandbox == 0,
                StatusCodeRevision.is_current == 1,
                StatusCodeRevision.lifecycle_status.in_(("APPROVED", "RELEASED")),
            )
            .all()
        )
        totals = dict(
            db.session.query(CodeOccurrence.status_code_id, func.sum(CodeOccurrence.occurrences))
            .group_by(CodeOccurrence.status_code_id)
            .all()
        )
        seen_in = dict(
            db.session.query(CodeOccurrence.status_code_id, func.count(CodeOccurrence.id))
            .group_by(CodeOccurrence.status_code_id)
            .all()
        )
        never, frequent = [], []
        for revision, code in live:
            total = int(totals.get(code.id, 0) or 0)
            entry = {
                "status_code_id": code.id,
                "status_code_identifier": code.status_code_identifier,
                "title": revision.title,
                "occurrences": total,
                "periods_reported": int(seen_in.get(code.id, 0)),
                "has_remedy_text": bool(revision.remedy_text),
            }
            if total == 0 and periods >= min_periods:
                never.append(entry)
            elif total > 0:
                frequent.append(entry)
        frequent.sort(key=lambda e: (-e["occurrences"], e["status_code_identifier"] or ""))
        never.sort(key=lambda e: e["status_code_identifier"] or "")
        top_rows = frequent[:top]
        return {
            "periods_imported": periods,
            "live_codes": len(live),
            "never_fired": never,
            "most_frequent": top_rows,
            "needs_remedy_text": [e for e in top_rows if not e["has_remedy_text"]],
        }
