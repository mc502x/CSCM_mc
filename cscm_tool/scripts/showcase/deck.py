"""Slides as plain data, written to .pptx and to preview images (FR-186).

A slide is a list of elements (rectangle, ellipse, chevron, text, picture) at positions in inches on
a 13.333 x 7.5 inch page. Two writers read the same data: python-pptx makes the PowerPoint file and
Pillow draws a preview of every slide, which is also how the build checks its own layout - a text
that does not fit its box stops the build (no PowerPoint is needed to find an overflow). The preview
uses Noto Sans, which is wider than the Calibri of the .pptx, so a text that fits the preview fits
the deck. The previews also make the PDF copy of a deck."""

import re
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

PAGE_W, PAGE_H = 13.333, 7.5
FONT = "Calibri"
FONT_FILE = (
    Path(__file__).resolve().parents[2] / "app" / "static" / "fonts" / "NotoSans-Regular.ttf"
)
BOLD_FACTOR = 1.07  # bold text is wider; the preview widens it by this much


class LayoutError(Exception):
    """A text does not fit the box it was given."""


@dataclass
class Para:
    text: str  # **bold** is supported
    size: float | None = None
    color: str | None = None
    bold: bool = False
    bullet: bool = False
    after: float = 0.0  # extra space after, in points


@dataclass
class Element:
    kind: str  # rect | ellipse | chevron | text | picture
    x: float
    y: float
    w: float
    h: float
    fill: str | None = None
    line: str | None = None
    radius: float = 0.0  # inches, rounded corners of a rectangle
    paras: list[Para] = field(default_factory=list)
    size: float = 14
    color: str = "1A1D21"
    bold: bool = False
    align: str = "l"  # l | c | r
    valign: str = "t"  # t | m | b
    path: str | None = None
    alt: str = ""
    name: str = ""


@dataclass
class Slide:
    title: str
    bg: str = "FFFFFF"
    elements: list[Element] = field(default_factory=list)
    notes: str = ""

    def add(self, element: Element) -> Element:
        self.elements.append(element)
        return element

    def rect(self, x, y, w, h, fill=None, line=None, radius=0.0, name="") -> Element:
        return self.add(Element("rect", x, y, w, h, fill=fill, line=line, radius=radius, name=name))

    def ellipse(self, x, y, w, h, fill, text="", color="FFFFFF", size=16, name="") -> Element:
        paras = [Para(text, bold=True)] if text else []
        return self.add(
            Element(
                "ellipse",
                x,
                y,
                w,
                h,
                fill=fill,
                paras=paras,
                size=size,
                color=color,
                align="c",
                valign="m",
                name=name,
            )
        )

    def chevron(self, x, y, w, h, fill, text, color="FFFFFF", size=14, name="") -> Element:
        return self.add(
            Element(
                "chevron",
                x,
                y,
                w,
                h,
                fill=fill,
                paras=[Para(text, bold=True)],
                size=size,
                color=color,
                align="c",
                valign="m",
                name=name,
            )
        )

    def text(
        self, x, y, w, h, paras, size=14, color="1A1D21", bold=False, align="l", valign="t", name=""
    ) -> Element:
        if isinstance(paras, str):
            paras = [Para(paras)]
        return self.add(
            Element(
                "text",
                x,
                y,
                w,
                h,
                paras=list(paras),
                size=size,
                color=color,
                bold=bold,
                align=align,
                valign=valign,
                name=name,
            )
        )

    def picture(self, x, y, w, h, path, alt, frame=True, name="") -> Element:
        return self.add(
            Element(
                "picture",
                x,
                y,
                w,
                h,
                path=str(path),
                alt=alt,
                line="C9D3E0" if frame else None,
                name=name,
            )
        )


# ------------------------------------------------------------------------------------ measuring
_FONTS: dict[int, ImageFont.FreeTypeFont] = {}


def _font(points: float) -> ImageFont.FreeTypeFont:
    key = int(points * 10)
    if key not in _FONTS:
        _FONTS[key] = ImageFont.truetype(str(FONT_FILE), size=key)
    return _FONTS[key]


def _runs(text: str, bold: bool) -> list[tuple[str, bool]]:
    parts = re.split(r"(\*\*.+?\*\*)", text)
    return [
        (p[2:-2], True) if p.startswith("**") and p.endswith("**") else (p, bold)
        for p in parts
        if p
    ]


def _width(text: str, points: float, bold: bool) -> float:
    """Width of a text in inches."""
    return _font(points).getlength(text) / 10 / 72 * (BOLD_FACTOR if bold else 1.0)


def wrap(para: Para, size: float, bold: bool, width: float) -> list[list[tuple[str, bool]]]:
    """The lines of a paragraph as lists of (text, bold) runs, for a box `width` inches wide."""
    words: list[tuple[str, bool]] = []
    for text, is_bold in _runs(para.text, bold or para.bold):
        for word in re.findall(r"\S+|\s+", text):
            words.append((word, is_bold))
    lines: list[list[tuple[str, bool]]] = [[]]
    used = 0.0
    indent = 0.25 if para.bullet else 0.0
    for word, is_bold in words:
        w = _width(word, size, is_bold)
        if word.isspace():
            if lines[-1]:
                lines[-1].append((word, is_bold))
                used += w
            continue
        if lines[-1] and used + w > width - indent:
            while lines[-1] and lines[-1][-1][0].isspace():
                lines[-1].pop()
            lines.append([])
            used = 0.0
        lines[-1].append((word, is_bold))
        used += w
    while lines and lines[-1] and lines[-1][-1][0].isspace():
        lines[-1].pop()
    return [line for line in lines if line] or [[]]


LINE = 1.2  # line height as a multiple of the size


def layout(element: Element) -> list[tuple[Para, float, list[list[tuple[str, bool]]]]]:
    """[(paragraph, size, lines)] of a text element."""
    out = []
    for para in element.paras:
        size = para.size or element.size
        out.append((para, size, wrap(para, size, element.bold, element.w)))
    return out


def text_height(element: Element) -> float:
    total = 0.0
    for para, size, lines in layout(element):
        total += len(lines) * size * LINE / 72 + para.after / 72
    return total


def problems(slide: Slide) -> list[str]:
    found = []
    for element in slide.elements:
        if element.paras:
            needed = text_height(element)
            if needed > element.h + 0.02:
                first = element.paras[0].text[:50]
                found.append(
                    f"Slide '{slide.title}': the text '{first}' needs {needed:.2f} in, the box has {element.h:.2f} in"
                )
    return found


def check(slide: Slide) -> None:
    found = problems(slide)
    if found:
        raise LayoutError("; ".join(found))


def check_all(slides: list[Slide]) -> None:
    found = [message for slide in slides for message in problems(slide)]
    if found:
        raise LayoutError("; ".join(found))


# ------------------------------------------------------------------------------------ pptx
def _rgb(value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(value)


def to_pptx(
    slides: list[Slide], path: Path, title: str, subject: str, author: str = "CSCM Tool"
) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    check_all(slides)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(PAGE_W), Inches(PAGE_H)
    prs.core_properties.title = title
    prs.core_properties.subject = subject
    prs.core_properties.author = author
    blank = prs.slide_layouts[6]
    aligns = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
    anchors = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

    def fill_text(shape, element: Element) -> None:
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = anchors[element.valign]
        first = True
        for para in element.paras:
            p = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            p.alignment = aligns[element.align]
            if para.after:
                p.space_after = Pt(para.after)
            p.line_spacing = 1.0
            if para.bullet:
                properties = p._p.get_or_add_pPr()
                properties.set("marL", str(int(0.25 * 914400)))
                properties.set("indent", str(-int(0.25 * 914400)))
                bullet = properties.makeelement(qn("a:buChar"), {"char": "•"})
                properties.append(bullet)
            for text, bold in _runs(para.text, element.bold or para.bold):
                run = p.add_run()
                run.text = text
                run.font.name = FONT
                run.font.size = Pt(para.size or element.size)
                run.font.bold = bold
                run.font.color.rgb = _rgb(para.color or element.color)

    for slide in slides:
        page = prs.slides.add_slide(blank)
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = _rgb(slide.bg)
        for element in slide.elements:
            x, y, w, h = (Inches(v) for v in (element.x, element.y, element.w, element.h))
            if element.kind == "picture":
                with Image.open(element.path) as image:
                    ratio = image.width / image.height
                fit_w = min(element.w, element.h * ratio)
                fit_h = fit_w / ratio
                left = element.x + (element.w - fit_w) / 2
                top = element.y + (element.h - fit_h) / 2
                shape = page.shapes.add_picture(
                    element.path, Inches(left), Inches(top), Inches(fit_w), Inches(fit_h)
                )
                shape._element.nvPicPr.cNvPr.set("descr", element.alt)
                if element.name:
                    shape.name = element.name
                if element.line:
                    shape.line.color.rgb = _rgb(element.line)
                    shape.line.width = Pt(0.75)
                continue
            if element.kind == "text":
                shape = page.shapes.add_textbox(x, y, w, h)
                fill_text(shape, element)
            else:
                kind = {
                    "rect": MSO_SHAPE.ROUNDED_RECTANGLE if element.radius else MSO_SHAPE.RECTANGLE,
                    "ellipse": MSO_SHAPE.OVAL,
                    "chevron": MSO_SHAPE.CHEVRON,
                }[element.kind]
                shape = page.shapes.add_shape(kind, x, y, w, h)
                if element.kind == "rect" and element.radius:
                    shape.adjustments[0] = min(0.5, element.radius / min(element.w, element.h))
                if element.fill:
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = _rgb(element.fill)
                else:
                    shape.fill.background()
                if element.line:
                    shape.line.color.rgb = _rgb(element.line)
                    shape.line.width = Pt(0.75)
                else:
                    shape.line.fill.background()
                shape.shadow.inherit = False
                if element.paras:
                    fill_text(shape, element)
            if element.name:
                shape.name = element.name
        if slide.notes:
            page.notes_slide.notes_text_frame.text = slide.notes
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


# ------------------------------------------------------------------------------------ preview
def render(slide: Slide, dpi: int = 100) -> Image.Image:
    """A picture of one slide, drawn from the same data as the .pptx."""
    check(slide)
    final = (round(PAGE_W * dpi), round(PAGE_H * dpi))
    dpi *= 2  # drawn twice as large and scaled down: small text keeps its spacing
    image = Image.new("RGB", (round(PAGE_W * dpi), round(PAGE_H * dpi)), "#" + slide.bg)
    draw = ImageDraw.Draw(image)

    def px(value: float) -> int:
        return round(value * dpi)

    for element in slide.elements:
        box = (px(element.x), px(element.y), px(element.x + element.w), px(element.y + element.h))
        if element.kind == "picture":
            with Image.open(element.path) as source:
                source = source.convert("RGBA")
                ratio = source.width / source.height
                fit_w = min(element.w, element.h * ratio)
                fit_h = fit_w / ratio
                left = element.x + (element.w - fit_w) / 2
                top = element.y + (element.h - fit_h) / 2
                resized = source.resize((px(fit_w), px(fit_h)), Image.LANCZOS)
            image.paste(resized, (px(left), px(top)), resized)  # the logo has transparent corners
            if element.line:
                draw.rectangle(
                    (px(left), px(top), px(left + fit_w), px(top + fit_h)),
                    outline="#" + element.line,
                )
            continue
        if element.kind in ("rect", "ellipse", "chevron"):
            fill = "#" + element.fill if element.fill else None
            outline = "#" + element.line if element.line else None
            if element.kind == "ellipse":
                draw.ellipse(box, fill=fill, outline=outline)
            elif element.kind == "chevron":
                notch = min(element.h / 2, element.w / 3)
                points = [
                    (element.x, element.y),
                    (element.x + element.w - notch, element.y),
                    (element.x + element.w, element.y + element.h / 2),
                    (element.x + element.w - notch, element.y + element.h),
                    (element.x, element.y + element.h),
                    (element.x + notch, element.y + element.h / 2),
                ]
                draw.polygon([(px(a), px(b)) for a, b in points], fill=fill, outline=outline)
            elif element.radius:
                draw.rounded_rectangle(box, radius=px(element.radius), fill=fill, outline=outline)
            else:
                draw.rectangle(box, fill=fill, outline=outline)
        if element.paras:
            _draw_text(draw, element, dpi)
    return image.resize(final, Image.LANCZOS)


def _draw_text(draw: ImageDraw.ImageDraw, element: Element, dpi: int) -> None:
    blocks = layout(element)
    total = text_height(element)
    top = element.y
    if element.valign == "m":
        top = element.y + (element.h - total) / 2
    elif element.valign == "b":
        top = element.y + element.h - total
    y = top
    for para, size, lines in blocks:
        font = ImageFont.truetype(str(FONT_FILE), size=max(6, round(size / 72 * dpi)))
        color = "#" + (para.color or element.color)
        for number, line in enumerate(lines):
            width = sum(_width(t, size, b) for t, b in line)
            indent = 0.25 if para.bullet else 0.0
            if element.align == "c":
                x = element.x + (element.w - width) / 2
            elif element.align == "r":
                x = element.x + element.w - width
            else:
                x = element.x + indent
            if para.bullet and number == 0:
                draw.text((round(element.x * dpi), round(y * dpi)), "•", font=font, fill=color)
            for text, bold in line:
                pos = (round(x * dpi), round(y * dpi))
                draw.text(pos, text, font=font, fill=color)
                if bold:
                    draw.text((pos[0] + 1, pos[1]), text, font=font, fill=color)
                x += _width(text, size, bold)
            y += size * LINE / 72
        y += para.after / 72


def to_previews(slides: list[Slide], folder: Path, dpi: int = 100) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for number, slide in enumerate(slides, 1):
        target = folder / f"slide-{number:02d}.png"
        render(slide, dpi).save(target)
        paths.append(target)
    return paths


def to_pdf(slides: list[Slide], path: Path, dpi: int = 150) -> None:
    pages = [render(slide, dpi) for slide in slides]
    path.parent.mkdir(parents=True, exist_ok=True)
    pages[0].save(path, "PDF", resolution=dpi, save_all=True, append_images=pages[1:], quality=88)
