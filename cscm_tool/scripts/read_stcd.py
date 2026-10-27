"""Reads the status code (StCd) definitions from the generated Structured Text programs.

    python scripts/read_stcd.py mc_dir/ --hc hc_dir/ --format xlsx -o stcd.xlsx
    python scripts/read_stcd.py StCd_1000.st --hc StCd_14000.st --format json -o stcd.json
    python scripts/read_stcd.py --hc hc_dir/ --controller hc

Each file is one `PROGRAM StCd_<n>` written by StatuscodeGeneratorEuros. Per program the script
reads:

  * the header comment: legacy number and name (`StCd: 3000, 'Umrichter nicht bereit'`), the
    source database and the generator,
  * the attribute block between `(* Attribute aus DB *)` and `FirstRun := FALSE;`: the number
    (`StCdArUi[Idx]`), the repeat code and its legacy number, repeat limits, task frequency,
    operating-mode mask, alarm type, availability group, snapshot flag, the per-subsystem
    `*PrgArUs` settings and the read/write level of every user group,
  * the set condition (`StatusSetArGBo[Idx] := ...;`), the set and reset delays, the
    `StCd<n>DisableGpBo` parameter and the other StCd programs it references.

Main controller (MC) numbers in the files are the current ones; the old number is only in the
header. The numbers in Hub controller (HC) files are the legacy HC numbers: they are looked up in
the sheet HC of the mapping workbook seed_data/legacy_mapping.xlsx (`--mapping`; the application
keeps the mapping in its database, docs/03-srs.md FR-191). Paths given with `--hc` are HC; the other paths
count as HC only from `--hc-from` upwards, which cannot tell a four-digit legacy number from an
MC one, so pass HC files with `--hc`. Files are read as UTF-8, then CP1252 (the generator's
encoding). Directories are searched recursively for SUFFIXES. Attributes the script does not
know by name are kept under `other`. Needs openpyxl for the mapping and `--format xlsx`.

The reading itself lives in app/importers/stcd.py (shared with the application's status code import,
 docs/03-srs.md FR-048o); this file is only the command-line front end."""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from dataclasses import asdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.importers.stcd import (  # noqa: E402
    HC_FROM,
    SUFFIXES,
    StatusCode,
    apply_mapping,
    decode,
    load_mapping,
    parse_text,
)

MAPPING = Path(__file__).resolve().parents[1] / "seed_data" / "legacy_mapping.xlsx"


def _read(path: Path) -> str:
    return decode(path.read_bytes())


def _files(paths: list[Path]) -> list[Path]:
    out: list[Path] = []
    for path in paths:
        if path.is_dir():
            out += sorted(
                p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in SUFFIXES
            )
        else:
            out.append(path)
    return out


_FIRST = [
    "identifier",
    "controller",
    "legacy_number",
    "name",
    "source_database",
    "alarm",
    "alarm_type",
    "set_condition",
    "set_delay",
    "set_delay_ms",
    "reset_delay",
    "reset_delay_ms",
]


def _flat(code: StatusCode) -> dict[str, object]:
    row = asdict(code)
    for key in ("depends_on", "programs", "user_groups", "other"):
        row.pop(key)
    row["depends_on"] = " ".join(map(str, code.depends_on))
    row.update({f"prg_{k}": v for k, v in code.programs.items()})
    for group, levels in code.user_groups.items():
        row.update({f"ug{group}_{k}": v for k, v in levels.items()})
    row.update({f"other_{k}": v for k, v in code.other.items()})
    return row


def _columns(flat: list[dict[str, object]]) -> list[str]:
    """Key columns first, then the remaining fixed fields, then prg_/ug/other_ in natural order."""
    keys = {k for r in flat for k in r}
    first = [k for k in _FIRST if k in keys]
    dynamic = {k for k in keys if k.startswith(("prg_", "ug", "other_"))}
    fixed = [k for k in (flat[0] if flat else []) if k not in first and k not in dynamic]
    natural = lambda k: [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", k)]  # noqa: E731
    return first + fixed + sorted(dynamic, key=natural)


def _csv(rows: list[StatusCode]) -> str:
    flat = [_flat(r) for r in rows]
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=_columns(flat), lineterminator="\n", restval="")
    writer.writeheader()
    writer.writerows(flat)
    return buf.getvalue()


def _write_xlsx(rows: list[StatusCode], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def sheet(ws, header: list[str], body: list[list[object]]) -> None:
        ws.append(header)
        for row in body:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for i, name in enumerate(header, 1):
            width = max([len(name)] + [len(str(r[i - 1])) for r in body if r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

    flat = [_flat(r) for r in rows]
    columns = _columns(flat)
    wb = Workbook()
    sheet(wb.active, columns, [[r.get(c) for c in columns] for r in flat])
    wb.active.title = "Status codes"
    groups = [
        [r.identifier, r.controller, r.legacy_number, g, lv.get("read"), lv.get("write")]
        for r in rows
        for g, lv in r.user_groups.items()
    ]
    sheet(
        wb.create_sheet("User groups"),
        ["identifier", "controller", "legacy_number", "user_group", "read", "write"],
        groups,
    )
    wb.save(path)


def _table(rows: list[StatusCode]) -> str:
    head = (
        "StCd",
        "Ctrl",
        "Legacy",
        "Name",
        "Alarm",
        "Set delay",
        "Reset delay",
        "Repeat",
        "Source",
    )
    body = [
        (
            r.identifier,
            r.controller,
            str(r.legacy_number or ""),
            r.name,
            "" if r.alarm_type is None else str(r.alarm_type),
            r.set_delay,
            r.reset_delay,
            "" if r.repeat_code is None else str(r.repeat_code),
            r.source_database,
        )
        for r in rows
    ]
    widths = [max(len(c) for c in col) for col in zip(head, *body)]
    fmt = "  ".join(f"{{:<{w}}}" for w in widths)
    lines = [fmt.format(*head), fmt.format(*("-" * w for w in widths))]
    lines += [fmt.format(*row).rstrip() for row in body]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Read status code definitions from StCd_*.st files.")
    ap.add_argument("paths", nargs="*", type=Path, help="MC files or directories")
    ap.add_argument(
        "--hc",
        nargs="+",
        type=Path,
        default=[],
        metavar="PATH",
        help="files or directories holding Hub controller programs (legacy numbers)",
    )
    ap.add_argument("--format", choices=("table", "json", "csv", "xlsx"), default="table")
    ap.add_argument("-o", "--output", type=Path, help="write to this file instead of stdout")
    ap.add_argument("--controller", choices=("all", "mc", "hc"), default="all")
    ap.add_argument(
        "--hc-from",
        type=int,
        default=HC_FROM,
        help=f"first number of a non---hc path counted as HC (default {HC_FROM})",
    )
    ap.add_argument(
        "--mapping", type=Path, default=MAPPING, help="legacy HC number mapping workbook"
    )
    args = ap.parse_args(argv)

    inputs = [(p, None) for p in _files(args.paths)] + [(p, True) for p in _files(args.hc)]
    if not inputs:
        print("no input files found", file=sys.stderr)
        return 2
    if args.format == "xlsx" and not args.output:
        print("--format xlsx needs -o", file=sys.stderr)
        return 2

    rows: list[StatusCode] = []
    for path, hc in inputs:
        try:
            found = parse_text(_read(path), str(path), hc, args.hc_from)
        except OSError as exc:
            print(f"{path}: {exc}", file=sys.stderr)
            return 2
        if not found:
            print(f"{path}: no PROGRAM StCd_<n> found", file=sys.stderr)
        rows += found

    mapping: dict[int, int] = {}
    try:
        mapping = load_mapping(args.mapping)
    except (OSError, KeyError, ValueError, ImportError) as exc:
        if any(r.controller == "HC" for r in rows):
            print(f"mapping {args.mapping}: {exc}", file=sys.stderr)
            return 2
    for r in rows:
        if r.controller == "HC":
            apply_mapping(r, mapping)
        elif r.number in mapping:
            print(
                f"{r.file}: {r.number} is also a legacy HC number; pass the file with --hc "
                "if it is a Hub controller program",
                file=sys.stderr,
            )

    for r in rows:
        for warning in r.warnings:
            print(f"{r.file}:{r.line}: {warning}", file=sys.stderr)
    seen: dict[tuple[str, int], str] = {}
    for r in rows:
        key = (r.controller, r.number)
        if r.identifier and key in seen:
            print(f"{r.file}: {r.identifier} already defined in {seen[key]}", file=sys.stderr)
        seen.setdefault(key, r.file)
    if args.controller != "all":
        rows = [r for r in rows if r.controller == args.controller.upper()]
    rows.sort(key=lambda r: (r.controller != "MC", r.number))

    if args.format == "xlsx":
        _write_xlsx(rows, args.output)
        print(f"{len(rows)} status codes -> {args.output}", file=sys.stderr)
        return 0
    if args.format == "json":
        out = json.dumps([asdict(r) for r in rows], indent=2, ensure_ascii=False)
    elif args.format == "csv":
        out = _csv(rows)
    else:
        out = _table(rows)

    if args.output:
        args.output.write_text(out + ("" if out.endswith("\n") else "\n"), encoding="utf-8")
        print(f"{len(rows)} status codes -> {args.output}", file=sys.stderr)
    else:
        sys.stdout.reconfigure(errors="replace")
        print(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
