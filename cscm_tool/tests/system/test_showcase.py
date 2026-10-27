"""The design showcase decks (FR-186, scripts/build_showcase.py): the slide model and its layout
check, both decks from placeholder pictures, the PowerPoint files, the throw-away demo and the
choice of the browser. The screenshots themselves need Edge or Chrome and are not part of the tests.

TC-073. Requirements: FR-186, FR-187.
"""

import importlib.util
import sys
from pathlib import Path

import pytest
from PIL import Image

from app import branding, create_app

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "scripts"
LOGO = ROOT / "app" / "static" / "img" / "cscm-logo-512.png"


@pytest.fixture(scope="module")
def showcase():
    sys.path.insert(0, str(SCRIPTS))
    from showcase import capture, deck, decks, demo  # noqa: PLC0415

    return type("Showcase", (), {"capture": capture, "deck": deck, "decks": decks, "demo": demo})


@pytest.fixture(scope="module")
def images(showcase, tmp_path_factory):
    """A plain picture for every screenshot the decks name, and the logo."""
    folder = tmp_path_factory.mktemp("pictures")
    ids = {
        "released_code": 1,
        "pending_cr": 1,
        "change_set": 1,
        "open_release": 1,
        "mapping_change": 1,
    }
    found = {}
    for shot in showcase.decks.shots(ids):
        path = folder / f"{shot.name}.png"
        Image.new("RGB", shot.size, "#dde6f2").save(path)
        found[shot.name] = path
    found["logo"] = LOGO
    return found


def test_a_text_that_does_not_fit_stops_the_build(showcase):
    deck = showcase.deck
    slide = deck.Slide("Too much")
    slide.text(1, 1, 2, 0.4, [deck.Para("word " * 80)], size=14)
    with pytest.raises(deck.LayoutError, match="Too much"):
        deck.check_all([slide])
    fits = deck.Slide("Fine")
    fits.text(1, 1, 6, 1.5, [deck.Para("A **short** text that fits easily.")], size=14)
    deck.check_all([fits])


def test_both_decks_are_built_from_the_screenshots_and_every_picture_is_described(showcase, images):
    facts = showcase.decks.gather("9.9.9", "2026-01-01")
    assert facts.version == "9.9.9" and facts.tests != "n/a" and facts.requirements != "n/a"
    assert facts.coverage.endswith("%") and facts.delivered
    for build, count in ((showcase.decks.management, 11), (showcase.decks.team, 24)):
        slides = build(facts, images)
        showcase.deck.check_all(slides)  # every text fits its box
        assert len(slides) == count
        pictures = [e for s in slides for e in s.elements if e.kind == "picture"]
        assert pictures and all(p.alt.strip() for p in pictures)
        first = " ".join(p.text for e in slides[0].elements for p in e.paras)
        assert branding.OWNER_LINE in first and branding.LONG_NAME in first
        assert branding.PRODUCT in first and "Version 9.9.9" in first
        assert any(e.name == "Logo" for e in slides[0].elements)
        assert all(s.title for s in slides)


def test_the_decks_are_written_as_powerpoint_files(showcase, images, tmp_path):
    from pptx import Presentation

    facts = showcase.decks.gather("9.9.9", "2026-01-01")
    slides = showcase.decks.management(facts, images)
    target = tmp_path / "deck.pptx"
    showcase.deck.to_pptx(slides, target, "CSCM Tool - test", "test")
    prs = Presentation(str(target))
    assert len(prs.slides) == len(slides) and prs.core_properties.title == "CSCM Tool - test"
    assert abs(prs.slide_width / 914400 - 13.333) < 0.01
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert any("Version" in n for n in notes)
    described = [
        shape._element.nvPicPr.cNvPr.get("descr")
        for s in prs.slides
        for shape in s.shapes
        if shape.shape_type == 13
    ]
    assert described and all(described)
    previews = showcase.deck.to_previews(slides[:2], tmp_path / "previews", dpi=40)
    assert [p.name for p in previews] == ["slide-01.png", "slide-02.png"]
    showcase.deck.to_pdf(slides[:2], tmp_path / "deck.pdf", dpi=40)
    assert (tmp_path / "deck.pdf").read_bytes().startswith(b"%PDF")


def test_the_demo_has_its_data_and_the_application_has_no_demo_routes(showcase):
    assert not [r for r in create_app("testing").url_map.iter_rules() if "/_show" in r.rule]
    demo = showcase.demo.build()
    try:
        client = demo.app.test_client()
        signed_in = client.get("/_show/as/admin?next=/status-codes", follow_redirects=True)
        page = signed_in.get_data(as_text=True)
        assert signed_in.status_code == 200 and "Converter grid undervoltage trip" in page
        assert {"change_set", "release", "released_code", "pending_cr", "mapping_change"} <= set(
            demo.ids
        )
        dark = client.get("/_show/theme/dark?next=/login")
        assert dark.status_code == 302
        script = client.get("/static/js/theme.js").get_data(as_text=True)
        assert "setAttribute('data-theme','dark')" in script
        assert "Preview" in client.get("/_show/st-preview").get_data(as_text=True)
    finally:
        demo.stop()


def test_the_browser_is_found_or_named(showcase, monkeypatch, tmp_path):
    capture = showcase.capture
    fake = tmp_path / "browser.exe"
    fake.write_bytes(b"x")
    monkeypatch.setenv("CSCM_BROWSER", str(fake))
    assert capture.find_browser() == str(fake)
    monkeypatch.delenv("CSCM_BROWSER")
    monkeypatch.setattr(capture, "_WINDOWS", ())
    monkeypatch.setattr(capture.shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="CSCM_BROWSER"):
        capture.find_browser()


def test_the_script_refuses_an_unknown_deck(capsys):
    spec = importlib.util.spec_from_file_location("build_showcase", SCRIPTS / "build_showcase.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    assert module.main(["--only", "bogus"]) == 2
    assert "Unknown deck" in capsys.readouterr().out
