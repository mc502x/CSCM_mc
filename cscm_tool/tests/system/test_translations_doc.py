"""The translation catalogue document and the review sheets for native speakers
(docs/26-translations.md, scripts/gen_translations_doc.py). The document is generated from the
catalogues, so it cannot drift. The review itself is done in the application (S42-S44, tested in
tests/integration/test_translation_review.py).

TC-057. Requirements: FR-161, FR-110, FR-111, FR-155, FR-157.
"""

import io
import re
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

APP_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(APP_DIR / "scripts"))
import gen_translations_doc as gen  # noqa: E402

DOC = APP_DIR.parent / "docs" / "26-translations.md"


def test_the_document_is_exactly_what_the_generator_produces_from_the_code():
    assert gen.is_current(), "run: python scripts/gen_translations_doc.py --write"


def test_every_language_has_its_section_and_its_coverage_row():
    text = DOC.read_text(encoding="utf-8")
    table = text.split("### 9.1 Coverage", 1)[1].split("###", 1)[0]
    cat = gen.catalogue()
    total = len(gen._screens())
    for code in gen.languages():
        assert f"`{code}`" in table, code
        if code == "en":
            continue
        assert f"{gen.label(code)} — {gen.ENGLISH_NAMES[code]}" in text, code
        row = next(line for line in table.splitlines() if f"| `{code}` |" in line)
        assert f"| {len(cat[code])} |" in row and f"`{gen.fingerprint(cat[code])}`" in row
    assert f"**{total} strings**" in text
    assert cat["de"] and len(cat["de"]) == total  # German is complete


def test_string_ids_are_stable_and_never_collide():
    strings = list(gen._screens())
    ids = {gen.string_id(s) for s in strings}
    assert len(ids) == len(strings)
    assert gen.string_id("Log in") == gen.string_id("Log in")
    assert re.fullmatch(r"S[0-9a-f]{6}", gen.string_id("Log in"))


def test_the_fingerprint_changes_when_one_translation_changes():
    t = dict(gen.catalogue()["fr"])
    before = gen.fingerprint(t)
    assert re.fullmatch(r"[0-9a-f]{12}", before)
    t["Log in"] = t["Log in"] + "!"
    assert gen.fingerprint(t) != before
    assert gen.fingerprint(gen.catalogue()["fr"]) == before  # and it is reproducible


@pytest.mark.parametrize("code", [c for c in gen.languages() if c != "en"])
def test_the_review_sheet_of_every_language_is_complete(code):
    workbook = load_workbook(io.BytesIO(gen.build_sheet(code)))
    assert workbook.sheetnames == ["Review", "Missing", "Checklist"]
    cat = gen.catalogue()[code]
    screens = gen._screens()
    review = list(workbook["Review"].iter_rows(min_row=2, values_only=True))
    assert len(review) == len(cat)
    assert {r[0] for r in review} == {gen.string_id(e) for e in cat}
    by_english = {r[2]: r[3] for r in review}
    assert by_english == cat  # the current translation of every string is on the sheet
    missing = list(workbook["Missing"].iter_rows(min_row=2, values_only=True))
    assert len(missing) == len(screens) - len(cat)
    checklist = " ".join(
        str(c) for row in workbook["Checklist"].iter_rows(values_only=True) for c in row if c
    )
    assert (
        gen.fingerprint(cat) in checklist and "Meaning" in checklist and "Placeholders" in checklist
    )
    # the verdict column offers exactly the three answers
    validations = workbook["Review"].data_validations.dataValidation
    assert validations and gen.VERDICTS == ("OK", "CHANGE", "UNCLEAR")
    assert all(v in validations[0].formula1 for v in gen.VERDICTS)


@pytest.mark.parametrize("code,expected", [("ar", "right"), ("ur", "right"), ("fr", "left")])
def test_right_to_left_languages_get_a_right_to_left_translation_column(code, expected):
    sheet = load_workbook(io.BytesIO(gen.build_sheet(code)))["Review"]
    assert sheet["D2"].alignment.horizontal == expected


def test_the_command_line_writes_one_sheet_per_language_and_refuses_unknown_ones(tmp_path):
    assert gen.main(["--xlsx", "all", "-o", str(tmp_path)]) == 0
    names = sorted(p.name for p in tmp_path.glob("*.xlsx"))
    assert len(names) == len(gen.languages()) - 1 and "translation-review-ka.xlsx" in names
    assert gen.main(["--xlsx", "xx", "-o", str(tmp_path)]) == 2
    assert gen.main(["--xlsx", "en", "-o", str(tmp_path)]) == 2  # the source has no sheet
    assert gen.main(["--check"]) == 0


def test_the_concept_names_the_layers_the_checks_and_the_open_items():
    text = DOC.read_text(encoding="utf-8")
    for heading in (
        "## 3. The concept: how a native speaker checks a language",
        "### 3.3 The checklist",
        "## 4. Status code manuals in all designated languages",
        "### 4.3 Terminology shared by the interface and the content",
        "## 5. Text outside the interface catalogue (layer C)",
        "## 8. Decisions and open items",
        "## 9. Catalogue (generated)",
    ):
        assert heading in text, heading
    assert "scripts/gen_translations_doc.py" in text and "FR-161" in text
