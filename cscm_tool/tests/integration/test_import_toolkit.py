"""The migration toolkit around the library import (docs/21-improvement-ideas.md): I-05 template, dry
run, update mode, reconciliation and error report; I-12 duplicate warnings; I-32 background jobs.
The basic import contract (every row becomes a Draft) stays in test_import.py.

TC-043. Requirements: FR-048i, FR-048j, FR-048k, FR-048l, FR-048m, FR-048n, FR-093, BR-015.
"""

import csv
import io
import time

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.change_request import ChangeRequest
from app.models.status_code import StatusCode, StatusCodeRevision
from app.services.import_service import TEMPLATE_HEADERS
from tests.helpers import (
    PASSWORD,
    cr_id_for_revision,
    domain_sign,
    drive_to_approved,
    login_as,
    make_cast,
    release_revisions,
    submit,
)


def _csv(rows):
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def _row(legacy="2001", title="Imported grid trip", **overrides):
    row = {
        "functional_system_group": "WCNV",
        "functional_subgroup": "010xx",
        "turbine_platforms": "3XM;4XM",
        "title": title,
        "description": "Imported through the migration toolkit for tests.",
        "status_category": "ERROR",
        "available_group": "27",
        "set_delay": "10ms",
        "reset_delay": "2min",
        "legacy_reference_number": legacy,
    }
    row.update(overrides)
    return row


def _post(client, rows, **fields):
    return client.post(
        "/api/v1/imports",
        data={"file": (io.BytesIO(_csv(rows)), "library.csv"), **fields},
        content_type="multipart/form-data",
    )


@pytest.fixture()
def engineer(client, make_user):
    user = make_user("ENGINEER")
    login_as(client, user)
    return user


# ---------------------------------------------------------------- template (I-05b)
def test_template_csv_and_xlsx_carry_the_exact_headers_and_lookup_codes(client, engineer):
    csv_template = client.get("/api/v1/imports/template?format=csv")
    header = csv_template.get_data(as_text=True).splitlines()[0].lstrip("﻿").split(",")
    assert tuple(header) == TEMPLATE_HEADERS
    workbook = load_workbook(io.BytesIO(client.get("/api/v1/imports/template?format=xlsx").data))
    assert workbook.sheetnames == ["Import", "Instructions", "Lookups"]
    assert tuple(c.value for c in workbook["Import"][1]) == TEMPLATE_HEADERS
    lookups = [str(c.value) for row in workbook["Lookups"].iter_rows() for c in row if c.value]
    for code in ("WCNV", "010xx", "3XM", "ERROR", "SCADA_USER", "COVERT"):
        assert code in lookups
    assert client.get("/api/v1/imports/template?format=pdf").status_code == 422


def test_the_downloaded_template_imports_cleanly_as_is(client, engineer):
    template = client.get("/api/v1/imports/template?format=csv").data
    result = client.post(
        "/api/v1/imports",
        data={"file": (io.BytesIO(template), "template.csv")},
        content_type="multipart/form-data",
    ).get_json()
    assert result["created_count"] == 1 and result["error_count"] == 0  # the example row is valid


# ---------------------------------------------------------------- dry run (I-05a)
def test_dry_run_reports_every_problem_and_creates_nothing(app, client, engineer):
    rows = [
        _row("1", "Good row"),
        _row("2", "Bad delay", set_delay="soon", status_category="PANIC"),
        _row("3", "Unknown group", functional_system_group="NOPE", turbine_platforms="9XM"),
    ]
    report = _post(client, rows, dry_run="true").get_json()
    assert report["ok"] is False and report["would_create"] == 1 and report["error_count"] == 2
    bad = report["rows"][1]
    assert sorted(e["column"] for e in bad["errors"]) == [
        "set_delay",
        "status_category",
    ]  # all of them
    assert all(e["hint"] for e in bad["errors"])
    assert {e["column"]: e["value"] for e in bad["errors"]}["set_delay"] == "soon"
    columns = {e["column"] for e in report["rows"][2]["errors"]}
    assert {"functional_system_group", "turbine_platforms"} <= columns
    with app.app_context():
        assert db.session.query(StatusCode).count() == 0  # nothing was created
        from app.models.import_job import ImportJob

        assert db.session.query(ImportJob).count() == 0  # not even a job


def test_dry_run_warns_about_a_possible_duplicate_without_blocking(app, client, engineer):
    _post(client, [_row("10", "Converter overcurrent trip")])
    report = _post(client, [_row("11", "Converter over-current trip")], dry_run="true").get_json()
    assert report["ok"] is True
    assert report["warning_count"] == 1 and "alike" in report["rows"][0]["warnings"][0]


def test_missing_required_columns_are_named_up_front(client, engineer):
    response = client.post(
        "/api/v1/imports",
        data={"file": (io.BytesIO(b"title,description\nx,y\n"), "bad.csv")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 422
    assert "functional_system_group" in response.get_json()["message"]


# ---------------------------------------------------------------- error report (I-05e)
def test_error_report_has_row_column_value_and_a_fix_hint_as_csv_and_xlsx(client, engineer):
    body = _post(
        client, [_row("1"), _row("2", set_delay="soon", status_category="PANIC")]
    ).get_json()
    assert body["error_count"] == 1 and body["error_report_xlsx_available"] is True
    lines = list(
        csv.DictReader(
            io.StringIO(client.get(f"/api/v1/imports/{body['id']}/errors").get_data(as_text=True))
        )
    )
    assert sorted((r["row"], r["column"], r["value"]) for r in lines) == [
        ("3", "set_delay", "soon"),
        ("3", "status_category", "PANIC"),
    ]
    assert "ms" in next(r for r in lines if r["column"] == "set_delay")["how_to_fix"]
    workbook = load_workbook(
        io.BytesIO(client.get(f"/api/v1/imports/{body['id']}/errors?format=xlsx").data)
    )
    sheet = workbook["Errors"]
    assert [c.value for c in sheet[1]] == ["Row", "Column", "Value", "Problem", "How to fix"]
    assert sheet.max_row == 3


def test_formula_like_values_in_the_error_report_are_neutralised(client, engineer):
    body = _post(client, [_row("1", set_delay="=cmd|' /C calc'!A0")]).get_json()
    report = client.get(f"/api/v1/imports/{body['id']}/errors").get_data(as_text=True)
    assert "'=cmd" in report


# ---------------------------------------------------------------- create mode + duplicates (I-05, I-12)
def test_a_legacy_number_already_imported_is_refused_in_create_mode(client, engineer):
    assert _post(client, [_row("500")]).get_json()["created_count"] == 1
    again = _post(client, [_row("500", "Same number again")]).get_json()
    assert again["created_count"] == 0 and again["error_count"] == 1
    text = client.get(f"/api/v1/imports/{again['id']}/errors").get_data(as_text=True)
    assert "Already imported" in text and "update mode" in text
    same_file = _post(client, [_row("600"), _row("600", "Twice in one file")]).get_json()
    assert same_file["created_count"] == 1 and same_file["error_count"] == 1


def test_the_same_legacy_number_in_another_group_is_a_different_code(client, engineer):
    both = _post(
        client,
        [
            _row("7", "Converter one"),
            _row("7", "Generator one", functional_system_group="WGEN", functional_subgroup="020xx"),
        ],
    ).get_json()
    assert both["created_count"] == 2 and both["error_count"] == 0


def test_similar_titles_are_reported_as_warnings(client, engineer):
    _post(client, [_row("20", "Grid undervoltage trip")])
    job = _post(client, [_row("21", "Grid under-voltage trip")]).get_json()
    assert job["created_count"] == 1 and job["warning_count"] == 1


def test_new_content_columns_are_imported(app, client, engineer):
    _post(client, [_row("30", remedy_text="Reset the converter.", affected_component="DC link")])
    with app.app_context():
        revision = db.session.query(StatusCodeRevision).one()
        assert (
            revision.remedy_text == "Reset the converter."
            and revision.affected_component == "DC link"
        )


# ---------------------------------------------------------------- update mode (I-05c)
def test_update_mode_revises_the_open_draft_in_place(app, client, engineer):
    _post(client, [_row("40", "Original title")])
    report = _post(
        client, [_row("40", "Corrected title")], mode="UPDATE", dry_run="true"
    ).get_json()
    assert report["would_update"] == 1 and report["would_create"] == 0
    job = _post(client, [_row("40", "Corrected title")], mode="UPDATE").get_json()
    assert job["created_count"] == 1 and job["error_count"] == 0 and job["import_mode"] == "UPDATE"
    with app.app_context():
        assert db.session.query(StatusCode).count() == 1  # no duplicate
        revision = db.session.query(StatusCodeRevision).one()
        assert revision.title == "Corrected title" and revision.row_version == 2


def test_update_mode_skips_unchanged_rows_and_creates_missing_ones(app, client, engineer):
    _post(client, [_row("41", "Stable")])
    job = _post(client, [_row("41", "Stable"), _row("42", "Brand new")], mode="UPDATE").get_json()
    assert job["created_count"] == 1  # only the new row; the unchanged one is skipped
    with app.app_context():
        assert db.session.query(StatusCode).count() == 2


def test_update_mode_opens_a_new_revision_of_a_released_code(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    revision, _ = drive_to_approved(
        app, client, cast, wcnv_ids, title="Released legacy code", legacy_reference_number="50"
    )
    release_revisions(client, cast, [revision["id"]], label="MIG1")
    login_as(client, cast["engineer"])
    job = _post(
        client, [_row("50", "Released legacy code", set_delay="30ms")], mode="UPDATE"
    ).get_json()
    assert job["created_count"] == 1 and job["error_count"] == 0
    with app.app_context():
        revisions = (
            db.session.query(StatusCodeRevision).order_by(StatusCodeRevision.revision_number).all()
        )
        assert [(r.revision_number, r.lifecycle_status, r.set_delay) for r in revisions] == [
            (1, "RELEASED", "10ms"),
            (2, "DRAFT", "30ms"),
        ]
        cr = (
            db.session.query(ChangeRequest).filter_by(status_code_revision_id=revisions[1].id).one()
        )
        assert cr.cr_type == "REVISION" and cr.import_job_id == job["id"]


def test_update_mode_errors_are_specific(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    login_as(client, cast["engineer"])
    _post(client, [_row("60", "In review soon")])
    with app.app_context():
        revision_id = db.session.query(StatusCodeRevision).one().id
    cr_id = cr_id_for_revision(app, revision_id)
    domain_sign(client, cast, cr_id)
    assert submit(client, cast, cr_id).status_code == 200  # the imported draft is now in review
    login_as(client, cast["engineer"])
    locked = _post(
        client, [_row("60", "In review soon", set_delay="40ms")], mode="UPDATE"
    ).get_json()
    assert locked["error_count"] == 1
    assert "in review" in client.get(f"/api/v1/imports/{locked['id']}/errors").get_data(
        as_text=True
    )
    no_key = _post(client, [_row("", "No legacy number")], mode="UPDATE").get_json()
    assert no_key["error_count"] == 1
    moved = _post(
        client, [_row("60", "In review soon", functional_subgroup="012xx")], mode="UPDATE"
    ).get_json()
    assert moved["error_count"] == 1
    assert "cannot change" in client.get(f"/api/v1/imports/{moved['id']}/errors").get_data(
        as_text=True
    )
    assert client.post("/api/v1/imports", data={"mode": "X"}).status_code == 422


def test_update_of_someone_elses_draft_is_reported_on_the_row(app, client, make_user):
    first = make_user("ENGINEER")
    second = make_user("ENGINEER")
    login_as(client, first)
    _post(client, [_row("70", "Mine")])
    login_as(client, second)
    job = _post(client, [_row("70", "Theirs now")], mode="UPDATE").get_json()
    assert job["error_count"] == 1
    assert "owning Engineer" in client.get(f"/api/v1/imports/{job['id']}/errors").get_data(
        as_text=True
    )


# ---------------------------------------------------------------- reconciliation (I-05d)
def test_reconciliation_counts_states_and_names_the_missing_numbers(
    app, client, make_user, wcnv_ids
):
    cast = make_cast(make_user)
    released, _ = drive_to_approved(
        app, client, cast, wcnv_ids, title="Already live", legacy_reference_number="1"
    )
    release_revisions(client, cast, [released["id"]], label="MIG2")
    login_as(client, cast["engineer"])
    _post(client, [_row("2", "Imported draft")])

    report = client.get("/api/v1/imports/reconciliation?expected_total=5").get_json()
    assert report["imported_total"] == 2
    assert report["by_state"]["released"] == 1 and report["by_state"]["imported_draft"] == 1
    assert report["expected_total"] == 5 and report["missing_count"] == 3

    master = _csv([{"legacy_reference_number": n} for n in ("1", "2", "3", "4")])
    named = client.post(
        "/api/v1/imports/reconciliation",
        data={"file": (io.BytesIO(master), "master.csv")},
        content_type="multipart/form-data",
    ).get_json()
    assert named["missing"] == ["3", "4"] and named["percent_imported"] == 50.0
    assert client.post("/api/v1/imports/reconciliation").status_code == 422


# ---------------------------------------------------------------- background jobs (I-32)
def test_a_big_file_is_processed_in_the_background_with_progress(file_app, make_user):
    with file_app.app_context():
        from app.services.user_service import UserService

        UserService.create_user(
            username="eng",
            email="eng@example.com",
            full_name="Eng",
            role_code="ENGINEER",
            password=PASSWORD,
            actor=None,
        )
    client = file_app.test_client()
    assert (
        client.post(
            "/api/v1/auth/login", json={"username": "eng", "password": PASSWORD}
        ).status_code
        == 200
    )
    rows = [_row(str(1000 + i), f"Background row number {i}") for i in range(25)]
    accepted = client.post(
        "/api/v1/imports",
        data={"file": (io.BytesIO(_csv(rows)), "library.csv"), "async": "true"},
        content_type="multipart/form-data",
    )
    assert accepted.status_code == 202
    job = accepted.get_json()
    assert job["status"] in ("PENDING", "VALIDATING", "COMPLETE")
    deadline = time.time() + 60
    while job["status"] not in ("COMPLETE", "FAILED") and time.time() < deadline:
        time.sleep(0.3)
        job = client.get(f"/api/v1/imports/{job['id']}").get_json()
    assert job["status"] == "COMPLETE" and job["created_count"] == 25
    assert job["processed_count"] == 25 and job["progress_percent"] == 100


def test_files_above_the_threshold_go_to_the_background_automatically(app, client, engineer):
    app.config["IMPORT_ASYNC_ROW_THRESHOLD"] = 1
    job = _post(client, [_row("1"), _row("2", "Second row title")]).get_json()
    assert job["status"] == "COMPLETE"  # (tests run the background job inline)
    assert job["created_count"] == 2 and job["processed_count"] == 2


def test_import_ui_pages_render_and_the_validate_button_creates_nothing(app, client, engineer):
    assert client.get("/imports/new").status_code == 200
    assert client.get("/imports/template.csv").status_code == 200
    assert client.get("/imports/template.xlsx").status_code == 200
    assert client.get("/imports/reconciliation").status_code == 200
    checked = client.post(
        "/imports/new",
        data={
            "file": (io.BytesIO(_csv([_row("1"), _row("2", set_delay="soon")])), "x.csv"),
            "action": "validate",
            "mode": "CREATE",
        },
        content_type="multipart/form-data",
    )
    html = checked.get_data(as_text=True)
    assert checked.status_code == 200 and "set_delay" in html and "Nothing has been created" in html
    uploaded = client.post(
        "/imports/new",
        data={
            "file": (io.BytesIO(_csv([_row("3")])), "x.csv"),
            "action": "import",
            "mode": "CREATE",
        },
        content_type="multipart/form-data",
    )
    assert uploaded.status_code == 302
    detail = client.get(uploaded.headers["Location"])
    assert detail.status_code == 200 and b"Import Job" in detail.data
