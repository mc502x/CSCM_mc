"""The technical nomenclature: a vocabulary of its own, apart from the interface translations
(docs/26-translations.md, docs/03-srs.md FR-166).

TC-061. Requirements: FR-166, BR-018, AUD-077, SEC-010.
"""

import io
import json

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.audit import AuditLogEntry
from app.models.terminology import TerminologyTerm, TerminologyTranslation
from app.services.terminology_service import TerminologyService as Terms
from app.terminology_seed import LANGUAGES, terms
from app.translation_catalogue import catalogue, reviewable_languages
from tests.helpers import login_as, make_cast


@pytest.fixture()
def cast(client, make_user):
    return make_cast(make_user)


def _term(app, code):
    with app.app_context():
        return db.session.query(TerminologyTerm).filter_by(code=code).one().id


def _tr(app, code, language):
    with app.app_context():
        return (
            db.session.query(TerminologyTranslation)
            .join(TerminologyTerm, TerminologyTerm.id == TerminologyTranslation.term_id)
            .filter(TerminologyTerm.code == code, TerminologyTranslation.language_code == language)
            .one()
        )


# ------------------------------------------------------------------------------ the shipped terms
def test_the_vocabulary_ships_with_every_term_in_every_language_as_a_draft(app):
    shipped = terms()
    with app.app_context():
        assert db.session.query(TerminologyTerm).count() == len(shipped) == 78
        rows = db.session.query(TerminologyTranslation).all()
        translatable = [t for t in shipped if not t["keep_english"]]
        assert len(rows) == len(translatable) * len(LANGUAGES)
        assert {r.status for r in rows} == {"DRAFT"} and {r.proposed_by for r in rows} == {None}
        assert set(LANGUAGES) == set(reviewable_languages())
        for coverage in Terms.coverage():
            assert coverage["approved"] == 0 and coverage["missing"] == 0
            assert coverage["draft"] == coverage["total"] == len(translatable)


def test_the_shipped_terms_agree_with_the_interface_catalogue_where_both_have_the_word():
    """Role names are interface strings and terms: they must be the same word."""
    shipped = {t["english"]: t["translations"] for t in terms() if t["category"] == "ROLE"}
    problems = []
    for language, texts in catalogue().items():
        for english, per_language in shipped.items():
            if english in texts and texts[english] != per_language[language]:
                problems.append((language, english, texts[english], per_language[language]))
    assert problems == []


def test_terms_that_are_never_translated_have_no_translation(app):
    with app.app_context():
        keep = db.session.query(TerminologyTerm).filter_by(keep_english=1).all()
        assert {t.english for t in keep} >= {"StCd", "MC / HC", "2XM / 3XM / 4XM", "PrgIdx_GUi"}
        assert (
            db.session.query(TerminologyTranslation)
            .filter(TerminologyTranslation.term_id.in_([t.id for t in keep]))
            .count()
            == 0
        )
        assert Terms.term_text("keep-stcd", "fr") == "StCd"
        assert Terms.term_text("status-code", "fr") == "Status Code"  # a draft is not used yet


# ------------------------------------------------------------------------------ who may do what
def test_everybody_reads_but_only_the_technical_roles_propose(client, cast):
    for who in ("viewer", "engineer", "reviewer", "chief_engineer", "stcd_owner", "admin"):
        login_as(client, cast[who])
        assert client.get("/terminology").status_code == 200, who
        assert client.get("/terminology/fr").status_code == 200, who
    login_as(client, cast["viewer"])
    page = client.get("/terminology/fr").get_data(as_text=True)
    assert "Save translations" not in page and "readonly" in page
    assert client.post("/terminology/fr/save", data={"term": "1", "t_1": "x"}).status_code == 403
    assert client.post(
        "/terminology/translations/1/approve", data={"language": "fr"}
    ).status_code in (
        200,
        302,
    )  # refused with a message, nothing changes
    assert "Save translations" in _as(client, cast, "engineer", "/terminology/fr")


def _as(client, cast, who, path):
    login_as(client, cast[who])
    return client.get(path).get_data(as_text=True)


def test_an_interface_reviewer_has_no_say_in_the_terminology(app, client, cast):
    login_as(client, cast["admin"])
    client.post("/translations/fr/assign", data={"user_id": cast["viewer"].id})
    login_as(client, cast["viewer"])
    term_id = _term(app, "alarm")
    assert (
        client.post(
            "/terminology/fr/save", data={"term": term_id, f"t_{term_id}": "Alerte"}
        ).status_code
        == 403
    )
    assert Terms.can_propose(type("A", (), {"role_code": "VIEWER"})()) is False


# ------------------------------------------------------------------------------ propose and approve
def test_a_proposal_is_a_draft_until_another_person_approves_it(app, client, cast):
    term_id = _term(app, "alarm")
    login_as(client, cast["reviewer"])
    done = client.post(
        "/terminology/fr/save",
        data={
            "term": term_id,
            f"t_{term_id}": "Signal d'alarme",
            f"n_{term_id}": "standard wording",
        },
        follow_redirects=True,
    )
    assert "1 translation(s) saved as drafts." in done.get_data(as_text=True)
    tr = _tr(app, "alarm", "fr")
    assert (
        tr.text == "Signal d'alarme"
        and tr.status == "DRAFT"
        and tr.proposed_by == cast["reviewer"].id
    )
    # the proposer cannot approve their own proposal
    own = client.post(
        f"/terminology/translations/{tr.id}/approve", data={"language": "fr"}, follow_redirects=True
    )
    assert "cannot approve your own proposal" in own.get_data(as_text=True)
    assert _tr(app, "alarm", "fr").status == "DRAFT"
    # a Chief Engineer can
    login_as(client, cast["chief_engineer"])
    ok = client.post(
        f"/terminology/translations/{tr.id}/approve", data={"language": "fr"}, follow_redirects=True
    )
    assert "Translation approved." in ok.get_data(as_text=True)
    approved = _tr(app, "alarm", "fr")
    assert approved.status == "APPROVED" and approved.approved_by == cast["chief_engineer"].id
    with app.app_context():
        assert Terms.term_text("alarm", "fr") == "Signal d'alarme"
        assert Terms.approved("fr")["alarm"] == "Signal d'alarme"
        actions = {a.action for a in db.session.query(AuditLogEntry).all()}
        assert {"TERM_TRANSLATE", "TERM_APPROVE"} <= actions
    # changing an approved text makes it a draft again, the same text again changes nothing
    login_as(client, cast["engineer"])
    client.post("/terminology/fr/save", data={"term": term_id, f"t_{term_id}": "Alarme"})
    assert _tr(app, "alarm", "fr").status == "DRAFT"
    stamp = _tr(app, "alarm", "fr").proposed_at
    client.post("/terminology/fr/save", data={"term": term_id, f"t_{term_id}": "Alarme"})
    assert _tr(app, "alarm", "fr").proposed_at == stamp  # the same text again changes nothing


def test_a_shipped_draft_can_be_approved_by_anybody_who_may_approve_and_all_at_once(
    app, client, cast
):
    login_as(client, cast["engineer"])
    tid = _term(app, "set-delay")
    assert client.post(
        f"/terminology/translations/{_tr(app, 'set-delay', 'de').id}/approve",
        data={"language": "de"},
    ).status_code in (200, 302)
    assert _tr(app, "set-delay", "de").status == "DRAFT"  # an Engineer proposes, does not approve
    login_as(client, cast["chief_engineer"])
    page = client.post("/terminology/de/approve-all", follow_redirects=True).get_data(as_text=True)
    with app.app_context():
        approved = (
            db.session.query(TerminologyTranslation)
            .filter_by(language_code="de", status="APPROVED")
            .count()
        )
        assert approved == 71 and f"{approved} translation(s) approved." in page
        assert (
            db.session.query(TerminologyTranslation)
            .filter_by(language_code="fr", status="APPROVED")
            .count()
            == 0
        )
    assert tid


def test_send_back_needs_a_note_and_reopens_an_approved_translation(app, client, cast):
    login_as(client, cast["chief_engineer"])
    client.post("/terminology/de/approve-all")
    tr = _tr(app, "brake-program", "de")
    term_id = tr.term_id
    login_as(client, cast["reviewer"])
    refused = client.post(
        f"/terminology/translations/{tr.id}/reject?term={term_id}",
        data={"language": "de", f"n_{term_id}": ""},
        follow_redirects=True,
    )
    assert "Say why in the note." in refused.get_data(as_text=True)
    assert _tr(app, "brake-program", "de").status == "APPROVED"
    client.post(
        f"/terminology/translations/{tr.id}/reject?term={term_id}",
        data={"language": "de", f"n_{term_id}": "we say Bremsenprogramm"},
    )
    back = _tr(app, "brake-program", "de")
    assert (
        back.status == "DRAFT"
        and back.note == "we say Bremsenprogramm"
        and back.approved_by is None
    )


def test_validation_of_a_translation(app, client, cast):
    login_as(client, cast["engineer"])
    keep = _term(app, "keep-stcd")
    with app.app_context():
        with pytest.raises(Exception, match="never translated"):
            Terms.propose(
                keep,
                "fr",
                "x",
                None,
                type("A", (), {"id": cast["engineer"].id, "role_code": "ENGINEER"})(),
            )
    term_id = _term(app, "alarm")
    long = client.post(
        "/terminology/fr/save",
        data={"term": term_id, f"t_{term_id}": "x" * 201},
        follow_redirects=True,
    )
    assert "at most 200" in long.get_data(as_text=True)
    assert client.get("/terminology/xx").status_code == 404
    assert (
        client.get("/terminology/en").status_code == 404
    )  # the source language has no translation


# ------------------------------------------------------------------------------ managing terms
def test_only_the_managers_add_and_change_terms_and_a_new_english_reopens_approvals(
    app, client, cast
):
    for who in ("engineer", "reviewer", "viewer"):
        login_as(client, cast[who])
        assert "Add a term" not in client.get("/terminology").get_data(as_text=True), who
        out = client.post(
            "/terminology/terms",
            data={"code": "x-term", "english": "X", "category": "PRODUCT"},
            follow_redirects=True,
        )
        assert "Only" in out.get_data(as_text=True), who
    login_as(client, cast["stcd_owner"])
    made = client.post(
        "/terminology/terms",
        data={
            "code": "pitch-system",
            "english": "Pitch system",
            "category": "PRODUCT",
            "definition": "Turns the rotor blades.",
        },
        follow_redirects=True,
    )
    assert "Term added." in made.get_data(as_text=True)
    dup = client.post(
        "/terminology/terms",
        data={"code": "pitch-system", "english": "Pitch", "category": "PRODUCT"},
        follow_redirects=True,
    )
    assert "exists already" in dup.get_data(as_text=True)
    bad = client.post(
        "/terminology/terms",
        data={"code": "Bad Code!", "english": "x", "category": "PRODUCT"},
        follow_redirects=True,
    )
    assert "lower-case letters" in bad.get_data(as_text=True)
    term_id = _term(app, "pitch-system")
    login_as(client, cast["engineer"])
    client.post("/terminology/de/save", data={"term": term_id, f"t_{term_id}": "Pitchsystem"})
    login_as(client, cast["reviewer"])
    client.post(
        f"/terminology/translations/{_tr(app, 'pitch-system', 'de').id}/approve",
        data={"language": "de"},
    )
    assert _tr(app, "pitch-system", "de").status == "APPROVED"
    login_as(client, cast["chief_engineer"])
    client.post(
        f"/terminology/terms/{term_id}",
        data={
            "english": "Blade pitch system",
            "category": "PRODUCT",
            "definition": "x",
            "is_active": "on",
        },
    )
    assert _tr(app, "pitch-system", "de").status == "DRAFT"  # the meaning changed: approve again
    with app.app_context():
        update = db.session.query(AuditLogEntry).filter_by(action="TERM_UPDATE").one()
        assert json.loads(update.after_value)["approvals_reopened"] == 1
    client.post(
        f"/terminology/terms/{term_id}",
        data={"english": "Blade pitch system", "category": "PRODUCT"},
    )  # off
    login_as(client, cast["viewer"])
    assert "Blade pitch system" not in client.get("/terminology").get_data(as_text=True)


# ------------------------------------------------------------------------------ check and export
def test_the_check_finds_interface_strings_that_do_not_use_the_term(app, client, cast):
    login_as(client, cast["engineer"])
    term_id = _term(app, "alarm")
    client.post("/terminology/de/save", data={"term": term_id, f"t_{term_id}": "Zzyzx"})
    login_as(client, cast["viewer"])
    page = client.get("/terminology/de/check").get_data(as_text=True)
    assert "Zzyzx" in page and "to look at" in page
    with app.app_context():
        findings = Terms.check("de")
        assert findings and all(f["term"] == "Alarm" for f in findings if f["expected"] == "Zzyzx")
    # a clean language has few or no findings for a term that the strings really use
    with app.app_context():
        assert all(
            f["expected"] != "Statuscode" or "Statuscode" not in f["translation"].casefold()
            for f in Terms.check("de")
        )


def test_the_vocabulary_is_exported_per_language_and_for_all(app, client, cast):
    login_as(client, cast["viewer"])
    one = client.get("/terminology/ar/export.xlsx")
    assert one.status_code == 200 and "spreadsheetml" in one.mimetype
    sheet = load_workbook(io.BytesIO(one.data))["Terminology"]
    header = [c.value for c in sheet[1]]
    assert header[:6] == ["code", "category", "english", "definition", "keep_english", "ar"]
    rows = {r[0].value: r for r in sheet.iter_rows(min_row=2)}
    assert rows["status-code"][5].value == "رمز الحالة" and rows["status-code"][6].value == "DRAFT"
    assert rows["status-code"][5].alignment.horizontal == "right"  # Arabic reads right to left
    csv = client.get("/terminology/export.csv").get_data(as_text=True)
    assert csv.startswith("﻿code,category,english") and ",de," in csv.splitlines()[0] + ","
    assert "status-code" in csv and csv.splitlines()[0].count(",") >= 40
    assert client.get("/terminology/xx/export.csv").status_code == 404


def test_a_formula_in_a_cell_is_neutralised_in_the_csv(app, client, cast):
    login_as(client, cast["admin"])
    client.post(
        "/terminology/terms",
        data={"code": "formula-test", "english": "=1+1", "category": "PRODUCT", "definition": "@x"},
    )
    csv = client.get("/terminology/export.csv").get_data(as_text=True)
    assert "'=1+1" in csv and ",=1+1," not in csv
