"""FR-032, FR-045, FR-046, FR-054, FR-064, BR-008, SEC-023, TC-023, TC-035:
catalogue exports (released set, working set, formats), revision diff, and
pre-publish Release deletion."""

import csv
import io

from openpyxl import load_workbook

from app.extensions import db
from app.models.status_code import StatusCodeRevision
from app.services.revision_diff_service import RevisionDiffService
from tests.helpers import (
    create_draft,
    drive_to_approved,
    login_as,
    make_cast,
    release_revisions,
    valid_payload,
)


def _csv_rows(response):
    return list(csv.DictReader(io.StringIO(response.get_data(as_text=True))))


def _released_code(app, client, cast, wcnv_ids, **overrides):
    revision, _ = drive_to_approved(app, client, cast, wcnv_ids, **overrides)
    release_revisions(client, cast, [revision["id"]], label=f"v-{revision['id']}")
    return revision


# ---- FR-045: all Released codes ------------------------------------------------------------


def test_released_export_contains_only_released_non_sandbox_codes(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    released = _released_code(app, client, cast, wcnv_ids)
    create_draft(app, client, cast, wcnv_ids, title="Still a Draft Code")  # not released
    login_as(client, cast["admin"])
    client.post("/api/v1/status-codes", json=valid_payload(*wcnv_ids, is_sandbox=True))

    login_as(client, cast["viewer"])  # Viewers may export Released data (BR-009)
    response = client.get("/api/v1/exports/released?format=csv")
    assert response.status_code == 200
    rows = _csv_rows(response)
    assert [r["title"] for r in rows] == [released["title"]]
    assert "distribution_notice" not in rows[0]
    assert all(r["lifecycle_status"] == "RELEASED" for r in rows)


def test_released_export_filters_and_formats(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    _released_code(app, client, cast, wcnv_ids)
    login_as(client, cast["viewer"])

    other_group = client.get("/api/v1/lookups/functional-system-groups").get_json()[-1]["id"]
    empty = client.get(f"/api/v1/exports/released?format=csv&functional_system_group={other_group}")
    assert _csv_rows(empty) == []

    for fmt, content_type in (
        ("json", "application/json"),
        ("xlsx", "spreadsheetml"),
        ("pdf", "application/pdf"),
    ):
        response = client.get(f"/api/v1/exports/released?format={fmt}")
        assert response.status_code == 200, fmt
        assert content_type in response.headers["Content-Type"]
    assert client.get("/api/v1/exports/released?format=docx").status_code == 422


def test_xlsx_export_has_same_rows_as_csv(app, client, make_user, wcnv_ids):
    """TC-035."""
    cast = make_cast(make_user)
    _released_code(app, client, cast, wcnv_ids)
    login_as(client, cast["viewer"])

    csv_rows = _csv_rows(client.get("/api/v1/exports/released?format=csv"))
    sheet = load_workbook(
        io.BytesIO(client.get("/api/v1/exports/released?format=xlsx").data)
    ).active
    header = [c.value for c in sheet[1]]
    xlsx_rows = [
        dict(zip(header, [c.value for c in row], strict=False))
        for row in sheet.iter_rows(min_row=2)
    ]

    assert header == list(csv_rows[0].keys())
    assert len(xlsx_rows) == len(csv_rows) == 1
    assert xlsx_rows[0]["status_code_identifier"] == csv_rows[0]["status_code_identifier"]
    assert xlsx_rows[0]["title"] == csv_rows[0]["title"]


def test_json_export_is_flat_rows(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    _released_code(app, client, cast, wcnv_ids)
    login_as(client, cast["viewer"])
    body = client.get("/api/v1/exports/released?format=json").get_json()
    assert isinstance(body, list) and body[0]["functional_system_group"] == "WCNV"
    assert not any(isinstance(v, dict | list) for v in body[0].values())


# ---- FR-046: working set ------------------------------------------------------------------


def test_working_set_is_labelled_and_excludes_sandbox(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    create_draft(app, client, cast, wcnv_ids, title="Unapproved Draft Code")
    login_as(client, cast["admin"])
    client.post("/api/v1/status-codes", json=valid_payload(*wcnv_ids, is_sandbox=True))

    login_as(client, cast["reviewer"])
    response = client.get("/api/v1/exports/working-set?format=csv")
    assert response.status_code == 200
    assert "NOT_FOR_DISTRIBUTION" in response.headers["Content-Disposition"]
    rows = _csv_rows(response)
    assert [r["title"] for r in rows] == ["Unapproved Draft Code"]  # TC-023: no sandbox
    assert rows[0]["distribution_notice"] == "Not for distribution - unapproved data"
    assert rows[0]["lifecycle_status"] == "DRAFT"


def test_working_set_filters_by_lifecycle_state(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    create_draft(app, client, cast, wcnv_ids, title="Unapproved Draft Code")
    login_as(client, cast["engineer"])
    none = client.get("/api/v1/exports/working-set?format=csv&lifecycle_status=REVIEW")
    assert _csv_rows(none) == []
    some = client.get("/api/v1/exports/working-set?format=csv&lifecycle_status=DRAFT&q=Unapproved")
    assert len(_csv_rows(some)) == 1


def test_working_set_not_available_to_viewers_or_anonymous(app, client, make_user):
    cast = make_cast(make_user)
    assert client.get("/api/v1/exports/working-set?format=csv").status_code == 401
    login_as(client, cast["viewer"])
    assert client.get("/api/v1/exports/working-set?format=csv").status_code == 403


def test_working_set_pdf_and_xlsx_render(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    create_draft(app, client, cast, wcnv_ids)
    login_as(client, cast["chief_engineer"])
    assert client.get("/api/v1/exports/working-set?format=pdf").status_code == 200
    xlsx = client.get("/api/v1/exports/working-set?format=xlsx")
    sheet = load_workbook(io.BytesIO(xlsx.data)).active
    assert "distribution_notice" in [c.value for c in sheet[1]]


# ---- SEC-023: formula injection ---------------------------------------------------------------


def test_exports_neutralise_formula_injection(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    create_draft(app, client, cast, wcnv_ids, title='=HYPERLINK("http://evil","x")')
    login_as(client, cast["engineer"])

    rows = _csv_rows(client.get("/api/v1/exports/working-set?format=csv"))
    assert rows[0]["title"].startswith("'=")
    sheet = load_workbook(
        io.BytesIO(client.get("/api/v1/exports/working-set?format=xlsx").data)
    ).active
    header = [c.value for c in sheet[1]]
    title_cell = sheet.cell(row=2, column=header.index("title") + 1).value
    assert title_cell.startswith("'=")


# ---- FR-032 / FR-064: revision diff --------------------------------------------------------


def test_revision_diff_endpoint_and_cr_baseline_diff(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    first = _released_code(app, client, cast, wcnv_ids)

    login_as(client, cast["engineer"])
    second = client.post(
        f"/api/v1/status-codes/{first['status_code_id']}/revisions",
        json=valid_payload(
            *wcnv_ids, title="Converter Grid Undervoltage Trip v2", set_delay="20ms"
        ),
    ).get_json()
    assert second["revision_number"] == 2

    diff = client.get(f"/api/v1/status-codes/{first['status_code_id']}/revisions/1/diff/2")
    assert diff.status_code == 200
    body = diff.get_json()
    assert (body["from_revision"], body["to_revision"]) == (1, 2)
    changes = {c["field"]: (c["old_value"], c["new_value"]) for c in body["field_changes"]}
    assert changes["title"] == (
        "Converter Grid Undervoltage Trip",
        "Converter Grid Undervoltage Trip v2",
    )
    assert changes["set_delay"] == ("10ms", "20ms")
    assert "revision_number" not in changes and "lifecycle_status" not in changes
    assert "description" not in changes  # unchanged fields are omitted

    with app.app_context():
        revision_row = db.session.get(StatusCodeRevision, second["id"])
        baseline = RevisionDiffService.diff_against_baseline(revision_row)
        assert baseline["is_new_code"] is False
        assert baseline["baseline_revision_number"] == 1
        assert {c["field"] for c in baseline["field_changes"]} == set(changes)


def test_new_code_diff_is_marked_new(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    revision, _ = create_draft(app, client, cast, wcnv_ids)
    with app.app_context():
        diff = RevisionDiffService.diff_against_baseline(
            db.session.get(StatusCodeRevision, revision["id"])
        )
    assert diff["is_new_code"] is True
    assert {c["field"] for c in diff["field_changes"]} >= {"title", "description"}


def test_diff_hidden_from_viewers_for_open_revisions_and_unknown_revisions(
    app, client, make_user, wcnv_ids
):
    cast = make_cast(make_user)
    first = _released_code(app, client, cast, wcnv_ids)
    login_as(client, cast["engineer"])
    client.post(
        f"/api/v1/status-codes/{first['status_code_id']}/revisions",
        json=valid_payload(*wcnv_ids, title="Converter Grid Undervoltage Trip v2"),
    )

    login_as(client, cast["viewer"])  # BR-009: revision 2 is a Draft
    url = f"/api/v1/status-codes/{first['status_code_id']}/revisions/1/diff/2"
    assert client.get(url).status_code == 404
    assert (
        client.get(f"/api/v1/status-codes/{first['status_code_id']}/revisions/1/diff/9").status_code
        == 404
    )


def test_diff_card_renders_on_change_request_page(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    _, cr_id = create_draft(app, client, cast, wcnv_ids, title="Diff Card Code")
    login_as(client, cast["engineer"])
    page = client.get(f"/change-requests/{cr_id}")
    assert page.status_code == 200
    assert b"(new code)" in page.data and b"Diff Card Code" in page.data


# ---- BR-008: pre-publish Release deletion ----------------------------------------------------


def test_release_can_be_deleted_before_publish_only(app, client, make_user, wcnv_ids):
    cast = make_cast(make_user)
    revision, _ = drive_to_approved(app, client, cast, wcnv_ids)

    login_as(client, cast["admin"])
    building = client.post("/api/v1/releases", json={"name": "B", "version_label": "b1"}).get_json()
    client.put(
        f"/api/v1/releases/{building['id']}/items",
        json={"status_code_revision_ids": [revision["id"]]},
    )
    client.get(f"/api/v1/releases/{building['id']}/export?format=csv")  # export_job references it

    login_as(client, cast["engineer"])
    assert client.delete(f"/api/v1/releases/{building['id']}").status_code == 403

    login_as(client, cast["admin"])
    assert client.delete(f"/api/v1/releases/{building['id']}").status_code == 204
    assert client.get(f"/api/v1/releases/{building['id']}").status_code == 404
    assert client.delete(f"/api/v1/releases/{building['id']}").status_code == 404

    # The Approved revision is untouched and still a candidate for a new Release.
    with app.app_context():
        assert db.session.get(StatusCodeRevision, revision["id"]).lifecycle_status == "APPROVED"

    published_id = release_revisions(client, cast, [revision["id"]], label="p1")
    login_as(client, cast["admin"])
    assert client.delete(f"/api/v1/releases/{published_id}").status_code == 409


def test_export_file_names_carry_a_timestamp(app, client, make_user, wcnv_ids):
    """Every export file name ends with the generation time, YYYY-MM-DD_HH-MM (UTC)."""
    import re

    cast = make_cast(make_user)
    revision = _released_code(app, client, cast, wcnv_ids)
    stamp = r"\d{4}-\d{2}-\d{2}_\d{2}-\d{2}"

    def name(response):
        return re.search(r"filename=(\S+)", response.headers["Content-Disposition"]).group(1)

    login_as(client, cast["admin"])
    released = name(client.get("/api/v1/exports/released?format=csv"))
    assert re.fullmatch(rf"cscm_released_{stamp}\.csv", released), released
    working = name(client.get("/api/v1/exports/working-set?format=xlsx"))
    assert re.fullmatch(rf"cscm_working-set_NOT_FOR_DISTRIBUTION_{stamp}\.xlsx", working), working
    backup = name(client.post("/api/v1/exports/full-database"))
    assert re.fullmatch(rf"cscm_full-database-backup_{stamp}\.db", backup), backup
    release_id = client.get("/api/v1/releases").get_json()[0]["id"]
    per_release = name(client.get(f"/api/v1/releases/{release_id}/export?format=pdf"))
    assert re.fullmatch(rf"cscm_release_v-{revision['id']}_{stamp}\.pdf", per_release), per_release
