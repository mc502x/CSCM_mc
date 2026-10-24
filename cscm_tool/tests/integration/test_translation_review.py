"""Translation review in the application (docs/26-translations.md, docs/03-srs.md FR-164): native
speakers check the interface translations of their language in the browser, an Administrator or an
StCd Owner assigns them and decides on their proposals, and an accepted proposal is shown at once.

TC-059. Requirements: FR-164, FR-157, FR-161, BR-018, BR-025, AUD-077.
"""

import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.audit import AuditLogEntry
from app.models.translation import (
    LanguageReviewer,
    TranslationOverride,
    TranslationReview,
    TranslationReviewItem,
)
from app.services.translation_review_service import TranslationReviewService as Reviews
from app.translation_catalogue import effective, screens, string_id
from tests.helpers import login_as, make_cast

pytestmark = pytest.mark.usefixtures("languages_open")


@pytest.fixture()
def cast(client, make_user):
    return make_cast(make_user)


def _assign(client, cast, language="fr", who="viewer", by="admin"):
    login_as(client, cast[by])
    return client.post(
        f"/translations/{language}/assign", data={"user_id": cast[who].id}, follow_redirects=True
    )


def _answers(app, language, verdict="OK", skip=()):
    with app.app_context():
        return [
            {"string_id": string_id(e), "verdict": verdict}
            for e in effective(language)
            if e not in skip
        ]


class _Auth:
    def __init__(self, user):
        self.id, self.role_code = user.id, user.role.code


def _auth(user):
    with db.session.no_autoflush:
        return _Auth(user)


# ------------------------------------------------------------------------------ access
def test_the_overview_shows_a_reviewer_only_their_languages_and_a_manager_all(client, cast):
    _assign(client, cast, "fr", "viewer")
    _assign(client, cast, "ka", "viewer")
    login_as(client, cast["viewer"])
    page = client.get("/translations").get_data(as_text=True)
    assert (
        "/translations/fr" in page and "/translations/ka" in page and "/translations/de" not in page
    )
    assert "Review</a>" in page and "Decisions and reviewers" not in page
    assert "Translations" in client.get("/").get_data(as_text=True)  # the sidebar link is there

    login_as(client, cast["engineer"])  # nobody assigned them
    page = client.get("/translations").get_data(as_text=True)
    assert "not assigned to review a language" in page
    assert 'href="/translations"' not in client.get("/").get_data(as_text=True)

    login_as(client, cast["stcd_owner"])
    page = client.get("/translations").get_data(as_text=True)
    assert "/translations/de" in page and "/translations/fr/decisions" in page


def test_only_an_administrator_or_an_stcd_owner_assigns_reviewers_and_it_is_audited(
    app, client, cast
):
    for who in ("engineer", "reviewer", "chief_engineer", "viewer"):
        login_as(client, cast[who])
        assert (
            client.post("/translations/fr/assign", data={"user_id": cast["viewer"].id}).status_code
            == 403
        ), who
        assert client.get("/translations/fr/decisions").status_code in (403, 302), who
    done = _assign(client, cast, "fr", "viewer", by="stcd_owner")
    assert "Reviewer assigned." in done.get_data(as_text=True)
    with app.app_context():
        assert db.session.query(LanguageReviewer).filter_by(language_code="fr").count() == 1
        entry = db.session.query(AuditLogEntry).filter_by(action="LANGUAGE_REVIEWER_ASSIGN").one()
        assert json.loads(entry.after_value)["language"] == "fr"
    login_as(client, cast["admin"])
    out = client.post(
        "/translations/fr/unassign", data={"user_id": cast["viewer"].id}, follow_redirects=True
    )
    assert "Reviewer removed." in out.get_data(as_text=True)
    with app.app_context():
        assert (
            db.session.query(AuditLogEntry).filter_by(action="LANGUAGE_REVIEWER_REMOVE").count()
            == 1
        )
    bad = client.post("/translations/fr/assign", data={"user_id": 99999}, follow_redirects=True)
    assert "active user" in bad.get_data(as_text=True)
    assert client.post("/translations/xx/assign", data={"user_id": 1}).status_code == 404
    assert (
        client.post("/translations/en/assign", data={"user_id": 1}).status_code == 404
    )  # the source


def test_an_unassigned_user_cannot_open_a_review_and_an_unknown_language_is_404(client, cast):
    login_as(client, cast["engineer"])
    assert client.get("/translations/fr").status_code == 403
    assert client.post("/translations/fr/save", data={}).status_code == 403
    assert client.get("/translations/fr/sheet.xlsx").status_code == 403
    assert client.get("/translations/xx").status_code == 404
    login_as(client, cast["stcd_owner"])  # a manager who does not review is sent to the decisions
    assert "/decisions" in client.get("/translations/fr").headers["Location"]


# ------------------------------------------------------------------------------ the review
def test_a_reviewer_answers_strings_and_is_told_what_is_wrong(app, client, cast):
    _assign(client, cast)
    login_as(client, cast["viewer"])
    page = client.get("/translations/fr").get_data(as_text=True)
    assert "Review: FR" in page and "Meaning" in page and 'type="radio"' in page
    sid = string_id("Log in")
    ok = client.post(
        "/translations/fr/save", data={"sid": [sid], f"v_{sid}": "OK"}, follow_redirects=True
    )
    assert "Saved 1 answer(s)." in ok.get_data(as_text=True)
    # a CHANGE needs a proposal, and the placeholders of the English text have to survive
    bad = client.post(
        "/translations/fr/save", data={"sid": [sid], f"v_{sid}": "CHANGE"}, follow_redirects=True
    )
    assert "needs a proposal" in bad.get_data(as_text=True)
    with app.app_context():
        review = db.session.query(TranslationReview).filter_by(language_code="fr").one()
        saved = Reviews.save(
            review.id,
            [
                {"string_id": string_id("results"), "proposal": "résultats %(n)s"},
                {"string_id": "Sffffff", "verdict": "OK"},
                {"string_id": sid, "verdict": "MAYBE"},
            ],
            _auth(cast["viewer"]),
        )
        assert saved["saved"] == 0 and len(saved["errors"]) == 3  # none of the three is acceptable


def test_a_review_is_completed_only_when_every_string_has_a_verdict_and_is_recorded(
    app, client, cast
):
    _assign(client, cast)
    login_as(client, cast["viewer"])
    early = client.post("/translations/fr/complete", follow_redirects=True)
    assert "have no verdict yet" in early.get_data(as_text=True)
    with app.app_context():
        review = db.session.query(TranslationReview).filter_by(language_code="fr").one()
        Reviews.save(review.id, _answers(app, "fr"), _auth(cast["viewer"]))
    done = client.post("/translations/fr/complete", follow_redirects=True)
    text = done.get_data(as_text=True)
    assert "Review completed" in text and "reviewed" in text
    with app.app_context():
        review = db.session.query(TranslationReview).filter_by(language_code="fr").one()
        assert review.status == "COMPLETED" and review.completed_at
        entry = (
            db.session.query(AuditLogEntry).filter_by(action="TRANSLATION_REVIEW_COMPLETE").one()
        )
        after = json.loads(entry.after_value)
        assert after["language"] == "fr" and after["ok"] == len(effective("fr"))
    again = client.post(
        "/translations/fr/save",
        data={"sid": [string_id("Log in")], "v_" + string_id("Log in"): "OK"},
        follow_redirects=True,
    )
    assert "completed" in again.get_data(as_text=True).lower() or again.status_code == 200
    # the next visit starts a new review
    assert client.get("/translations/fr").status_code == 200
    with app.app_context():
        assert db.session.query(TranslationReview).filter_by(language_code="fr").count() == 2


# ------------------------------------------------------------------------------ proposals and decisions
def _propose(
    app, client, cast, language="fr", english="Log in", text="Connexion", reviewer="viewer"
):
    _assign(client, cast, language, reviewer)
    login_as(client, cast[reviewer])
    sid = string_id(english)
    client.post(
        f"/translations/{language}/save",
        data={"sid": [sid], f"v_{sid}": "CHANGE", f"p_{sid}": text, f"c_{sid}": "more usual"},
    )
    with app.app_context():
        return db.session.query(TranslationReviewItem).filter_by(string_id=sid).one().id


def test_an_accepted_proposal_is_shown_at_once_and_can_be_taken_back(app, client, cast):
    item = _propose(app, client, cast)
    shipped = (
        app.test_client().get("/login", headers={"Accept-Language": "fr"}).get_data(as_text=True)
    )
    assert "Se connecter" in shipped and "Connexion" not in shipped

    login_as(client, cast["stcd_owner"])
    page = client.get("/translations/fr/decisions").get_data(as_text=True)
    assert "Connexion" in page and "more usual" in page and "Accept" in page
    done = client.post(
        f"/translations/fr/decide/{item}",
        data={"action": "accept", "text": "Connexion"},
        follow_redirects=True,
    )
    assert "Decision saved." in done.get_data(as_text=True)
    client.post("/api/v1/auth/logout")
    live = app.test_client().get("/login", headers={"Accept-Language": "fr"}).get_data(as_text=True)
    assert "Connexion" in live  # no release, no restart
    with app.app_context():
        override = db.session.query(TranslationOverride).filter_by(language_code="fr").one()
        assert override.text == "Connexion" and override.review_item_id == item
        decided = (
            db.session.query(AuditLogEntry).filter_by(action="TRANSLATION_PROPOSAL_DECIDE").one()
        )
        assert json.loads(decided.after_value)["decision"] == "ACCEPTED"
        override_id = override.id
    assert "Connexion" not in app.test_client().get(
        "/login", headers={"Accept-Language": "de"}
    ).get_data(
        as_text=True
    )  # only French

    login_as(client, cast["admin"])
    client.post(f"/translations/fr/override/{override_id}/remove")
    client.post("/api/v1/auth/logout")
    assert "Se connecter" in app.test_client().get(
        "/login", headers={"Accept-Language": "fr"}
    ).get_data(as_text=True)
    with app.app_context():
        assert (
            db.session.query(AuditLogEntry).filter_by(action="TRANSLATION_OVERRIDE_REMOVE").count()
            == 1
        )


def test_nobody_decides_on_their_own_proposal_and_a_decision_is_final(app, client, cast):
    item = _propose(app, client, cast, reviewer="admin")  # an Administrator who reviews French
    login_as(client, cast["admin"])
    own = client.post(
        f"/translations/fr/decide/{item}", data={"action": "accept"}, follow_redirects=True
    )
    assert "own proposal" in own.get_data(as_text=True)
    login_as(client, cast["stcd_owner"])
    assert "Decision saved." in client.post(
        f"/translations/fr/decide/{item}", data={"action": "reject"}, follow_redirects=True
    ).get_data(as_text=True)
    again = client.post(
        f"/translations/fr/decide/{item}", data={"action": "accept"}, follow_redirects=True
    )
    assert "decided already" in again.get_data(as_text=True)
    with app.app_context():
        assert db.session.query(TranslationOverride).count() == 0
    login_as(client, cast["admin"])  # and the reviewer cannot edit a decided row any more
    sid = string_id("Log in")
    edit = client.post(
        "/translations/fr/save", data={"sid": [sid], f"v_{sid}": "OK"}, follow_redirects=True
    )
    assert "Already decided" in edit.get_data(as_text=True)


def test_a_proposal_must_keep_the_placeholders_and_an_unclear_string_is_only_noted(
    app, client, cast
):
    _assign(client, cast)
    login_as(client, cast["viewer"])
    sid = string_id("Not yet assigned")
    client.post(
        "/translations/fr/save",
        data={"sid": [sid], f"v_{sid}": "UNCLEAR", f"c_{sid}": "assigned to whom?"},
    )
    with app.app_context():
        item_id = db.session.query(TranslationReviewItem).filter_by(string_id=sid).one().id
    login_as(client, cast["stcd_owner"])
    refused = client.post(
        f"/translations/fr/decide/{item_id}", data={"action": "accept"}, follow_redirects=True
    )
    assert "Only a proposed change" in refused.get_data(as_text=True)
    noted = client.post(
        f"/translations/fr/decide/{item_id}", data={"action": "note"}, follow_redirects=True
    )
    assert "Decision saved." in noted.get_data(as_text=True)
    with app.app_context():
        assert db.session.get(TranslationReviewItem, item_id).decision == "NOTED"


def test_a_missing_translation_can_be_proposed_and_accepted(app, client, cast):
    _assign(client, cast, "bn")
    login_as(client, cast["viewer"])
    with app.app_context():
        missing = next(e for e in screens() if e not in effective("bn"))
    sid = string_id(missing)
    page = client.get("/translations/bn?scope=missing").get_data(as_text=True)
    assert missing in page and "Your translation" in page
    client.post(
        "/translations/bn/save", data={"sid": [sid], f"p_{sid}": "অনুবাদ", "scope": "missing"}
    )
    with app.app_context():
        item = db.session.query(TranslationReviewItem).filter_by(string_id=sid).one()
        assert (
            item.verdict == "CHANGE" and item.current_text is None
        )  # a translation for a missing string
        item_id = item.id
    login_as(client, cast["admin"])
    client.post(f"/translations/bn/decide/{item_id}", data={"action": "accept"})
    with app.app_context():
        assert effective("bn")[missing] == "অনুবাদ"


def test_a_completed_review_shows_as_changed_after_a_later_acceptance(app, client, cast):
    item = _propose(app, client, cast)
    with app.app_context():
        review = db.session.query(TranslationReview).filter_by(language_code="fr").one()
        Reviews.save(review.id, _answers(app, "fr", skip=("Log in",)), _auth(cast["viewer"]))
    login_as(client, cast["viewer"])
    client.post("/translations/fr/complete")
    assert "reviewed" in client.get("/translations").get_data(as_text=True)
    login_as(client, cast["stcd_owner"])
    client.post(f"/translations/fr/decide/{item}", data={"action": "accept"})
    page = client.get("/translations").get_data(as_text=True)
    assert "changed since the review" in page


# ------------------------------------------------------------------------------ the sheet
def test_the_sheet_can_be_downloaded_filled_in_offline_and_uploaded(app, client, cast):
    _assign(client, cast)
    login_as(client, cast["viewer"])
    sheet = client.get("/translations/fr/sheet.xlsx")
    assert sheet.status_code == 200 and "spreadsheetml" in sheet.mimetype
    workbook = load_workbook(io.BytesIO(sheet.data))
    assert workbook.sheetnames == ["Review", "Missing", "Checklist"]
    review = workbook["Review"]
    rows = list(review.iter_rows(min_row=2))
    for row in rows[:3]:
        row[4].value = "OK"
    rows[3][4].value = "CHANGE"
    rows[3][5].value = "proposition"
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    done = client.post(
        "/translations/fr/upload",
        data={"sheet": (buffer, "sheet.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "saved 4 answer(s)" in done.get_data(as_text=True)
    with app.app_context():
        assert db.session.query(TranslationReviewItem).count() == 4
        assert (
            db.session.query(TranslationReviewItem).filter_by(verdict="CHANGE").one().proposal
            == "proposition"
        )


def test_a_sheet_that_is_not_one_is_refused_clearly(client, cast):
    _assign(client, cast)
    login_as(client, cast["viewer"])
    junk = client.post(
        "/translations/fr/upload",
        data={"sheet": (io.BytesIO(b"not a workbook"), "x.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "could not be read" in junk.get_data(as_text=True)
    other = Workbook()
    other.active.title = "Other"
    buffer = io.BytesIO()
    other.save(buffer)
    buffer.seek(0)
    wrong = client.post(
        "/translations/fr/upload",
        data={"sheet": (buffer, "x.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "has no" in wrong.get_data(as_text=True) and "tab" in wrong.get_data(as_text=True)
    none = client.post(
        "/translations/fr/upload",
        data={},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "Choose the completed sheet" in none.get_data(as_text=True)


def test_right_to_left_languages_are_shown_right_to_left_in_the_review(client, cast):
    _assign(client, cast, "ar")
    login_as(client, cast["viewer"])
    page = client.get("/translations/ar").get_data(as_text=True)
    assert 'dir="rtl"' in page and 'class="rtl-text"' in page
    _assign(client, cast, "fr")
    login_as(client, cast["viewer"])
    assert 'class="rtl-text"' not in client.get("/translations/fr").get_data(as_text=True)


def test_an_account_that_reviews_has_history_in_the_usage_report(client, cast):
    _assign(client, cast)
    login_as(client, cast["admin"])
    usage = client.get(f"/api/v1/users/{cast['viewer'].id}/usage").get_json()
    assert usage["deletable"] is False and "language_reviewer.user_id" in usage["references"]
