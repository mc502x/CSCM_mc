"""Status code import from the generated StCd_<n>.st programs and the legacy -> new mapping
(docs/03-srs.md section 1.16). Six real programs - three Main Controller (1000-1002) and three Hub
Controller (14000, 14005, 14025) - are in tests/fixtures/st/.

TC-053. Requirements: FR-048o, FR-048p, FR-048q, FR-048r, FR-048s, FR-048t, FR-048u, FR-048v,
FR-048w, BR-023, SEC-031.
"""

import io
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.importers import stcd
from app.models.change_request import ChangeRequest, ChangeSet
from app.models.import_job import ImportJob, ImportSourceItem
from app.models.status_code import StatusCode
from app.services.numbering_service import NumberingService
from tests.helpers import cr_id_for_revision, domain_sign, login_as, make_cast, submit

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "st"
MC_FILES = ["StCd_1000.st", "StCd_1001.st", "StCd_1002.st"]
HC_FILES = ["StCd_14000.st", "StCd_14005.st", "StCd_14025.st"]
WORKBOOK = Path(__file__).resolve().parents[3] / "docs" / "artifacts" / "hc-numbering-proposal.xlsx"


def _read(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def _upload(names, **fields):
    data = {"files": [(io.BytesIO(_read(n)), n) for n in names]}
    data.update(fields)
    return data


def _preview(client, names=MC_FILES + HC_FILES, **fields):
    return client.post(
        "/api/v1/imports/st/preview",
        data=_upload(names, **fields),
        content_type="multipart/form-data",
    )


def _confirm(client, token, selected=None, **extra):
    body = {"token": token, **extra}
    if selected is not None:
        body["selected"] = selected
    return client.post("/api/v1/imports/st", json=body)


@pytest.fixture()
def engineer(client, make_user):
    user = make_user("ENGINEER")
    login_as(client, user)
    return user


def _by_program(preview):
    return {(i["controller"], i["program_number"]): i for i in preview["items"]}


# ------------------------------------------------------------------------------ the reader
def test_the_reader_takes_the_attributes_of_a_main_controller_program():
    (code,), _sha = stcd.parse_bytes(_read("StCd_1000.st"), "StCd_1000.st")
    assert code.identifier == "StCd-01000" and code.controller == "MC"
    assert code.legacy_number == 3000 and code.name == "Umrichter nicht bereit"
    assert code.prg_idx == 14066
    assert (code.set_delay, code.set_delay_ms) == ("T#10ms", 10)
    assert code.reset_delay_ms == 120_000
    assert (code.repeat_code, code.max_repeats, code.repeat_time_s) == (9085, 4, 86400)
    assert code.availability_group == 20 and code.alarm is False and code.snapshot is True
    assert code.depends_on == [1045]
    assert code.programs["AutoReset"] == 10 and code.programs["Brake"] == 40
    assert len(code.user_groups) == 15


def test_a_hub_program_keeps_its_legacy_number_and_gets_the_new_one_from_the_table():
    mapping = stcd.load_bundled_mapping()
    for proposal, expected in ((1, "StCd-12501"), (2, None)):
        (code,), _ = stcd.parse_bytes(_read("StCd_14000.st"), "StCd_14000.st", hc=True)
        stcd.apply_mapping(code, mapping, proposal)
        assert code.legacy_number == 14000 and code.controller == "HC"
        assert code.identifier == (expected or f"StCd-{mapping[14000][1]:05d}")
    # a variable instead of a time literal is kept as text, with no milliseconds
    assert code.set_delay == "StCd14000SetDelayGpTi" and code.set_delay_ms is None


def test_a_legacy_hub_number_missing_from_the_table_leaves_no_identifier():
    (code,), _ = stcd.parse_bytes(_read("StCd_14000.st"), "x.st", hc=True)
    stcd.apply_mapping(code, {}, 1)
    assert code.identifier == "" and "not in the mapping table" in code.warnings[0]


@pytest.mark.skipif(not WORKBOOK.exists(), reason="numbering workbook not in this checkout")
def test_the_bundled_hub_table_equals_the_workbook():
    assert stcd.load_bundled_mapping() == stcd.load_mapping(WORKBOOK)


def test_a_file_without_programs_or_too_large_is_refused_cleanly():
    assert stcd.parse_text("just text", "x.st") == []
    with pytest.raises(stcd.StParseError):
        stcd.parse_bytes(b"x" * (stcd.MAX_FILE_BYTES + 1), "big.st")


# ------------------------------------------------------------------------------ preview
def test_the_preview_lists_every_program_and_creates_nothing(app, client, engineer):
    response = _preview(client)
    assert response.status_code == 200, response.get_json()
    preview = response.get_json()
    summary = preview["summary"]
    assert summary["programs"] == 6 and summary["files"] == 6
    assert summary["ready"] == 6 and summary["blocked"] == 0 and summary["exists"] == 0
    assert summary["with_legacy"] == 6 and summary["without_legacy"] == 0
    items = _by_program(preview)
    mc = items[("MC", 1000)]
    assert mc["identifier"] == "StCd-01000" and mc["legacy_number"] == 3000
    assert mc["group"] and mc["subgroup"] and mc["prg_idx"] == 14066
    hc = items[("HC", 14000)]
    assert hc["identifier"] == "StCd-12501" and hc["legacy_number"] == 14000
    assert any("parameter StCd14000SetDelayGpTi" in w for w in hc["warnings"])
    with app.app_context():
        assert db.session.query(StatusCode).count() == 0
        assert db.session.query(ImportJob).count() == 0


def test_proposal_two_places_the_hub_programs_inside_the_main_scheme(client, engineer):
    items = _by_program(_preview(client, HC_FILES, hc_proposal="2").get_json())
    mapping = stcd.load_bundled_mapping()
    assert items[("HC", 14000)]["identifier"] == f"StCd-{mapping[14000][1]:05d}"


def test_the_controller_can_be_forced_for_a_file(client, engineer):
    forced = _preview(client, ["StCd_1000.st"], controller="HC").get_json()
    item = forced["items"][0]
    assert item["controller"] == "HC" and item["state"] == "blocked"
    assert "mapping table" in item["issues"][0]["message"]


def test_a_zip_and_loose_files_are_read_alike_and_other_files_are_noted(client, engineer):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name in MC_FILES:
            archive.writestr(f"export/mc/{name}", _read(name))
        archive.writestr("export/readme.docx", b"not a program")
    response = client.post(
        "/api/v1/imports/st/preview",
        data={"files": [(io.BytesIO(buffer.getvalue()), "all.zip")]},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    assert response.get_json()["summary"]["programs"] == 3


def test_a_number_defined_twice_in_one_upload_blocks_the_second(client, engineer):
    data = {
        "files": [
            (io.BytesIO(_read("StCd_1000.st")), "a/StCd_1000.st"),
            (io.BytesIO(_read("StCd_1000.st")), "b/StCd_1000.st"),
        ]
    }
    preview = client.post(
        "/api/v1/imports/st/preview", data=data, content_type="multipart/form-data"
    ).get_json()
    states = [i["state"] for i in preview["items"]]
    assert states == ["ready", "blocked"]
    assert "also defined" in preview["items"][1]["issues"][0]["message"]


def test_nothing_usable_is_a_clear_error(client, engineer):
    response = client.post(
        "/api/v1/imports/st/preview",
        data={"files": [(io.BytesIO(b"hello"), "notes.st")]},
        content_type="multipart/form-data",
    )
    assert response.status_code == 422 and "status code program" in response.get_json()["message"]
    assert client.post("/api/v1/imports/st/preview").status_code == 422


def test_only_engineers_and_administrators_import(client, make_user):
    login_as(client, make_user("REVIEWER"))
    assert _preview(client).status_code == 403
    assert client.get("/api/v1/legacy-mapping").status_code == 200  # internal roles may read it


# ------------------------------------------------------------------------------ confirm
def test_confirm_creates_drafts_that_reserve_their_numbers_in_one_change_set(app, client, engineer):
    preview = _preview(client).get_json()
    items = _by_program(preview)
    chosen = [items[("MC", 1000)]["key"], items[("MC", 1002)]["key"], items[("HC", 14005)]["key"]]
    response = _confirm(client, preview["token"], chosen, change_set_name="Batch A")
    assert response.status_code == 202, response.get_json()
    job = response.get_json()
    assert job["source_format"] == "ST" and job["status"] == "COMPLETE"
    assert job["created_count"] == 3 and job["error_count"] == 0 and job["change_set_id"]

    with app.app_context():
        codes = db.session.query(StatusCode).order_by(StatusCode.id).all()
        assert [c.reserved_identifier for c in codes] == ["StCd-01000", "StCd-01002", "StCd-12502"]
        assert all(c.status_code_identifier is None for c in codes)  # issued only at Submit
        assert [c.legacy_reference_number for c in codes] == ["3000", "3010", "14005"]
        crs = db.session.query(ChangeRequest).all()
        assert {cr.cr_type for cr in crs} == {"IMPORT"}
        change_set = db.session.get(ChangeSet, job["change_set_id"])
        assert change_set.name == "Batch A" and len(change_set.change_requests) == 3
        assert db.session.query(ImportSourceItem).count() == 3
        item = db.session.query(ImportSourceItem).filter_by(program_number=1000).one()
        assert len(item.source_sha256) == 64 and item.controller == "MC"
        parsed = json.loads(item.parsed_json)
        assert parsed["prg_idx"] == 14066 and len(parsed["user_groups"]) == 15
        from sqlalchemy import text

        assert (
            db.session.execute(
                text("SELECT prg_idx FROM status_code WHERE reserved_identifier = 'StCd-01000'")
            ).scalar()
            == 14066
        )
        assert (Path(job_dir(app, job["id"])) / "StCd_1000.st").is_file()


def job_dir(app, job_id):
    return Path(app.instance_path) / "imports" / str(job_id)


def test_the_draft_carries_the_attributes_of_the_program(app, client, engineer):
    preview = _preview(client, ["StCd_1000.st"], default_category="WARNING").get_json()
    _confirm(client, preview["token"])
    with app.app_context():
        revision = db.session.query(StatusCode).one().revisions[0]
        assert revision.title == "Umrichter nicht bereit"
        assert revision.status_category == "WARNING" and revision.available_group == 20
        assert (revision.set_delay, revision.reset_delay) == ("10ms", "2min")
        assert revision.number_of_times_repeated == 4 and revision.repeated_over_the_course == "1d"
        assert revision.status_code_for_repeated_error == 9085
        assert revision.trigger_snapshot == 1 and revision.loadless_spinning_permitted == 1
        assert revision.auto_reset_program == 10 and revision.brake_program == 40
        assert revision.related_status_codes == "StCd-01045"
        assert "StCd_1000.st" in revision.description and revision.lifecycle_status == "DRAFT"
        assert sorted(p.turbine_platform.code for p in revision.platforms)  # all platforms


def test_confirming_without_a_selection_takes_every_ready_program(client, engineer):
    preview = _preview(client, MC_FILES).get_json()
    assert _confirm(client, preview["token"]).get_json()["created_count"] == 3


def test_a_second_preview_shows_what_is_already_there_and_confirm_refuses_it(client, engineer):
    first = _preview(client, MC_FILES).get_json()
    _confirm(client, first["token"])
    second = _preview(client, MC_FILES + ["StCd_14000.st"]).get_json()
    assert second["summary"]["exists"] == 3 and second["summary"]["ready"] == 1
    existing = [i for i in second["items"] if i["state"] == "exists"][0]
    assert existing["existing"]["identifier"] == existing["identifier"]
    refused = _confirm(client, second["token"], [existing["key"]])
    assert refused.status_code == 422 and "cannot be imported" in refused.get_json()["message"]
    assert _confirm(client, second["token"]).get_json()["created_count"] == 1  # only the new one


def test_a_legacy_number_that_belongs_to_another_code_blocks_the_program(client, engineer):
    first = _preview(client, ["StCd_1000.st"]).get_json()
    _confirm(client, first["token"])
    # the same legacy number arrives under a different new number
    text = _read("StCd_1000.st").decode("cp1252").replace("StCd_1000", "StCd_1050")
    text = text.replace("StCdArUi[Idx] := 1000", "StCdArUi[Idx] := 1050")
    again = client.post(
        "/api/v1/imports/st/preview",
        data={"files": [(io.BytesIO(text.encode("cp1252")), "StCd_1050.st")]},
        content_type="multipart/form-data",
    ).get_json()
    item = again["items"][0]
    assert item["identifier"] == "StCd-01050" and item["state"] == "blocked"
    assert "already belongs to" in item["issues"][0]["message"]


def test_a_used_prg_idx_is_a_warning_and_is_not_stored_twice(app, client, engineer):
    first = _preview(client, ["StCd_1000.st"]).get_json()
    _confirm(client, first["token"])
    text = _read("StCd_1000.st").decode("cp1252")
    text = text.replace("StCd_1000", "StCd_1050").replace(
        "StCdArUi[Idx] := 1000", "StCdArUi[Idx] := 1050"
    )
    text = text.replace("StCd: 3000,", "StCd: 3999,")
    again = client.post(
        "/api/v1/imports/st/preview",
        data={"files": [(io.BytesIO(text.encode("cp1252")), "StCd_1050.st")]},
        content_type="multipart/form-data",
    ).get_json()
    item = again["items"][0]
    assert item["state"] == "ready" and any("PrgIdx_GUi 14066" in w for w in item["warnings"])
    assert _confirm(client, again["token"]).get_json()["created_count"] == 1


def test_a_preview_belongs_to_its_uploader_and_expires(app, client, make_user):
    one = make_user("ENGINEER")
    login_as(client, one)
    token = _preview(client, MC_FILES).get_json()["token"]
    login_as(client, make_user("ENGINEER"))
    assert _confirm(client, token).status_code == 403
    login_as(client, one)
    assert _confirm(client, "0" * 32).status_code == 404
    assert _confirm(client, "../../etc").status_code == 404
    assert client.post("/api/v1/imports/st", json={}).status_code == 422
    assert _confirm(client, token, ["9:9"]).status_code == 422  # not part of this preview


# ------------------------------------------------------------------------------ governance
def test_submitting_issues_exactly_the_reserved_number(app, client, make_user):
    cast = make_cast(make_user)
    login_as(client, cast["engineer"])
    preview = _preview(client, ["StCd_1001.st"]).get_json()
    assert _confirm(client, preview["token"]).status_code == 202
    with app.app_context():
        code = db.session.query(StatusCode).one()
        code_id, revision_id = code.id, code.revisions[0].id
    cr_id = cr_id_for_revision(app, revision_id)
    domain_sign(client, cast, cr_id)
    response = submit(client, cast, cr_id)
    assert response.status_code == 200, response.get_json()
    with app.app_context():
        assert db.session.get(StatusCode, code_id).status_code_identifier == "StCd-01001"
        from app.models.audit import AuditLogEntry

        entry = (
            db.session.query(AuditLogEntry)
            .filter_by(action="SUBMIT_FOR_REVIEW")
            .order_by(AuditLogEntry.id.desc())
            .first()
        )
        assert json.loads(entry.after_value)["identifier_source"] == "reserved"


def test_an_administrator_cannot_override_a_reserved_identifier(app, client, make_user):
    cast = make_cast(make_user)
    login_as(client, cast["admin"])
    preview = _preview(client, ["StCd_1001.st"]).get_json()
    assert _confirm(client, preview["token"]).status_code == 202
    with app.app_context():
        cr_id = db.session.query(ChangeRequest).one().id
    domain_sign(client, cast, cr_id)
    response = submit(
        client, cast, cr_id, who="admin", identifier_override="StCd-01099", override_reason="x"
    )
    assert response.status_code == 422 and "reserved" in response.get_json()["message"]


def test_allocation_skips_a_reserved_number_and_a_manual_override_cannot_take_it(
    app, client, engineer, wcnv_ids
):
    _confirm(client, _preview(client, ["StCd_1000.st"]).get_json()["token"])
    group_id, subgroup_id = wcnv_ids
    with app.app_context():
        # StCd-01000 is the first number of the Converter band (1000-1199) and is reserved
        assert NumberingService.allocate_next_identifier(group_id, subgroup_id) == "StCd-01001"
        with pytest.raises(ValueError, match="reserved by an imported"):
            NumberingService.validate_manual_identifier("StCd-01000", group_id, subgroup_id)
        # the code itself may be issued its own number
        code = db.session.query(StatusCode).one()
        assert (
            NumberingService.validate_reserved_identifier(
                "StCd-01000", group_id, subgroup_id, code.id
            )
            == "StCd-01000"
        )


def test_the_database_refuses_a_second_reservation_of_a_number(app, client, engineer):
    preview = _preview(client, ["StCd_1000.st"]).get_json()
    _confirm(client, preview["token"])
    from sqlalchemy import text
    from sqlalchemy.exc import IntegrityError

    with app.app_context():
        code = db.session.query(StatusCode).one()
        with pytest.raises(IntegrityError):
            db.session.execute(
                text(
                    "INSERT INTO status_code (functional_system_group_id, functional_subgroup_id, "
                    "owner_id, reserved_identifier) VALUES (:g, :s, :o, 'StCd-01000')"
                ),
                {
                    "g": code.functional_system_group_id,
                    "s": code.functional_subgroup_id,
                    "o": code.owner_id,
                },
            )
        db.session.rollback()


# ------------------------------------------------------------------------------ legacy mapping
def test_the_mapping_lists_legacy_to_new_and_counts_codes_without_a_legacy_number(
    app, client, make_user, wcnv_ids
):
    login_as(client, make_user("ENGINEER"))
    preview = _preview(client).get_json()
    assert _confirm(client, preview["token"]).get_json()["created_count"] == 6
    body = client.get("/api/v1/legacy-mapping").get_json()
    pairs = {(e["controller"], e["legacy_number"]): e["new_identifier"] for e in body["entries"]}
    assert pairs[("MC", "3000")] == "StCd-01000" and pairs[("MC", "3010")] == "StCd-01002"
    assert pairs[("HC", "14000")] == "StCd-12501" and pairs[("HC", "14025")] == "StCd-12503"
    assert all(e["identifier_state"] == "reserved" for e in body["entries"])
    assert body["summary"]["with_legacy_number"] == 6 and body["summary"]["new_without_legacy"] == 0
    assert body["summary"]["by_controller"] == {"MC": 3, "HC": 3}

    # a code created by hand has no legacy number: it is counted, not mapped ("n/a")
    from tests.helpers import valid_payload

    created = client.post(
        "/api/v1/status-codes",
        json=valid_payload(*wcnv_ids, legacy_reference_number=None),
    )
    assert created.status_code == 201, created.get_json()
    assert client.get("/api/v1/legacy-mapping").get_json()["summary"]["new_without_legacy"] == 1

    hub_only = client.get("/api/v1/legacy-mapping?controller=HC").get_json()["entries"]
    assert {e["controller"] for e in hub_only} == {"HC"} and len(hub_only) == 3
    assert [
        e["legacy_number"]
        for e in client.get("/api/v1/legacy-mapping?q=01002").get_json()["entries"]
    ] == ["3010"]
    assert client.get("/api/v1/legacy-mapping?controller=XX").status_code == 422
    assert client.get("/api/v1/legacy-mapping?state=nope").status_code == 422


def test_the_mapping_exports_as_csv_and_xlsx(client, make_user):
    login_as(client, make_user("ENGINEER"))
    _confirm(client, _preview(client, MC_FILES).get_json()["token"])
    csv_body = client.get("/api/v1/legacy-mapping/export").get_data(as_text=True)
    assert (
        csv_body.lstrip("﻿").splitlines()[0].startswith("controller,legacy_number,new_identifier")
    )
    assert "3000,StCd-01000,reserved" in csv_body
    workbook = load_workbook(
        io.BytesIO(client.get("/api/v1/legacy-mapping/export?format=xlsx").data)
    )
    assert workbook.sheetnames == ["Legacy to new", "Note"]
    assert [c.value for c in workbook["Legacy to new"][2]][:3] == ["MC", "3000", "StCd-01000"]
    assert client.get("/api/v1/legacy-mapping/export?format=pdf").status_code == 422


def test_the_mapping_is_not_for_viewers(client, make_user):
    login_as(client, make_user("VIEWER"))
    assert client.get("/api/v1/legacy-mapping").status_code == 403
    assert client.get("/api/v1/legacy-mapping/export").status_code == 403


def test_the_source_of_an_imported_code_can_be_read(app, client, engineer):
    _confirm(client, _preview(client, ["StCd_1000.st"]).get_json()["token"])
    with app.app_context():
        code_id = db.session.query(StatusCode).one().id
    source = client.get(f"/api/v1/status-codes/{code_id}/import-source").get_json()
    assert source["source_file"] == "StCd_1000.st" and source["program_number"] == 1000
    assert source["attributes"]["legacy_number"] == 3000
    assert client.get("/api/v1/status-codes/99999/import-source").status_code == 404


def test_an_expired_preview_is_refused_and_removed(app, client, engineer):
    token = _preview(client, MC_FILES).get_json()["token"]
    staged = Path(app.instance_path) / "imports" / "staging" / token
    manifest = staged / "manifest.json"
    data = json.loads(manifest.read_text(encoding="utf-8"))
    data["created"] -= 25 * 3600
    manifest.write_text(json.dumps(data), encoding="utf-8")
    response = _confirm(client, token)
    assert response.status_code == 404 and "expired" in response.get_json()["message"]
    assert not staged.exists()


def test_a_zip_that_unpacks_too_large_is_refused(client, engineer, monkeypatch):
    from app.services import st_import_service

    monkeypatch.setattr(st_import_service, "MAX_UNZIPPED_BYTES", 10_000)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name in MC_FILES:
            archive.writestr(name, _read(name))
    response = client.post(
        "/api/v1/imports/st/preview",
        data={"files": [(io.BytesIO(buffer.getvalue()), "big.zip")]},
        content_type="multipart/form-data",
    )
    assert response.status_code == 422 and "unpacks to more than" in response.get_json()["message"]
    garbage = client.post(
        "/api/v1/imports/st/preview",
        data={"files": [(io.BytesIO(b"not a zip"), "x.zip")]},
        content_type="multipart/form-data",
    )
    assert garbage.status_code == 422 and "ZIP" in garbage.get_json()["message"]
