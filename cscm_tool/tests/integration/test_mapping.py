"""The mapping of legacy to new status codes kept in the database and changed only with the three
approvals of the StCd Owner, the Chief Engineer and an Administrator (FR-191..FR-198), and what the
.st import takes from it.

TC-076..TC-079. Requirements: FR-191, FR-192, FR-193, FR-194, FR-195, FR-196, FR-197, FR-198, BR-027.
"""

import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.mapping import MappingChange, StatusCodeMapping
from app.models.status_code import StatusCode
from app.services.mapping_service import COLUMNS, MappingService
from tests.helpers import login_as, make_cast

SEED = Path(__file__).resolve().parents[2] / "seed_data" / "legacy_mapping.xlsx"
FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "st"


@pytest.fixture()
def cast(client, make_user):
    return make_cast(make_user)


def _row_fields(**overrides):
    fields = {
        "family": "MC",
        "legacy_number": 3100,
        "new_number": 1150,
        "description": "Converter test code",
        "functional_group": "Converter",
        "logical_node": "WCNV",
        "status_type": "ERROR",
        "up_down_counter": "No",
    }
    fields.update(overrides)
    return fields


def _propose_add(client, user, **overrides):
    login_as(client, user)
    return client.post(
        "/api/v1/mapping-changes",
        json={
            "title": "New code",
            "reason": "The code is needed for the new converter.",
            "items": [{"action": "ADD", "fields": _row_fields(**overrides)}],
        },
    )


def _decide(client, user, change_id, decision="APPROVE", comment=None):
    login_as(client, user)
    return client.post(
        f"/api/v1/mapping-changes/{change_id}/decisions",
        json={"decision": decision, "comment": comment},
    )


def _row(new_number):
    return db.session.query(StatusCodeMapping).filter_by(new_number=new_number).first()


# ------------------------------------------------------------------------------ the seed
def test_the_seed_workbook_loads_into_an_empty_table_once(app):
    with app.app_context():
        loaded = MappingService.seed(SEED)
        assert loaded > 0
        total = MappingService.count()
        assert (
            total == loaded and MappingService.seed(SEED) == 0 and MappingService.count() == total
        )
        summary = MappingService.summary()
        assert summary["by_family"]["MC"] + summary["by_family"]["HC"] == total


def test_the_seed_workbook_is_consistent(app):
    with app.app_context():
        MappingService.seed(SEED)
        rows = db.session.query(StatusCodeMapping).all()
        assert len({r.new_number for r in rows}) == len(rows)
        legacy = [(r.family, r.legacy_number) for r in rows if r.legacy_number is not None]
        assert len(set(legacy)) == len(legacy)
        assert all(r.description and r.functional_group and r.logical_node for r in rows)


def test_a_bad_seed_workbook_is_refused_as_a_whole(app, tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "MC"
    sheet.append(list(COLUMNS))
    sheet.append([1, 1000, "Fine", "Converter", "WCNV", "Error", "No"])
    sheet.append([2, 1000, "Twice", "Converter", "WCNV", "Error", "No"])
    sheet.append([3, 1001, "Odd type", "Converter", "WCNV", "Fatal", "No"])
    book.create_sheet("HC").append(list(COLUMNS))
    path = tmp_path / "bad.xlsx"
    book.save(path)
    with app.app_context():
        with pytest.raises(Exception, match="type must be Error, Warning or Info"):
            MappingService.seed(path)
        assert MappingService.count() == 0  # nothing half loaded


# ------------------------------------------------------------------------------ the three approvals
def test_a_new_row_takes_effect_only_after_all_three_approvals(app, client, cast):
    created = _propose_add(client, cast["engineer"])
    assert created.status_code == 201, created.get_json()
    change_id = created.get_json()["id"]
    with app.app_context():
        assert _row(1150) is None  # a request changes nothing

    first = _decide(client, cast["stcd_owner"], change_id)
    assert first.status_code == 200 and first.get_json()["state"] == "OPEN"
    second = _decide(client, cast["chief_engineer"], change_id)
    assert second.get_json()["state"] == "OPEN"
    with app.app_context():
        assert _row(1150) is None  # two of three are not enough
    third = _decide(client, cast["admin"], change_id)
    assert third.status_code == 200 and third.get_json()["state"] == "APPLIED"
    with app.app_context():
        row = _row(1150)
        assert row.status == "APPROVED" and row.status_type == "ERROR" and row.up_down_counter == 0
        assert row.source == f"mapping change {change_id}" and row.version == 1


def test_one_rejection_closes_the_request_and_needs_a_reason(app, client, cast):
    change_id = _propose_add(client, cast["engineer"]).get_json()["id"]
    assert _decide(client, cast["admin"], change_id, "REJECT").status_code == 422  # no reason given
    rejected = _decide(client, cast["admin"], change_id, "REJECT", "Not needed")
    assert rejected.get_json()["state"] == "REJECTED"
    late = _decide(client, cast["stcd_owner"], change_id)
    assert late.status_code == 409  # closed
    with app.app_context():
        assert _row(1150) is None


def test_only_the_three_roles_decide_and_each_only_once(client, cast):
    change_id = _propose_add(client, cast["engineer"]).get_json()["id"]
    for who in ("engineer", "reviewer", "viewer"):
        assert _decide(client, cast[who], change_id).status_code == 403
    assert _decide(client, cast["stcd_owner"], change_id).status_code == 200
    again = _decide(client, cast["stcd_owner"], change_id)
    assert again.status_code == 409 and "decided already" in again.get_json()["message"]


def test_a_second_administrator_cannot_replace_the_first_and_the_requester_cannot_decide(
    client, cast, make_user
):
    change_id = _propose_add(client, cast["engineer"]).get_json()["id"]
    assert _decide(client, cast["admin"], change_id).status_code == 200
    other_admin = make_user("ADMINISTRATOR")
    assert _decide(client, other_admin, change_id).status_code == 409  # one Administrator's say
    # an Administrator who made the request cannot approve it
    own = _propose_add(client, cast["admin"], new_number=1151, legacy_number=3101).get_json()["id"]
    refused = _decide(client, cast["admin"], own)
    assert refused.status_code == 403 and "another person" in refused.get_json()["message"]


def test_a_decision_asks_for_the_password_again(client, cast):
    change_id = _propose_add(client, cast["engineer"]).get_json()["id"]
    login_as(client, cast["stcd_owner"])
    response = client.post(
        f"/api/v1/mapping-changes/{change_id}/decisions",
        json={"decision": "APPROVE"},
        no_step_up=True,
    )
    assert response.status_code == 403 and response.get_json()["error"] == "step_up_required"
    wrong = client.post(
        f"/api/v1/mapping-changes/{change_id}/decisions",
        json={"decision": "APPROVE", "current_password": "not the password at all"},
        no_step_up=True,
    )
    assert wrong.get_json()["error"] == "step_up_failed"


def test_the_requester_can_withdraw_and_nobody_else(app, client, cast):
    change_id = _propose_add(client, cast["engineer"]).get_json()["id"]
    login_as(client, cast["reviewer"])
    assert client.post(f"/api/v1/mapping-changes/{change_id}/withdraw").status_code == 403
    login_as(client, cast["engineer"])
    assert client.post(f"/api/v1/mapping-changes/{change_id}/withdraw").get_json()["state"] == (
        "WITHDRAWN"
    )
    assert _decide(client, cast["admin"], change_id).status_code == 409


def test_a_viewer_neither_reads_nor_proposes(client, cast):
    login_as(client, cast["viewer"])
    assert client.get("/api/v1/mapping").status_code == 403
    assert _propose_add(client, cast["viewer"]).status_code == 403


# ------------------------------------------------------------------------------ the checks
@pytest.mark.parametrize(
    ("overrides", "message"),
    [
        ({"description": ""}, "description"),
        ({"status_type": "Fatal"}, "type must be Error, Warning or Info"),
        ({"up_down_counter": "maybe"}, "up/down counter must be Yes or No"),
        ({"new_number": "abc"}, "new number must be a whole number"),
        ({"new_number": 100000}, "between 1 and 99999"),
        ({"parameter_set_delay": "soon"}, "neither a time"),
        ({"development_access": "Nobody"}, "not a named level"),
        ({"new_number": 12501}, "belongs to the"),  # a Hub number in a Main row
        ({"status_type": None}, "cannot be approved without type"),
        ({"logical_node": None}, "without logical node"),
    ],
)
def test_a_row_that_does_not_pass_the_checks_makes_no_request(client, cast, overrides, message):
    response = _propose_add(client, cast["engineer"], **overrides)
    assert response.status_code == 422, response.get_json()
    assert message in response.get_json()["message"]


def test_numbers_stay_unique(app, client, cast, st_mapping):
    duplicate = _propose_add(client, cast["engineer"], new_number=1000)
    assert duplicate.status_code == 422 and "two rows" in duplicate.get_json()["message"]
    legacy = _propose_add(client, cast["engineer"], legacy_number=3000)
    assert legacy.status_code == 422 and "legacy MC number 3000" in legacy.get_json()["message"]


def test_a_change_confirm_and_reject_work_on_stored_rows(app, client, cast, st_mapping):
    with app.app_context():
        row = _row(1000)
        row.status = "PROPOSED"
        db.session.commit()
        row_id = row.id
    login_as(client, cast["engineer"])
    confirm = client.post(
        "/api/v1/mapping-changes",
        json={
            "title": "Approve 1000",
            "reason": "Checked",
            "items": [{"action": "CONFIRM", "mapping_id": row_id}],
        },
    )
    assert confirm.status_code == 201, confirm.get_json()
    for who in ("stcd_owner", "chief_engineer", "admin"):
        _decide(client, cast[who], confirm.get_json()["id"])
    with app.app_context():
        assert _row(1000).status == "APPROVED"

    login_as(client, cast["engineer"])
    change = client.post(
        "/api/v1/mapping-changes",
        json={
            "title": "New title",
            "reason": "Better wording",
            "items": [
                {
                    "action": "CHANGE",
                    "mapping_id": row_id,
                    "fields": {"description": "Converter is not ready"},
                }
            ],
        },
    ).get_json()
    assert change["items"][0]["changed_fields"] == ["description"]
    for who in ("stcd_owner", "chief_engineer", "admin"):
        _decide(client, cast[who], change["id"])
    with app.app_context():
        assert _row(1000).description == "Converter is not ready"

    login_as(client, cast["engineer"])
    nothing = client.post(
        "/api/v1/mapping-changes",
        json={"title": "x", "reason": "y", "items": [{"action": "CHANGE", "mapping_id": row_id}]},
    )
    assert nothing.status_code == 422 and "nothing is changed" in nothing.get_json()["message"]


def test_a_row_changed_in_between_makes_the_older_request_unusable(app, client, cast, st_mapping):
    with app.app_context():
        row_id = _row(1001).id
    login_as(client, cast["engineer"])

    def change(description):
        response = client.post(
            "/api/v1/mapping-changes",
            json={
                "title": "t",
                "reason": "r",
                "items": [
                    {
                        "action": "CHANGE",
                        "mapping_id": row_id,
                        "fields": {"description": description},
                    }
                ],
            },
        )
        return response.get_json()["id"]

    older, newer = change("Fuse A tripped"), change("Fuse B tripped")
    for who in ("stcd_owner", "chief_engineer", "admin"):
        _decide(client, cast[who], newer)
    for who in ("stcd_owner", "chief_engineer"):
        assert _decide(client, cast[who], older).status_code == 200
    last = _decide(client, cast["admin"], older)
    assert last.status_code == 422 and "changed after the request" in last.get_json()["message"]
    with app.app_context():
        assert _row(1001).description == "Fuse B tripped"
        assert db.session.get(MappingChange, older).state == "OPEN"


def test_a_number_a_code_of_the_catalogue_carries_cannot_change(app, client, cast, st_mapping):
    login_as(client, cast["engineer"])
    preview = client.post(
        "/api/v1/imports/st/preview",
        data={"files": [(io.BytesIO((FIXTURES / "StCd_1000.st").read_bytes()), "StCd_1000.st")]},
        content_type="multipart/form-data",
    ).get_json()
    assert client.post("/api/v1/imports/st", json={"token": preview["token"]}).status_code == 202
    with app.app_context():
        assert _row(1000).status == "ISSUED"  # the import marks the number
        row_id = _row(1000).id
        assert db.session.query(StatusCode).count() == 1
    renumber = client.post(
        "/api/v1/mapping-changes",
        json={
            "title": "t",
            "reason": "r",
            "items": [{"action": "CHANGE", "mapping_id": row_id, "fields": {"new_number": 1199}}],
        },
    )
    assert renumber.status_code == 422 and "numbers cannot change" in renumber.get_json()["message"]
    reject = client.post(
        "/api/v1/mapping-changes",
        json={"title": "t", "reason": "r", "items": [{"action": "REJECT", "mapping_id": row_id}]},
    )
    assert reject.status_code == 422 and "cannot be rejected" in reject.get_json()["message"]


# ------------------------------------------------------------------------------ the workbook
def test_the_export_is_the_layout_the_upload_reads(app, client, cast, st_mapping):
    login_as(client, cast["engineer"])
    export = client.get("/api/v1/mapping/export")
    assert export.status_code == 200
    book = load_workbook(io.BytesIO(export.data))
    assert book.sheetnames == ["MC", "HC"]
    assert [c.value for c in book["MC"][1]] == list(COLUMNS)
    with app.app_context():
        plan = MappingService.plan_from_workbook(export.data)
        assert plan["items"] == [] and plan["unchanged"] == 6 and plan["problems"] == []


def test_a_workbook_makes_one_request_with_the_new_and_the_different_rows(
    app, client, cast, st_mapping
):
    login_as(client, cast["engineer"])
    book = load_workbook(io.BytesIO(client.get("/api/v1/mapping/export").data))
    sheet = book["MC"]
    header = [c.value for c in sheet[1]]
    sheet.cell(row=2, column=header.index("description") + 1).value = "Converter is not ready"
    new = [None] * len(header)
    for key, value in {
        "legacy_number": 3100,
        "new_number": 1150,
        "description": "Converter test code",
        "functional_group": "Converter",
        "logical_node": "WCNV",
        "type": "Info",
        "up_down_counter": "Yes",
    }.items():
        new[header.index(key)] = value
    sheet.append(new)
    buffer = io.BytesIO()
    book.save(buffer)
    response = client.post(
        "/api/v1/mapping-changes/workbook",
        data={
            "file": (io.BytesIO(buffer.getvalue()), "m.xlsx"),
            "title": "Corrected",
            "reason": "Review",
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 201, response.get_json()
    body = response.get_json()
    assert body["counts"] == {"ADD": 1, "CHANGE": 1, "CONFIRM": 0, "REJECT": 0}
    assert body["unchanged"] == 5 and body["change"]["origin"] == "WORKBOOK"
    for who in ("stcd_owner", "chief_engineer", "admin"):
        _decide(client, cast[who], body["change"]["id"])
    with app.app_context():
        assert _row(1150).status_type == "INFO" and _row(1150).up_down_counter == 1
        assert _row(1000).description == "Converter is not ready"


def test_a_workbook_with_a_wrong_row_names_it_and_requests_nothing(app, client, cast):
    login_as(client, cast["engineer"])
    book = Workbook()
    mc = book.active
    mc.title = "MC"
    mc.append(list(COLUMNS))
    mc.append([3100, 1150, "", "Converter", "WCNV", "Fatal"])
    book.create_sheet("HC").append(list(COLUMNS))
    buffer = io.BytesIO()
    book.save(buffer)
    response = client.post(
        "/api/v1/mapping-changes/workbook",
        data={"file": (io.BytesIO(buffer.getvalue()), "m.xlsx"), "title": "t", "reason": "r"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 422
    assert any("MC row 2" in p for p in response.get_json()["problems"])
    garbage = client.post(
        "/api/v1/mapping-changes/workbook",
        data={"file": (io.BytesIO(b"not a workbook"), "m.xlsx"), "title": "t", "reason": "r"},
        content_type="multipart/form-data",
    )
    assert garbage.status_code == 422 and "readable .xlsx" in garbage.get_json()["message"]
    with app.app_context():
        assert db.session.query(MappingChange).count() == 0


def test_the_export_never_starts_a_formula(app, client, cast, st_mapping):
    with app.app_context():
        _row(1000).description = "=1+1"
        db.session.commit()
    login_as(client, cast["engineer"])
    book = load_workbook(io.BytesIO(client.get("/api/v1/mapping/export").data))
    assert book["MC"].cell(row=2, column=3).value == "'=1+1"


# ------------------------------------------------------------------------------ the pages
def test_the_pages_show_the_mapping_the_requests_and_the_three_slots(app, client, cast, st_mapping):
    change_id = _propose_add(client, cast["engineer"]).get_json()["id"]
    login_as(client, cast["stcd_owner"])
    assert "StCd-01000" in client.get("/mapping").get_data(as_text=True)
    page = client.get(f"/mapping-changes/{change_id}").get_data(as_text=True)
    assert "StCd Owner" in page and "Chief Engineer" in page and "Administrator" in page
    assert 'name="decision" value="APPROVE"' in page  # the owner can decide
    login_as(client, cast["engineer"])
    own = client.get(f"/mapping-changes/{change_id}").get_data(as_text=True)
    assert 'name="decision"' not in own and "Withdraw" in own
    listing = client.get("/mapping-changes").get_data(as_text=True)
    assert "New code" in listing
    row = client.get("/mapping/1").get_data(as_text=True)
    assert "Request this change" in row


def test_the_pages_make_requests_and_decide_with_the_form_posts(app, client, cast, st_mapping):
    login_as(client, cast["engineer"])
    with app.app_context():
        row_id = _row(1002).id
    made = client.post(
        f"/mapping/{row_id}",
        data={"action": "CHANGE", "description": "Converter comm. fault", "reason": "Wording"},
    )
    assert made.status_code == 302
    change_id = int(made.headers["Location"].rsplit("/", 1)[1])
    ticked = client.post("/mapping/select", data={"ids": [row_id], "reason": "Checked"})
    assert ticked.status_code == 302
    for who in ("stcd_owner", "chief_engineer", "admin"):
        login_as(client, cast[who])
        done = client.post(f"/mapping-changes/{change_id}/decisions", data={"decision": "APPROVE"})
        assert done.status_code == 302
    with app.app_context():
        assert _row(1002).description == "Converter comm. fault"
    login_as(client, cast["admin"])
    assert (
        client.post("/mapping/seed").status_code == 302
    )  # the table is not empty: nothing happens


def test_the_menu_counts_the_requests_that_wait_for_a_decision(app, client, cast):
    from types import SimpleNamespace

    change_id = _propose_add(client, cast["engineer"]).get_json()["id"]

    def waiting(user):
        with app.app_context():
            return MappingService.awaiting(SimpleNamespace(id=user.id, role_code=user.role.code))

    assert waiting(cast["stcd_owner"]) == 1 and waiting(cast["chief_engineer"]) == 1
    assert waiting(cast["engineer"]) == 0  # not a decider
    _decide(client, cast["stcd_owner"], change_id)
    assert waiting(cast["stcd_owner"]) == 0 and waiting(cast["admin"]) == 1
    login_as(client, cast["chief_engineer"])
    assert "Mapping changes" in client.get("/mapping").get_data(as_text=True)


def test_a_workbook_that_unpacks_to_too_much_or_has_too_many_rows_is_refused(
    app, client, cast, monkeypatch
):
    from app.services import mapping_service

    login_as(client, cast["engineer"])
    book = Workbook()
    mc = book.active
    mc.title = "MC"
    mc.append(list(COLUMNS))
    for _n in range(1, 30):
        mc.append([None] * 30)  # blank rows count towards the limit of a sheet
    book.create_sheet("HC").append(list(COLUMNS))
    buffer = io.BytesIO()
    book.save(buffer)

    def upload():
        return client.post(
            "/api/v1/mapping-changes/workbook",
            data={"file": (io.BytesIO(buffer.getvalue()), "m.xlsx"), "title": "t", "reason": "r"},
            content_type="multipart/form-data",
        )

    monkeypatch.setattr(mapping_service, "MAX_SHEET_ROWS", 10)
    assert "has more than 10 rows" in upload().get_json()["message"]
    monkeypatch.setattr(mapping_service, "MAX_SHEET_ROWS", 20_000)
    monkeypatch.setattr(mapping_service, "MAX_WORKBOOK_UNPACKED", 100)
    assert "unpacks to more than" in upload().get_json()["message"]


def test_a_number_held_by_the_catalogue_is_found_in_the_identifiers(app, make_user, st_mapping):
    from app.models.lookup import LookupFunctionalSubgroup, LookupFunctionalSystemGroup
    from app.models.status_code import StatusCode
    from app.services.mapping_service import numbers_in_catalogue

    owner_id = make_user("ENGINEER").id
    with app.app_context():
        assert numbers_in_catalogue() == set()
        group = db.session.query(LookupFunctionalSystemGroup).first()
        subgroup = db.session.query(LookupFunctionalSubgroup).first()
        db.session.add(
            StatusCode(
                functional_system_group_id=group.id,
                functional_subgroup_id=subgroup.id,
                owner_id=owner_id,
                created_at="2026-10-27T10:00:00.000Z",
                reserved_identifier="StCd-01000",
            )
        )
        db.session.commit()
        assert numbers_in_catalogue() == {1000}


def test_a_withdrawn_import_is_blocked_when_its_mapping_row_is_gone(app, client, cast, st_mapping):
    login_as(client, cast["engineer"])
    data = {"files": [(io.BytesIO((FIXTURES / "StCd_1001.st").read_bytes()), "StCd_1001.st")]}
    preview = client.post(
        "/api/v1/imports/st/preview", data=data, content_type="multipart/form-data"
    ).get_json()
    assert client.post("/api/v1/imports/st", json={"token": preview["token"]}).status_code == 202
    from app.models.change_request import ChangeRequest

    with app.app_context():
        cr_id = db.session.query(ChangeRequest).one().id
    assert (
        client.post("/api/v1/change-requests/withdraw", json={"cr_ids": [cr_id]}).status_code == 200
    )
    with app.app_context():
        db.session.query(StatusCodeMapping).filter_by(new_number=1001).delete()
        db.session.commit()
    again = client.post(
        "/api/v1/imports/st/preview",
        data={"files": [(io.BytesIO((FIXTURES / "StCd_1001.st").read_bytes()), "StCd_1001.st")]},
        content_type="multipart/form-data",
    ).get_json()
    item = again["items"][0]
    assert item["state"] == "blocked" and "is not in the mapping" in item["issues"][0]["message"]


def test_show_all_puts_every_row_on_one_page_so_that_select_all_covers_them(
    app, client, cast, st_mapping
):
    from app.models.mapping import StatusCodeMapping
    from app.utils import utcnow_iso

    with app.app_context():
        now = utcnow_iso()
        for n in range(2001, 2121):  # 120 more rows: 126 in all, more than a page of 100
            db.session.add(
                StatusCodeMapping(
                    family="MC",
                    new_number=n,
                    description=f"Row {n}",
                    status="APPROVED",
                    created_at=now,
                    updated_at=now,
                )
            )
        db.session.commit()
    login_as(client, cast["engineer"])
    paged = client.get("/mapping").get_data(as_text=True)
    assert paged.count('class="map-row-check"') == 100
    assert "Show all (126)" in paged and "/mapping?all=1" in paged
    everything = client.get("/mapping?all=1").get_data(as_text=True)
    assert everything.count('class="map-row-check"') == 126
    assert "Show all (" not in everything and "?all=1" not in everything.split("<tbody>")[1]
    filtered = client.get("/mapping?all=1&family=HC").get_data(as_text=True)
    assert filtered.count('class="map-row-check"') == 3  # the filter still applies


def test_a_database_without_the_mapping_tables_says_what_to_run(app, client, cast):
    from sqlalchemy import text

    login_as(client, cast["engineer"])
    with app.app_context():
        for table in (
            "mapping_approval",
            "mapping_change_item",
            "mapping_change",
            "status_code_mapping",
        ):
            db.session.execute(text(f"DROP TABLE {table}"))  # nosec B608
        db.session.commit()
    page = client.get("/mapping")
    assert page.status_code == 503
    assert "flask db upgrade" in page.get_data(as_text=True)
    api = client.get("/api/v1/mapping")
    assert api.status_code == 503 and api.get_json()["error"] == "database_not_upgraded"
    assert client.get("/").status_code == 200  # the rest of the tool keeps working


def test_the_sign_off_of_a_domain_is_below_the_changes_and_the_status_above(
    app, client, cast, wcnv_ids
):
    from tests.helpers import create_draft

    revision, cr_id = create_draft(app, client, cast, wcnv_ids)
    page = client.get(f"/change-requests/{cr_id}").get_data(as_text=True)
    assert page.index("Sign-Off Status") < page.index("Changes") < page.index("Actions")
    if "Sign Off This Domain" in page:
        assert page.index("Sign Off This Domain") > page.index("Actions")
