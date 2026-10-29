"""The corpus report (I-68): what would the import do with every program under a folder, and why not.

TC-083. Requirements: FR-201, FR-048o, FR-192, FR-193.
"""

import io
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.auth.providers import AuthenticatedUser
from app.extensions import db
from app.models.mapping import StatusCodeMapping
from app.models.status_code import StatusCode
from app.models.user import User
from app.services.corpus_service import CorpusService, normalise

pytestmark = pytest.mark.usefixtures("st_mapping")
FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "st"


@pytest.fixture()
def corpus(tmp_path):
    for sub, names in (
        ("mc", ["StCd_1000.st", "StCd_1001.st", "StCd_1002.st"]),
        ("hc/old", ["StCd_14000.st", "StCd_14005.st", "StCd_14025.st"]),
    ):
        (tmp_path / sub).mkdir(parents=True)
        for name in names:
            shutil.copy(FIXTURES / name, tmp_path / sub / name)
    (tmp_path / "notes.st").write_text("this is not a program", encoding="utf-8")
    (tmp_path / "readme.docx").write_bytes(b"ignored")
    return tmp_path


def _actor(make_user):
    """The Administrator as the services see them (built inside an application context)."""
    uid = make_user("ADMINISTRATOR").id
    return lambda: AuthenticatedUser.from_model(db.session.get(User, uid))


def test_the_numbers_and_names_are_taken_out_of_a_reason():
    assert normalise("StCd-01001 (legacy MC number 3005) is not in the mapping") == (
        "StCd-# (legacy MC number #) is not in the mapping"
    )
    assert (
        normalise("The program 'Umrichter nicht bereit' has no timer")
        == "The program '…' has no timer"
    )


def test_every_program_of_a_folder_is_checked_and_nothing_is_created(app, corpus, make_user):
    user = _actor(make_user)
    with app.app_context():
        result = CorpusService.run(corpus, user())
        assert result["files"] == 7 and result["programs"] == 6  # the .docx is not read at all
        assert result["states"] == {"ready": 6}
        assert result["controllers"] == {"MC": 3, "HC": 3}
        assert [e["file"] for e in result["file_errors"]] == ["notes.st"]
        assert db.session.query(StatusCode).count() == 0
        warns = [r for r in result["reasons"] if r["kind"] == "warns"]
        assert any("parameter" in r["reason"] for r in warns)  # 14000 sets its delay by a parameter


def test_a_program_the_mapping_does_not_know_is_one_line_of_the_reasons(app, corpus, make_user):
    user = _actor(make_user)
    with app.app_context():
        db.session.query(StatusCodeMapping).filter(
            StatusCodeMapping.new_number.in_((1001, 1002))
        ).delete()
        db.session.commit()
        result = CorpusService.run(corpus, user())
        assert result["states"] == {"ready": 4, "blocked": 2}
        blocks = [r for r in result["reasons"] if r["kind"] == "blocks"]
        assert blocks[0]["programs"] == 2 and "is not in the mapping" in blocks[0]["reason"]


def test_the_report_is_a_workbook_with_one_row_per_program(app, corpus, make_user):
    user = _actor(make_user)
    with app.app_context():
        data = CorpusService.workbook(CorpusService.run(corpus, user()))
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == ["Summary", "Reasons", "Programs", "Not programs"]
    assert book["Programs"].max_row == 7 and book["Not programs"].max_row == 2


def test_the_command_prints_the_summary_and_writes_the_report(app, corpus, make_user, tmp_path):
    make_user("ADMINISTRATOR")
    out = tmp_path / "report.xlsx"
    result = app.test_cli_runner().invoke(args=["corpus-report", str(corpus), "--out", str(out)])
    assert result.exit_code == 0, result.output
    assert "7 files, 6 programs" in result.output and "ready" in result.output
    assert out.exists() and out.stat().st_size > 0
    missing = app.test_cli_runner().invoke(args=["corpus-report", str(corpus), "--user", "nobody"])
    assert missing.exit_code != 0 and "No such active user" in missing.output
